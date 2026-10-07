"""库浏览器窗口（ADR-0048 第 3 条）：在 FreeCAD 界面中筛选、看详情、放入布局，与对话面板共用场景。"""

from __future__ import annotations

import os
import time
from pathlib import Path

import pytest

from kb.library import JsonLibrary

ROOT = Path(__file__).resolve().parent.parent
REAL = JsonLibrary(ROOT / "data" / "library")


@pytest.fixture(scope="module")
def gui():
    import FreeCAD  # noqa: F401
    import FreeCADGui

    if not os.environ.get("DISPLAY"):
        pytest.skip("没有显示环境")
    FreeCADGui.showMainWindow()
    from PySide import QtWidgets

    return QtWidgets.QApplication.instance()


def _pump_until(app, cond, timeout=120):
    deadline = time.time() + timeout
    while not cond() and time.time() < deadline:
        app.processEvents()
        time.sleep(0.005)
    assert cond(), "超时"


@pytest.mark.freecad
def test_browse_filter_detail_and_place(gui):
    import FreeCAD

    from freecad_addon.gui.browser import LibraryBrowser
    from freecad_addon.gui.host import ToolsHost

    host = ToolsHost(lambda: REAL, note="组件库：真实厂商组件（测试）")
    dock = LibraryBrowser(host=host)
    try:
        dock.category.setCurrentIndex(dock.category.findData("reducer"))
        dock.text.setText("csf 25")
        assert [r.model for r in dock._rows] == ["CSF-25-100-2UH"]
        dock.table.selectRow(0)
        html = dock.detail.toHtml()
        assert "CSF-25-100-2UH" in html and "src-hd-csg-csf-gear-units" in html and "曲线 24 点" in html
        assert dock.place_selected() == "reducer"
        _pump_until(gui, lambda: "已放入" in dock.status.text() or "没有放入" in dock.status.text())
        assert "已放入 reducer" in dock.status.text()
        doc = FreeCAD.getDocument("FapLayout")
        assert any(o.Name.startswith("reducer") for o in doc.Objects)
        # 第二次放入同品类时自动编号；AI 经同一宿主放入的实例也计入
        box = []  # 界面线程不能直接调工具（会与主线程执行器互等），AI 也是在后台线程调用
        host.call_in_background("place_component", {"instance": "motor", "component_id":
                                "servo_motor.yaskawa-electric-corporation.sgm7j-02a-6"}, box.append, source="ai")
        _pump_until(gui, lambda: box)
        assert not box[0].is_error and dock.instances == {"reducer", "motor"}
        assert dock.place_selected() == "reducer2"
        _pump_until(gui, lambda: "reducer2" in dock.status.text() and "正在" not in dock.status.text())
    finally:
        dock.close()
        host.tools.close()


@pytest.mark.freecad
def test_document_groups_instances_with_port_markers_on_their_frames(gui):
    import FreeCAD

    from freecad_addon.core.scene import Scene
    from freecad_addon.fc import snapshot

    s = Scene()
    s.place("motor", REAL.get("servo_motor.yaskawa-electric-corporation.sgm7j-02a-6"))
    s.place("reducer", REAL.get("reducer.harmonic-drive-llc.csf-25-100-2uh"))
    s.connect("motor.shaft", "reducer.input_bore")
    doc = snapshot.build_document(s)
    group = doc.getObject("reducer_group")
    names = {o.Name for o in group.Group}
    assert {"reducer", "reducer__input_bore", "reducer__output_flange", "reducer__motor_flange"} <= names
    part = doc.getObject("reducer")
    assert part.FapComponent == "reducer.harmonic-drive-llc.csf-25-100-2uh" and part.FapCategory == "reducer"
    mk = doc.getObject("reducer__output_flange")
    assert mk.Label == "reducer.output_flange" and mk.FapPortType == "mechanical.flange"
    assert doc.getObject("motor__shaft").FapConnected == "是" and mk.FapConnected == "否"
    # 标记圆盘中心在端口坐标系原点
    _, origin, _ = s.port_world("reducer.output_flange")
    c = mk.Shape.Solids[0].CenterOfMass  # 圆盘
    assert (c.x, c.y) == pytest.approx(tuple(origin[:2]), abs=1e-6) and abs(c.z - origin[2]) < 1.0
    FreeCAD.closeDocument(doc.Name)


@pytest.mark.freecad
def test_manual_assembly_connect_and_live_check(gui):
    """ADR-0048 第 5 条：人选两个接口标记连接，校验窗口自动校验并着色；与 AI 用同一场景与工具。"""
    import FreeCAD
    import FreeCADGui

    from freecad_addon.core import live_check
    from freecad_addon.gui.check_panel import CheckPanel, PortConnector
    from freecad_addon.gui.host import ToolsHost

    host = ToolsHost(lambda: REAL, note="组件库：真实厂商组件（测试）")
    panel, conn = CheckPanel(host=host), PortConnector(host)
    try:
        assert panel.result["mode"] == "empty"
        box = []
        for n, (inst, cid) in enumerate((("motor", "servo_motor.yaskawa-electric-corporation.sgm7j-02a"),
                                         ("reducer", "reducer.harmonic-drive-llc.csf-25-100-2uh")), start=1):
            host.call_in_background("place_component", {"instance": inst, "component_id": cid}, box.append)
            _pump_until(gui, lambda n=n: len(box) == n)
        _pump_until(gui, lambda: len(panel.layout_state["instances"]) == 2 and panel.result["mode"] == "interface")
        FreeCADGui.Selection.clearSelection()
        assert conn.connect_selected() is False and "选两个接口标记" in conn.last  # 没选就提示
        doc = FreeCAD.getDocument("FapLayout")
        FreeCADGui.Selection.addSelection(doc.getObject("motor__shaft"))
        FreeCADGui.Selection.addSelection(doc.getObject("reducer__input_bore"))
        assert conn.selected_refs() == ["motor.shaft", "reducer.input_bore"]
        assert conn.connect_selected() is True
        _pump_until(gui, lambda: conn.last.startswith("已连接") or conn.last.startswith("没有连接"))
        assert conn.last == "已连接 motor.shaft ↔ reducer.input_bore"
        _pump_until(gui, lambda: panel.layout_state["connections"])
        r = panel.result
        assert r["mode"] == "interface" and r["overall"] == "fail"  # 光轴配键槽孔
        assert r["ports"]["motor.shaft"] == "fail"
        gdoc = FreeCADGui.getDocument("FapLayout")
        assert tuple(gdoc.getObject("motor__shaft").ShapeColor[:3]) == pytest.approx(live_check.STATUS_COLORS["fail"], abs=0.01)
        assert "C2 圆柱配合：不通过" in panel.report.toPlainText()
        # 填需求后按需求校验全部 11 项
        panel.inputs["output_torque_cont_nm"].setText("25")
        panel.inputs["output_speed_rpm"].setText("30")
        assert panel.apply_form() and panel.result["mode"] == "full"
    finally:
        panel.close()
        host.tools.close()


@pytest.mark.freecad
def test_project_is_saved_in_fcstd_and_restored(gui, tmp_path):
    """ADR-0049 第 1、2 条：布局文档就地更新；方案记录随 .FCStd 保存；打开后恢复方案，场景一致。"""
    import FreeCAD

    from freecad_addon.core import project
    from freecad_addon.fc import project_doc, snapshot
    from freecad_addon.gui.check_panel import CheckPanel
    from freecad_addon.gui.host import ToolsHost

    host = ToolsHost(lambda: REAL, note="组件库：真实厂商组件（测试）")
    panel = CheckPanel(host=host)
    try:
        box = []
        for n, (name, args) in enumerate((
                ("place_component", {"instance": "motor",
                                     "component_id": "servo_motor.yaskawa-electric-corporation.sgm7j-02a-6"}),
                ("place_component", {"instance": "reducer", "component_id": "reducer.harmonic-drive-llc.csf-25-100-2uh",
                                     "position_mm": [0, 0, 200]}),
                ("connect_ports", {"a": "motor.shaft", "b": "reducer.input_bore"})), start=1):
            host.call_in_background(name, args, box.append)
            _pump_until(gui, lambda n=n: len(box) == n)
        _pump_until(gui, lambda: panel.layout_state["connections"])
        doc = FreeCAD.getDocument(snapshot._target["name"])
        user_box = doc.addObject("Part::Box", "UserBox")  # 用户自己的对象：就地更新时保留
        host.call_in_background("place_component", {"instance": "drive", "component_id":
                                "drive.yaskawa-america-inc.sgd7s-1r6a", "position_mm": [150, 0, 0]}, box.append)
        _pump_until(gui, lambda: len(box) == 4 and len(panel.layout_state["instances"]) == 3)
        assert doc.getObject("UserBox") is user_box
        panel.set_requirement({"output_torque_cont_nm": 25, "output_speed_rpm": 30})
        rec = project.loads(project_doc.read(doc))
        assert rec["requirement"]["output_speed_rpm"] == 30 and len(rec["layout"]["instances"]) == 3
        before = panel.layout_state
        path = str(tmp_path / "joint.FCStd")
        doc.saveAs(path)
        FreeCAD.closeDocument(doc.Name)
        # 场景被清空后，打开文件恢复
        for k, inst in enumerate(("motor", "reducer", "drive"), start=len(box) + 1):
            host.call_in_background("place_component", {"instance": inst, "remove": True}, box.append)
            _pump_until(gui, lambda k=k: len(box) == k)
        opened = FreeCAD.openDocument(path)
        FreeCAD.setActiveDocument(opened.Name)
        notes = []
        assert panel.restore(notes.append) is None
        _pump_until(gui, lambda: notes)
        _pump_until(gui, lambda: len(panel.layout_state["instances"]) == 3)
        assert notes == [[]] and not project.diff(before, panel.layout_state)["changed"]
        assert snapshot._target["name"] == opened.Name and opened.getObject("motor").FapCategory == "servo_motor"
        assert panel.requirement["output_speed_rpm"] == 30
        FreeCAD.closeDocument(opened.Name)
    finally:
        snapshot.set_target(snapshot.DOC_NAME)
        panel.close()
        host.tools.close()


@pytest.mark.freecad
def test_dragging_a_part_writes_back_to_the_scene(gui):
    """ADR-0049 第 3 条：用户移动零件后场景随之改变，已连接的刚性组一起动；同步造成的改动不回写。"""
    import FreeCAD

    from freecad_addon.fc import snapshot
    from freecad_addon.gui.check_panel import CheckPanel
    from freecad_addon.gui.host import ToolsHost

    host = ToolsHost(lambda: REAL, note="组件库：真实厂商组件（测试）")
    panel = CheckPanel(host=host)
    try:
        box = []
        for n, (name, args) in enumerate((
                ("place_component", {"instance": "motor",
                                     "component_id": "servo_motor.yaskawa-electric-corporation.sgm7j-02a-6"}),
                ("place_component", {"instance": "reducer", "component_id": "reducer.harmonic-drive-llc.csf-25-100-2uh",
                                     "position_mm": [0, 0, 200]}),
                ("connect_ports", {"a": "motor.shaft", "b": "reducer.input_bore"}),
                ("place_component", {"instance": "drive", "component_id": "drive.yaskawa-america-inc.sgd7s-1r6a",
                                     "position_mm": [150, 0, 0]})), start=1):
            host.call_in_background(name, args, box.append)
            _pump_until(gui, lambda n=n: len(box) == n)
        _pump_until(gui, lambda: len(panel.layout_state["instances"]) == 3)
        gui.processEvents()
        assert panel.drag is not None and not panel.drag.pending  # 同步造成的改动没有排队
        pos = {i["instance"]: i["position_mm"] for i in panel.layout_state["instances"]}
        doc = FreeCAD.getDocument(snapshot._target["name"])
        # 用户把电机挪到 x=50：减速器与它是一个刚性组，一起挪；驱动器不动
        pl = doc.getObject("motor").Placement
        pl.Base = FreeCAD.Vector(50, 0, 0)
        doc.getObject("motor").Placement = pl
        assert "motor" in panel.drag.pending
        _pump_until(gui, lambda: {i["instance"]: i["position_mm"] for i in panel.layout_state["instances"]}["motor"]
                    == pytest.approx([50, 0, 0]))
        new = {i["instance"]: i["position_mm"] for i in panel.layout_state["instances"]}
        assert new["reducer"] == pytest.approx([pos["reducer"][0] + 50, pos["reducer"][1], pos["reducer"][2]])
        assert new["drive"] == pytest.approx(pos["drive"])
        # 拖动接口标记不回写
        mk = doc.getObject("drive__mount")
        mpl = mk.Placement
        mpl.Base = FreeCAD.Vector(999, 0, 0)
        mk.Placement = mpl
        assert not panel.drag.pending
    finally:
        panel.close()
        host.tools.close()


@pytest.mark.freecad
def test_ai_changes_become_a_proposal_to_accept_or_reject(gui):
    """ADR-0049 第 4 条：AI 一轮改动后列为提案；拒绝恢复原状，接受保留；未处理时手动改动视为接受。"""
    import sys

    sys.path.insert(0, str(ROOT / "tests"))
    from test_chat_panel import ScriptedLLM, reply, use

    from freecad_addon.gui.check_panel import CheckPanel
    from freecad_addon.gui.host import ToolsHost
    from freecad_addon.gui.panel import ChatPanel

    M = "servo_motor.yaskawa-electric-corporation.sgm7j-02a-6"
    R = "reducer.harmonic-drive-llc.csf-25-100-2uh"

    def script():
        return [reply(use(1, "place_component", {"instance": "motor", "component_id": M}),
                      use(2, "place_component", {"instance": "reducer", "component_id": R,
                                                 "position_mm": [0, 0, 200]})),
                reply(use(3, "connect_ports", {"a": "motor.shaft", "b": "reducer.input_bore"})),
                reply({"type": "text", "text": "电机与减速器已对接。"}, stop="end_turn")]

    host = ToolsHost(lambda: REAL, note="组件库：真实厂商组件（测试）")
    checks = CheckPanel(host=host)
    chat = ChatPanel(llm_factory=lambda: ScriptedLLM(script() + script() + [
        reply(use(4, "place_component", {"instance": "drive", "component_id": "drive.yaskawa-america-inc.sgd7s-1r6a",
                                         "position_mm": [150, 0, 0]})),
        reply({"type": "text", "text": "加了驱动器。"}, stop="end_turn")]), host=host)
    try:
        assert chat.review.isChecked()
        # 第一轮：拒绝 → 恢复为空
        assert chat.send("装配电机与减速器")
        _pump_until(gui, lambda: not chat.busy and host.pending is not None)
        _pump_until(gui, lambda: not checks.proposal_box.isHidden() and not checks.reject_button.isHidden())
        text = checks.proposal_text.text()
        assert "增加：motor、reducer" in text and "新连接：motor.shaft ↔ reducer.input_bore" in text
        assert "电机与减速器已对接" in text
        notes = []
        assert checks.reject(notes.append)
        _pump_until(gui, lambda: notes)
        assert notes == [[]] and host.pending is None and host.layout["instances"] == []
        _pump_until(gui, lambda: "已恢复原状" in checks.proposal_text.text())
        # 第二轮：接受 → 保留
        assert chat.send("再装一次")
        _pump_until(gui, lambda: not chat.busy and host.pending is not None)
        checks.accept_button.click()
        assert host.pending is None and len(host.layout["instances"]) == 2
        # 第三轮：提案未处理时手动改动视为接受
        assert chat.send("加一个驱动器")
        _pump_until(gui, lambda: not chat.busy and host.pending is not None)
        box = []
        host.call_in_background("place_component", {"instance": "drive", "position_mm": [180, 0, 0]}, box.append)
        _pump_until(gui, lambda: box)
        assert host.pending is None and len(host.layout["instances"]) == 3
    finally:
        chat.close()
        checks.close()
        host.tools.close()


@pytest.mark.freecad
def test_selection_panel_finds_and_places_a_candidate(gui):
    """ADR-0051 第 1 条：不经对话，在选型面板填需求、查找、放入布局；校验窗口按全部 11 项校验。"""
    import json

    from freecad_addon.core import selection
    from freecad_addon.core.system_layout import mechanical_connections
    from freecad_addon.gui.check_panel import CheckPanel
    from freecad_addon.gui.host import ToolsHost
    from freecad_addon.gui.selection_panel import SelectionPanel

    req = json.loads((ROOT / "examples" / "joint-real-requirement.json").read_text(encoding="utf-8"))
    host = ToolsHost(lambda: REAL, note="组件库：真实厂商组件（测试）")
    checks, panel = CheckPanel(host=host), SelectionPanel(host=host)
    try:
        assert panel.find() is False and "至少要填" in panel.status.text()  # 空表单
        panel.set_requirement(req)
        assert panel.find() is True
        _pump_until(gui, lambda: panel.find_button.isEnabled())
        assert panel.rows and panel.table.rowCount() == len(panel.rows)
        row = panel.selected()
        assert row is not None and row.rank == 1
        assert "SGM7J-02A" in selection.cells(row)[2]
        _pump_until(gui, lambda: checks.requirement is not None)  # 查找时的需求同步到校验窗口
        # 先放一个无关零件，放入方案时应被清掉
        box = []
        host.call_in_background("place_component", {"instance": "stray", "component_id":
                                                    "reducer.harmonic-drive-llc.csf-14-50-2uh"}, box.append)
        _pump_until(gui, lambda: box)
        assert panel.place_selected(confirm=False) is True
        _pump_until(gui, lambda: panel.find_button.isEnabled())
        assert panel.status.text().startswith("已放入布局"), panel.status.text()
        lay = host.layout
        assert sorted(i["instance"] for i in lay["instances"]) == sorted(c["instance"] for c in row.system["components"])
        ordered, _ = mechanical_connections(row.system, REAL.get)
        assert {(c["a"], c["b"]) for c in lay["connections"]} == {(c["a"], c["b"]) for c in ordered}
        boxes = {i["instance"]: i["bbox_mm"] for i in lay["instances"]}
        cats = {c["instance"]: REAL.get(c["component"])["category"] for c in row.system["components"]}
        drive = next(n for n, c in cats.items() if c == "drive")
        others = max(b["max_mm"][0] for n, b in boxes.items() if n != drive)
        assert boxes[drive]["min_mm"][0] == pytest.approx(others + selection.GAP_MM, abs=1e-6)
        _pump_until(gui, lambda: len(checks.layout_state["instances"]) == len(lay["instances"]))
        r = checks.recheck()
        assert r["mode"] == "full" and r["overall"] == "unknown"  # 只有 C3（可配置法兰）待确认
        assert not [c for c in r["checks"] if c["status"] == "fail"]
    finally:
        panel.close()
        checks.close()
        host.tools.close()


@pytest.mark.freecad
def test_bom_panel_follows_the_scene_and_exports(gui, tmp_path):
    """ADR-0051 第 2 条：BOM 面板随场景刷新；导出的 CSV 与 export_system 一致。"""
    from freecad_addon.core.live_check import system_from_layout
    from freecad_addon.gui.bom_panel import BomPanel
    from freecad_addon.gui.host import ToolsHost
    from mcp_server import tools

    host = ToolsHost(lambda: REAL, note="组件库：真实厂商组件（测试）")
    dock = BomPanel(host=host)
    try:
        assert dock.result["rows"] == [] and not dock.csv_button.isEnabled()
        assert dock.export("csv", str(tmp_path / "x.csv")) is None and "还没有零件" in dock.status.text()
        box = []
        for n, (inst, cid) in enumerate((("motor", "servo_motor.yaskawa-electric-corporation.sgm7j-02a-6"),
                                         ("reducer", "reducer.harmonic-drive-llc.csf-25-100-2uh"),
                                         ("drive", "drive.yaskawa-america-inc.sgd7s-1r6a")), start=1):
            host.call_in_background("place_component", {"instance": inst, "component_id": cid}, box.append)
            _pump_until(gui, lambda n=n: len(box) == n)
        _pump_until(gui, lambda: dock.table.rowCount() == 3)
        assert dock.csv_button.isEnabled() and dock.total.text().startswith("共 3 种、3 件")
        assert dock.custom.isHidden()  # 还没连接，没有定制要求
        host.call_in_background("connect_ports", {"a": "motor.mount_flange", "b": "reducer.motor_flange"},
                                box.append)
        _pump_until(gui, lambda: not dock.custom.isHidden())
        assert "须按" in dock.custom.toPlainText() and "第 65 页" in dock.custom.toPlainText()
        path = dock.export("csv", str(tmp_path / "bom.csv"))
        system, _ = system_from_layout(dock.layout_state, REAL.get)
        want = tools.export_system(REAL, system, format="bom_csv")["content"]
        assert Path(path).read_text(encoding="utf-8") == want
        host.call_in_background("place_component", {"instance": "drive", "remove": True}, box.append)
        _pump_until(gui, lambda: dock.table.rowCount() == 2)
        assert not dock.custom.isHidden()
    finally:
        dock.close()
        host.tools.close()


@pytest.mark.freecad
def test_swap_suggestion_fixes_a_failed_shaft_fit(gui):
    """ADR-0052：光轴电机配键槽孔减速器 C2 不通过；换件建议给出带键电机，换上后 C2 通过。"""
    from freecad_addon.gui.check_panel import CheckPanel
    from freecad_addon.gui.host import ToolsHost

    plain = "servo_motor.yaskawa-electric-corporation.sgm7j-02a"
    keyed = "servo_motor.yaskawa-electric-corporation.sgm7j-02a-6"
    host = ToolsHost(lambda: REAL, note="组件库：真实厂商组件（测试）")
    panel = CheckPanel(host=host)
    try:
        box = []
        steps = [("place_component", {"instance": "motor", "component_id": plain}),
                 ("place_component", {"instance": "reducer", "component_id": "reducer.harmonic-drive-llc.csf-25-100-2uh"}),
                 ("connect_ports", {"a": "motor.mount_flange", "b": "reducer.motor_flange"}),
                 ("connect_ports", {"a": "motor.shaft", "b": "reducer.input_bore"})]
        for n, (name, args) in enumerate(steps, start=1):
            host.call_in_background(name, args, box.append)
            _pump_until(gui, lambda n=n: len(box) == n)
        _pump_until(gui, lambda: len(panel.layout_state["connections"]) == 2)
        r = panel.recheck()
        assert {c["check"]: c["status"] for c in r["checks"]}["C2"] == "fail"
        sb = panel.swap
        assert sb.instance.currentText() == "motor"  # 默认选结论最差的
        assert sb.find() is True
        _pump_until(gui, lambda: not sb.busy)
        assert sb.suggestions and sb.suggestions[0]["component_id"] == keyed
        assert "C2 圆柱配合：不通过 → 通过" in sb.detail.text()
        assert sb.swap_selected() is True
        _pump_until(gui, lambda: not sb.busy)
        assert sb.status.text().startswith("已换上"), sb.status.text()
        _pump_until(gui, lambda: any(i["component"] == keyed for i in panel.layout_state["instances"])
                    and len(panel.layout_state["connections"]) == 2)
        r = panel.recheck()
        assert {c["check"]: c["status"] for c in r["checks"]}["C2"] == "pass"
    finally:
        panel.close()
        host.tools.close()


@pytest.mark.freecad
def test_sentence_is_parsed_into_the_form_without_searching(gui):
    """ADR-0054：一句话需求经解析（录制回放）只填表单，不自动查找；缺密钥时提示可直接填表。"""
    import json

    from engine import requirement_parse as rp
    from freecad_addon.gui.host import ToolsHost
    from freecad_addon.gui.selection_panel import SelectionPanel
    from ingest.llm_extract import RecordedClient

    joint2 = json.loads((ROOT / "examples" / "joint2-requirement.json").read_text(encoding="utf-8"))
    host = ToolsHost(lambda: REAL, note="组件库：真实厂商组件（测试）")
    panel = SelectionPanel(host=host, client_factory=lambda: RecordedClient(rp.RECORDINGS_DIR))
    try:
        panel.sentence.setPlainText(joint2["statement"])
        assert panel.parse() is True
        _pump_until(gui, lambda: not panel.parsing)
        assert panel.parse_note.text().startswith("解析完成"), panel.parse_note.text()
        req, err = panel.requirement()
        assert err is None and req == {k: v for k, v in joint2.items() if k != "id"}
        assert panel.rows == [] and panel.table.rowCount() == 0  # 没有自动查找
        panel.clear()
        assert panel.extra == {} and panel.inputs["output_torque_cont_nm"].text() == ""

        def no_key():
            raise RuntimeError("未设置 ANTHROPIC_API_KEY")

        panel._client_factory = no_key
        panel.sentence.setPlainText("随便一句")
        assert panel.parse() is False and "可以直接在下表填数" in panel.parse_note.text()
    finally:
        panel.close()
        host.tools.close()


@pytest.mark.freecad
def test_enterprise_component_is_marked_in_browser_and_bom(gui, tmp_path):
    """ADR-0055：企业库叠加在公共库上，企业补过的组件在浏览器与 BOM 中标为企业库。"""
    from freecad_addon.gui.bom_panel import BomPanel
    from freecad_addon.gui.browser import LibraryBrowser
    from freecad_addon.gui.host import ToolsHost
    from kb.enterprise import EnterpriseLibrary
    from kb.library import LayeredLibrary

    ent = EnterpriseLibrary(tmp_path / "ent", create=True)
    spec = tmp_path / "spec.pdf"
    spec.write_bytes(b"%PDF fake")
    doc = ent.register_document(spec, "供应商规格书", "张工")["id"]
    cid = "reducer.harmonic-drive-llc.csf-25-100-2uh"
    ent.set_value(cid, "params/backlash_arcmin", {"value": 1, "method": "manual", "confidence": 1, "reviewed": True,
                                                  "source": {"doc": doc, "page": 2, "note": "原文“1 arcmin”；录入人：张工"}},
                  "张工", REAL.get)
    lib = LayeredLibrary(ent, REAL)
    host = ToolsHost(lambda: lib, note="企业库（测试）")
    browser, bom = LibraryBrowser(host=host), BomPanel(host=host)
    try:
        browser.text.setText("csf 25-100")
        rows = [r.id for r in browser._rows]
        browser.table.selectRow(rows.index(cid))
        assert "来自企业库" in browser.detail.toHtml()
        box = []
        host.call_in_background("place_component", {"instance": "reducer", "component_id": cid}, box.append)
        _pump_until(gui, lambda: bom.table.rowCount() == 1)
        assert bom.table.item(0, len(bom_view_headers()) - 1).text() == "企业"
    finally:
        browser.close()
        bom.close()
        host.tools.close()


def bom_view_headers():
    from freecad_addon.core import bom_view

    return bom_view.HEADERS


@pytest.mark.freecad
def test_fill_dialog_writes_enterprise_value_and_rechecks(gui, tmp_path):
    """M9 第二步：校验窗口“补数据”为减速器补电机法兰分度圆，写入企业库并重新校验；缺依据时不写。"""
    from freecad_addon.core import library_choice
    from freecad_addon.gui.check_panel import CheckPanel
    from freecad_addon.gui.host import ToolsHost
    from kb.enterprise import EnterpriseLibrary
    from kb.library import LayeredLibrary

    ent = EnterpriseLibrary(tmp_path / "ent", create=True)
    lib = LayeredLibrary(ent, REAL)
    host = ToolsHost(lambda: lib, note="企业库（测试）")
    panel = CheckPanel(host=host)
    old_settings = library_choice.SETTINGS
    library_choice.SETTINGS = tmp_path / "settings.json"  # 不写用户目录
    try:
        box = []
        steps = [("place_component", {"instance": "motor", "component_id":
                                      "servo_motor.yaskawa-electric-corporation.sgm7j-02a-6"}),
                 ("place_component", {"instance": "reducer", "component_id": "reducer.harmonic-drive-llc.csf-25-100-2uh"}),
                 ("connect_ports", {"a": "motor.mount_flange", "b": "reducer.motor_flange"})]
        for n, (name, args) in enumerate(steps, start=1):
            host.call_in_background(name, args, box.append)
            _pump_until(gui, lambda n=n: len(box) == n)
        _pump_until(gui, lambda: len(panel.layout_state["connections"]) == 1)
        panel.recheck()
        sb = panel.swap
        sb.instance.setCurrentText("reducer")
        dlg = sb.fill_dialog()
        # 待确认所涉端口（电机法兰，以及尚未连接的输入孔）的字段排在前面
        assert dlg.row().target.split("/")[:2] in (["ports", "motor_flange"], ["ports", "input_bore"])
        targets = [r.target for r in dlg.rows]
        assert targets.index("ports/motor_flange/pcd_mm") < targets.index("params/backlash_arcmin")
        dlg.field.setCurrentIndex(dlg.field.findData("ports/motor_flange/pcd_mm"))
        dlg.value.setText("70")
        assert dlg.save() is False and "资料、页码、原文与录入人" in dlg.status.text()  # 缺依据不写
        assert ent.get("reducer.harmonic-drive-llc.csf-25-100-2uh") is None
        spec = tmp_path / "order.pdf"
        spec.write_bytes(b"%PDF order")
        dlg.actor.setText("张工")
        doc = dlg.register(str(spec), "HD 订货确认单")
        assert doc and dlg.doc.currentData() == doc
        dlg.page.setValue(1)
        dlg.quote.setText("PCD 70")
        assert dlg.save() is True
        comp = lib.get("reducer.harmonic-drive-llc.csf-25-100-2uh")
        assert comp["ports"][1]["spec"]["pcd_mm"]["value"] == 70.0 and lib.origin(comp["id"]) == "enterprise"
        assert library_choice.read_settings(tmp_path / "settings.json")["actor"] == "张工"
        r = panel.recheck()
        c3 = next(c for c in r["checks"] if c["check"] == "C3")
        assert "缺少分度圆直径" not in [f["message"] for f in c3.get("findings", [])]
    finally:
        library_choice.SETTINGS = old_settings
        panel.close()
        host.tools.close()


@pytest.mark.freecad
def test_import_dialog_extracts_confirms_and_writes(gui, tmp_path):
    """M9 第三步：校验窗口“导入规格书”——Excel 选型表 → 模拟 AI 抽取 → 代码核对 → 取消一项 → 写入企业库并重新校验。"""
    import copy

    from openpyxl import Workbook

    from freecad_addon.core import library_choice
    from freecad_addon.gui.check_panel import CheckPanel
    from freecad_addon.gui.host import ToolsHost
    from ingest.llm_simulate import SimulatedClient, simulated_response
    from ingest.synthetic import datasheets, row_label
    from kb.enterprise import EnterpriseLibrary
    from kb.library import LayeredLibrary

    rid = "reducer.harmonic-drive-llc.csf-25-100-2uh"
    sheet = next(s for s in datasheets(0, ["reducer"]) if s.answer["layout"] == "two_col")
    answer = copy.deepcopy(sheet.answer)
    answer["component"]["model"] = REAL.get(rid)["model"]  # 让模拟响应报这个零件的型号
    a = answer
    wb = Workbook()
    wb.remove(wb.active)
    for page in sorted({e["page"] for e in a["fields"]}):
        ws = wb.create_sheet(f"P{page}")
        ws.append([f"Harmonic Drive {a['component']['model']}"])
        for e in a["fields"]:
            if e["page"] == page:
                ws.append([row_label(e, a["lang"], a["layout"]), e["text"]])
    xlsx = tmp_path / "selection.xlsx"
    wb.save(xlsx)

    ent = EnterpriseLibrary(tmp_path / "ent", create=True)
    lib = LayeredLibrary(ent, REAL)
    host = ToolsHost(lambda: lib, note="企业库（测试）")
    panel = CheckPanel(host=host)
    old_settings = library_choice.SETTINGS
    library_choice.SETTINGS = tmp_path / "settings.json"
    try:
        box = []
        host.call_in_background("place_component", {"instance": "reducer", "component_id": rid}, box.append)
        _pump_until(gui, lambda: len(box) == 1)
        _pump_until(gui, lambda: panel.layout_state.get("instances"))
        panel.recheck()
        sb = panel.swap
        sb.instance.setCurrentText("reducer")
        dlg = sb.import_dialog(lambda: SimulatedClient(simulated_response(answer)))
        assert dlg.extract() is False and "规格书文件" in dlg.status.text()
        dlg.file.setText(str(xlsx))
        dlg.actor.setText("张工")
        assert dlg.extract() is True
        _pump_until(gui, lambda: not dlg.busy)
        assert dlg.rows and dlg.save_button.isEnabled(), dlg.status.text()
        chosen = dlg.chosen()
        assert "params/rated_torque_nm" in chosen and all(t.startswith(("params/", "ports/")) for t in chosen)
        dlg.set_checked("params/rated_torque_nm", False)
        assert "params/rated_torque_nm" not in dlg.chosen()
        assert dlg.save() is True
        comp = lib.get(rid)
        assert lib.origin(rid) == "enterprise"
        written = set(dlg.written)
        assert written and "params/rated_torque_nm" not in written
        tgt = min(t for t in written if t.startswith("params/"))
        pv = comp["params"][tgt.split("/")[1]]
        assert pv["method"] == "extracted" and pv["reviewed"] and "确认人：张工" in pv["source"]["note"]
        assert comp["params"]["rated_torque_nm"] == REAL.get(rid)["params"]["rated_torque_nm"]  # 未勾选的不变
        assert panel.recheck() is not None
    finally:
        library_choice.SETTINGS = old_settings
        panel.close()
        host.tools.close()


@pytest.mark.freecad
def test_new_part_wizard_creates_selectable_part(gui, tmp_path):
    """M9 第四步：新建零件向导——Excel 规格书 → 模拟 AI 抽取 → 补录一项 → 生成零件存入企业库 → 放入布局、出 BOM。"""
    from openpyxl import Workbook

    from engine.export import bom_rows
    from freecad_addon.core import library_choice
    from freecad_addon.gui.host import ToolsHost
    from freecad_addon.gui.import_dialog import NewPartDialog
    from ingest.llm_simulate import SimulatedClient, simulated_response
    from ingest.synthetic import datasheets, row_label
    from kb.enterprise import EnterpriseLibrary
    from kb.library import LayeredLibrary

    sheet = next(s for s in datasheets(0, ["reducer"]) if s.answer["layout"] == "two_col")
    a = sheet.answer
    wb = Workbook()
    wb.remove(wb.active)
    for page in sorted({e["page"] for e in a["fields"]}):
        ws = wb.create_sheet(f"P{page}")
        ws.append([a["title"]])
        for e in a["fields"]:
            if e["page"] == page:
                ws.append([row_label(e, a["lang"], a["layout"]), e["text"]])
    xlsx = tmp_path / "new.xlsx"
    wb.save(xlsx)

    ent = EnterpriseLibrary(tmp_path / "ent", create=True)
    lib = LayeredLibrary(ent, REAL)
    host = ToolsHost(lambda: lib, note="企业库（测试）")
    old_settings = library_choice.SETTINGS
    library_choice.SETTINGS = tmp_path / "settings.json"
    try:
        dlg = NewPartDialog(lib, lambda: SimulatedClient(simulated_response(a)))
        dlg.category.setCurrentIndex(dlg.category.findData("reducer"))
        assert dlg.add_manual() is False and "先导入规格书" in dlg.status.text()
        dlg.file.setText(str(xlsx))
        dlg.actor.setText("张工")
        assert dlg.extract() is True
        _pump_until(gui, lambda: not dlg.busy)
        assert dlg.rows and all(r.writable for r in dlg.rows), dlg.status.text()
        assert dlg.vendor.text() and dlg.model.text()  # 由规格书预填
        for r in dlg.rows:
            dlg.set_checked(r.target, True)
        # 补录一项 AI 没抽到的值（须写页码与原文）
        got = {r.target for r in dlg.rows}
        extra = next(r for r in dlg.m_rows if r.target.startswith("params/") and r.kind == "number"
                     and not r.enum and r.target not in got)
        dlg.m_field.setCurrentIndex(dlg.m_field.findData(extra.target))
        dlg.m_value.setText("1.5")
        assert dlg.add_manual() is False  # 缺页码、原文
        dlg.m_page.setValue(1)
        dlg.m_quote.setText("printed 1.5")
        assert dlg.add_manual() is True and extra.target in dlg.manual
        assert dlg.save() is True, dlg.status.text()
        cid = dlg.saved["id"]
        assert lib.origin(cid) == "enterprise" and cid in {c["id"] for c in lib.all("reducer")}
        assert lib.get(cid)["params"][extra.target.split("/")[1]]["value"] == 1.5
        box = []
        host.call_in_background("place_component", {"instance": "reducer", "component_id": cid}, box.append)
        _pump_until(gui, lambda: len(box) == 1)
        system = {"id": "s", "components": [{"instance": "reducer", "component": cid}], "connections": []}
        assert bom_rows(system, lib.get)[0]["component_id"] == cid
    finally:
        library_choice.SETTINGS = old_settings
        host.tools.close()


@pytest.mark.freecad
def test_analysis_list_dialog(gui, monkeypatch):
    """ADR-0056：分析模块一览列出各模块的可用状态。"""
    import types

    import analysis
    from analysis.common import result
    from freecad_addon.gui.analysis_list import AnalysisListDialog

    fake = types.ModuleType("fake")
    fake.MODULE = {"id": "demo", "name": "演示模块", "python": ["no_such_pkg_xyz"], "install": "pip install demo"}
    fake.run = lambda setup: result.make("demo", "1", setup, {})
    monkeypatch.setattr(analysis, "MODULE_PATHS", ("x",))
    monkeypatch.setattr(analysis, "_load", lambda p: fake)
    dlg = AnalysisListDialog()
    assert dlg.table.rowCount() == 1
    assert dlg.table.item(0, 0).text() == "演示模块" and "no_such_pkg_xyz" in dlg.table.item(0, 1).text()
    assert dlg.table.item(0, 2).text() == "pip install demo"


@pytest.mark.freecad
def test_mbd_panel_computes_and_sends_to_selection(gui):
    """M10 第二步：多体动力学窗口——填机械臂、计算、看关节力矩表，送去选型窗口。"""
    pytest.importorskip("mujoco")
    from freecad_addon.gui.host import ToolsHost
    from freecad_addon.gui.mbd_panel import MbdPanel
    from freecad_addon.gui.selection_panel import SelectionPanel

    host = ToolsHost(lambda: REAL, note="组件库：真实厂商组件（测试）")
    sel = SelectionPanel(host=host)
    panel = MbdPanel(selection=lambda: sel, host=host)
    try:
        assert panel.compute() is False and "请填写" in panel.status.text()  # 质量不估：没填不算
        for key, text in {"base_height_mm": "160", "l1_mm": "300", "m1_kg": "3", "l2_mm": "250", "m2_kg": "1.2",
                          "elbow_kg": "1", "payload_kg": "1.5", "j1_deg": "90", "j2_deg": "30", "move_s": "1.2",
                          "duration_s": "1.5"}.items():
            panel.inputs[key].setText(text)
        assert panel.compute() is True
        _pump_until(gui, lambda: not panel.busy)
        assert panel.result and panel.result["status"] == "ok", panel.status.text()
        assert panel.table.rowCount() == 3 and panel.joint.currentText() == "j2"  # 最吃力的关节
        assert panel.send_to_selection() is True
        req, err = sel.requirement()
        assert err is None and req["output_torque_peak_nm"] > req["output_torque_cont_nm"] > 0
        assert req["load_inertia_kgm2"] > 0  # 表单外一并带入
        assert "多体动力学" in sel.parse_note.text()
        # 机构
        panel.model.setCurrentIndex(panel.model.findData("C-LNK-4BAR"))
        assert panel.compute() is True
        _pump_until(gui, lambda: not panel.busy)
        assert panel.table.rowCount() == 1 and panel.joint.currentText() == "crank"
    finally:
        panel.close()
        sel.close()
        host.tools.close()


@pytest.mark.freecad
def test_fea_panel_static_color_and_fatigue(gui, tmp_path):
    """M10 第三步：有限元窗口——FreeCAD 中的轴（Part 圆柱）→ 选端面固定、另一端加扭矩 → 计算 → 着色 → 疲劳。"""
    pytest.importorskip("gmsh")
    import FreeCAD
    import FreeCADGui
    import Part

    from analysis.common.solvers import find_ccx
    from freecad_addon.gui.fea_panel import FeaPanel

    if find_ccx() is None:
        pytest.skip("没有 CalculiX")
    doc = FreeCAD.newDocument("FeaTest")
    try:
        obj = doc.addObject("Part::Feature", "Shaft")
        obj.Shape = Part.makeCylinder(17.5, 100)
        doc.recompute()
        FreeCADGui.ActiveDocument.ActiveView.viewIsometric()
        panel = FeaPanel(mbd_result=lambda: {"data": {"series": {"j2": {"drive": [0, 100, 350, 200, 0]}}}})
        assert panel.set_part(obj) and len(panel.faces) == 3
        bottom = next(f"Face{i + 1}" for i, f in enumerate(obj.Shape.Faces)
                      if f.Surface.__class__.__name__ == "Plane" and abs(f.CenterOfMass.z) < 1e-6)
        top = next(f"Face{i + 1}" for i, f in enumerate(obj.Shape.Faces)
                   if f.Surface.__class__.__name__ == "Plane" and abs(f.CenterOfMass.z - 100) < 1e-6)
        assert panel.add_load("fixed", panel.face_ids([bottom]))
        assert panel.add_load("torque", panel.face_ids([top]), "350")
        panel.material.setCurrentIndex(panel.material.findData("45-QT"))
        panel.mesh.setText("4")
        assert panel.compute() is True
        _pump_until(gui, lambda: not panel.busy, timeout=300)
        assert panel.result and panel.result["status"] == "ok", panel.status.text()
        assert panel.result["data"]["vm_max_mpa"] == pytest.approx(72.0, rel=0.05)
        assert panel.node is not None and "安全系数" in panel.summary.text()
        assert panel.load_ref() == 350
        assert panel.compute_fatigue() is False and "β" in panel.fatigue_label.text()  # 不给缺省值
        panel.surface.setText("0.92")
        panel.size.setText("0.85")
        panel.series_kind.setCurrentIndex(panel.series_kind.findData("mbd"))
        panel.mbd_joint.setText("j2")
        assert panel.compute_fatigue() is True and panel.fatigue["data"]["cycles_per_block"] >= 1
        panel.close()
    finally:
        FreeCAD.closeDocument(doc.Name)


@pytest.mark.freecad
def test_fea_panel_thermal(gui):
    """M10 第四步：有限元窗口的温度场——一端定温、另一端发热，热平衡与公式一致，按温度着色。"""
    pytest.importorskip("gmsh")
    import FreeCAD
    import Part

    from analysis.common.solvers import find_ccx
    from freecad_addon.gui.fea_panel import FeaPanel

    if find_ccx() is None:
        pytest.skip("没有 CalculiX")
    doc = FreeCAD.newDocument("ThermTest")
    try:
        obj = doc.addObject("Part::Feature", "Wall")
        obj.Shape = Part.makeBox(100, 20, 20)
        doc.recompute()
        panel = FeaPanel()
        assert panel.set_part(obj)
        end0 = next(f"Face{i + 1}" for i, f in enumerate(obj.Shape.Faces) if abs(f.CenterOfMass.x) < 1e-6)
        end1 = next(f"Face{i + 1}" for i, f in enumerate(obj.Shape.Faces) if abs(f.CenterOfMass.x - 100) < 1e-6)
        panel.kind.setCurrentIndex(panel.kind.findData("thermal"))
        panel.material.setCurrentIndex(panel.material.findData("45-QT"))
        assert panel.add_load("temperature", panel.face_ids([end0]), "20")
        assert panel.add_load("heat_flux", panel.face_ids([end1]), "10")
        assert panel.add_load("convection", panel.face_ids([end1]), "abc") is False  # 对流写 h, 环境温度
        panel.mesh.setText("5")
        assert panel.compute() is True
        _pump_until(gui, lambda: not panel.busy, timeout=300)
        assert panel.result and panel.result["status"] == "ok", panel.status.text()
        assert panel.result["data"]["t_max_c"] == pytest.approx(70.2, rel=0.01)
        assert "℃" in panel.summary.text() and panel.node is not None
        panel.close()
    finally:
        FreeCAD.closeDocument(doc.Name)


@pytest.mark.freecad
def test_optimize_dialog_makes_shaft(gui):
    """M10 第四步：设计优化对话框——两段直径寻优，满足安全系数，生成最优的轴到 FreeCAD。"""
    pytest.importorskip("optuna")
    import FreeCAD

    from analysis.common.solvers import find_ccx
    from freecad_addon.gui.optimize_dialog import OptimizeDialog, parse_groups, parse_segments

    if find_ccx() is None:
        pytest.skip("没有 CalculiX")
    assert parse_segments("30x40, 35×12") == [[30, 40], [35, 12]] and parse_groups("0+4, 2") == [[0, 4], [2]]
    dlg = OptimizeDialog()
    assert dlg.start() is False  # 没填不算
    dlg.segments.setText("30x40, 40x60, 30x30")
    dlg.torque.setText("350")
    dlg.groups.setText("0+2")
    dlg.low.setText("26")
    dlg.high.setText("33")
    dlg.material.setCurrentIndex(dlg.material.findData("45-QT"))
    dlg.trials.setValue(6)
    assert dlg.start() is True
    _pump_until(gui, lambda: not dlg.busy, timeout=600)
    assert dlg.result and dlg.result["status"] == "ok", dlg.status.text()
    assert dlg.table.rowCount() >= 2 and dlg.make_button.isEnabled()
    doc = FreeCAD.newDocument("OptTest")
    try:
        obj = dlg.make_shaft()
        assert obj is not None and obj.Shape.isValid() and abs(obj.Shape.BoundBox.ZLength - 130) < 1e-6
    finally:
        FreeCAD.closeDocument(doc.Name)


@pytest.mark.freecad
def test_topo_dialog_makes_plate(gui):
    """M11 第一步：拓扑优化对话框——悬臂算例 → 密度图 → 板件生成到 FreeCAD，体积与轮廓面积 × 板厚一致。"""
    import FreeCAD

    from freecad_addon.gui.topo_dialog import TopoDialog

    dlg = TopoDialog()
    dlg.preset.setCurrentIndex(dlg.preset.findData("cantilever"))
    assert (dlg.nelx.value(), dlg.nely.value()) == (64, 32)
    dlg.nelx.setValue(32)
    dlg.nely.setValue(16)
    dlg.length.setText("64")
    dlg.thickness.setText("abc")
    assert dlg.start() is False
    dlg.thickness.setText("4")
    assert dlg.start() is True
    _pump_until(gui, lambda: not dlg.busy, timeout=120)
    assert dlg.result and dlg.result["status"] == "ok", dlg.status.text()
    assert dlg.image.pixmap() is not None and dlg.make_button.isEnabled()
    doc = FreeCAD.newDocument("TopoTest")
    try:
        obj = dlg.make_plate()
        assert obj is not None and obj.Shape.isValid()
        assert abs(obj.Shape.Volume - dlg.result["data"]["volume_mm3"]) < 1e-3 * dlg.result["data"]["volume_mm3"]
        assert abs(obj.Shape.BoundBox.ZLength - 4) < 1e-9
    finally:
        FreeCAD.closeDocument(doc.Name)


@pytest.mark.freecad
def test_cam_dialog(gui, tmp_path):
    """M11 第二步：数控编程对话框——车削生成两个程序并保存；铣削特征；打开程序文件检查。"""
    from freecad_addon.gui.cam_dialog import CamDialog

    dlg = CamDialog()
    assert dlg.turn_machine.count() >= 1 and dlg.mill_machine.count() >= 1 and dlg.kienzle.count() > 10
    dlg.tabs.setCurrentIndex(0)
    dlg.segments.setText("一段")
    assert dlg.start() is False
    dlg.segments.setText("30x40, 40x60, 30x30")
    dlg.kienzle.setCurrentIndex(dlg.kienzle.findData("C45E"))
    assert dlg.start() is True
    _pump_until(gui, lambda: not dlg.busy, timeout=120)
    assert dlg.result and dlg.result["data"]["verdict"] == "ok", dlg.status.text()
    assert "O1001" in dlg.gcode.toPlainText() and dlg.save_button.isEnabled()
    saved = dlg.save(str(tmp_path))
    assert [Path(p).name for p in saved] == ["O1001.nc", "O1002.nc"]
    dlg.tabs.setCurrentIndex(1)
    dlg.box.setText("0 0 60 40")
    dlg.features.setPlainText("型腔 10 10 50 30 0 3")
    assert dlg.start() is True
    _pump_until(gui, lambda: not dlg.busy, timeout=120)
    assert dlg.result["data"]["verdict"] == "warn"  # 尖角型腔转角留料
    assert dlg.open_file(saved[0]) and dlg.tabs.currentIndex() == 2
    dlg.check_machine.setCurrentIndex(dlg.check_machine.findData("EX-LATHE"))
    assert dlg.start() is True
    _pump_until(gui, lambda: not dlg.busy, timeout=120)
    assert dlg.result["data"]["mode"] == "check" and dlg.result["data"]["total_time_s"] > 0
    assert not dlg.save_button.isEnabled()


@pytest.mark.freecad
def test_cfd_dialog(gui):
    """M11 第三步：流体分析窗口——没装 OpenFOAM 时说明安装方法；装了时方块进虚拟风洞，按散热系数着色。"""
    pytest.importorskip("gmsh")
    import FreeCAD
    import Part

    import analysis
    from freecad_addon.gui.cfd_dialog import CfdDialog

    dlg = CfdDialog()
    st = next(m for m in analysis.modules() if m["id"] == "cfd")
    if not st["available"]:
        assert not dlg.run_button.isEnabled() and "OpenFOAM" in dlg.status.text() and dlg.compute() is False
        return
    doc = FreeCAD.newDocument("CfdTest")
    try:
        obj = doc.addObject("Part::Feature", "Block")
        obj.Shape = Part.makeBox(100, 40, 20)
        doc.recompute()
        assert dlg.compute() is False  # 没选零件
        assert dlg.set_part(obj) and len(dlg.faces) == 6
        dlg.velocity.setText("2")
        dlg.heat.setChecked(True)
        dlg.wall_t.setText("60")
        dlg.cells.setValue(20000)
        dlg.level.setValue(1)
        dlg.iterations.setValue(300)
        assert dlg.compute() is True
        _pump_until(gui, lambda: not dlg.busy, timeout=600)
        assert dlg.result and dlg.result["status"] == "ok", dlg.status.text()
        assert dlg.result["data"]["heat_w"] > 0 and dlg.table.rowCount() == 6
        vals = dlg.face_values("h_w_m2k")
        assert len(vals) == 6 and all(v is not None for v in vals)
        assert len(obj.ViewObject.DiffuseColor) == 6
    finally:
        FreeCAD.closeDocument(doc.Name)


@pytest.mark.freecad
def test_twin_panel(gui, tmp_path):
    """M10 第五步：数字孪生窗口——布局中的减速器 + 运行记录 CSV → 已用寿命、剩余、假如；缺参数时待确认。"""
    import copy

    from freecad_addon.gui.host import ToolsHost
    from freecad_addon.gui.twin_panel import TwinPanel
    from kb.library import LayeredLibrary

    rid = "reducer.harmonic-drive-llc.csf-25-100-2uh"

    class Top:  # 测试数据：在 CSF-25 上加额定寿命等（不进零件库）
        def __init__(self):
            c = copy.deepcopy(REAL.get(rid))
            src = {"doc": "src-ent-test", "page": 1, "note": "测试数据"}
            c["params"]["rated_life_h"] = {"value": 7000, "condition": "L10", "source": src, "method": "manual",
                                           "confidence": 1, "reviewed": True}
            c["params"]["rated_input_speed_rpm"] = {"value": 2000, "source": src, "method": "manual", "confidence": 1,
                                                    "reviewed": True}
            self.c = c

        def get(self, cid):
            return self.c if cid == rid else None

        def all(self, category=None):
            return [self.c] if category in (None, "reducer") else []

    csv = tmp_path / "run.csv"
    csv.write_text("instance,hours,output_torque_avg_nm,output_torque_max_nm,output_speed_rpm\n"
                   + "reducer,16,30,60,15\n" * 30, encoding="utf-8")
    for lib, expect in ((LayeredLibrary(Top(), REAL), "正常"), (REAL, "待确认")):
        host = ToolsHost(lambda lib=lib: lib, note="测试")
        panel = TwinPanel(host=host)
        try:
            assert panel.compute() is False and "没有减速器" in panel.status.text()
            box = []
            host.call_in_background("place_component", {"instance": "reducer", "component_id": rid}, box.append)
            _pump_until(gui, lambda box=box: len(box) == 1)
            assert panel.load_csv(str(csv)) and len(panel.records) == 30
            panel.load_factor.setText("0.8")
            assert panel.compute() is True
            assert panel.table.rowCount() == 1 and panel.table.item(0, 2).text() == expect
            if expect == "正常":
                p = panel.result["data"]["parts"][0]
                assert p["used_ratio"] > 0 and panel.result["data"]["whatif"]["parts"][0]["remaining_h"] > p["remaining_h"]
            else:
                assert "补数据" in panel.status.text()
        finally:
            panel.close()
            host.tools.close()


@pytest.mark.freecad
def test_publish_dialog_to_factory(gui, tmp_path):
    """M10 第 6c 项：发布到工厂——当前布局 → 本机假工厂（用工厂的 cad_link 检查内容）→ 第 1 版。"""
    import importlib.util
    import json
    import threading
    from http.server import BaseHTTPRequestHandler, HTTPServer

    from freecad_addon.core import library_choice
    from freecad_addon.gui.check_panel import CheckPanel
    from freecad_addon.gui.host import ToolsHost
    from freecad_addon.gui.publish_dialog import PublishDialog

    spec = importlib.util.spec_from_file_location("cad_link", ROOT / "factory" / "digital" / "hub" / "cad_link.py")
    cad = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(cad)
    got = []

    class Hub(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_POST(self):
            raw = self.rfile.read(int(self.headers["Content-Length"]))
            if self.path == "/api/login":
                out = {"token": "tok"}
            elif self.path == "/api/files":
                out = {"url": "/api/files/s", "sha256": "s", "size": len(raw)}
            else:
                body = json.loads(raw)
                cad.validate(body)
                got.append(body)
                out = {"code": body["code"], "revision": len(got), "name": body["name"]}
            b = json.dumps(out).encode()
            self.send_response(200)
            self.send_header("Content-Length", str(len(b)))
            self.end_headers()
            self.wfile.write(b)

    srv = HTTPServer(("127.0.0.1", 0), Hub)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    host = ToolsHost(lambda: REAL, note="测试")
    panel = CheckPanel(host=host)
    old = library_choice.SETTINGS
    library_choice.SETTINGS = tmp_path / "settings.json"
    try:
        dlg = PublishDialog(host, panel)
        dlg.name.setText("第 2 关节")
        dlg.author.setText("张工")
        assert dlg.publish() is False and "没有零件" in dlg.status.text()
        box = []
        for n, (inst, cid) in enumerate((("motor", "servo_motor.yaskawa-electric-corporation.sgm7j-02a-6"),
                                         ("reducer", "reducer.harmonic-drive-llc.csf-25-100-2uh")), start=1):
            host.call_in_background("place_component", {"instance": inst, "component_id": cid}, box.append)
            _pump_until(gui, lambda n=n: len(box) == n)
        dlg.url.setText(f"http://127.0.0.1:{srv.server_port}")
        assert dlg.publish() is True
        _pump_until(gui, lambda: not dlg.busy)
        assert dlg.result == {"code": dlg.code.text(), "revision": 1, "name": "第 2 关节"}, dlg.status.text()
        body = got[0]
        assert len(body["bom"]) == 2 and body["files"] and body["files"][0]["name"] == "layout.step"
        assert library_choice.read_settings(tmp_path / "settings.json")["factory_url"].startswith("http://127.0.0.1")
    finally:
        library_choice.SETTINGS = old
        srv.shutdown()
        panel.close()
        host.tools.close()
