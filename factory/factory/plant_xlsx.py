# -*- coding: utf-8 -*-
"""工厂清单的 Excel 模板与导入（freecad-ai-parts ADR-0064 第 3 条）。

模板 = 演示工厂填好的样例；企业照着改，导入时逐格读取，出错说明“哪张表第几行哪一列”。
表：说明、工厂、工作站、设备、工序、供应商、客户；第 2 版（第 45 轮）加物料、工艺路线、物料清单、成品、库存、测量特性。
区域、阶段、物料类别在表里写中文（也认英文代号）。
"""
import datetime as _dt
import re

from . import plant as P

AREA_ZH = {"office": "办公", "design": "设计", "machining": "机加工", "quality": "质检", "warehouse": "仓库",
           "logistics": "物流", "ai": "AI", "field": "现场"}
STAGE_ZH = {"blank": "下料", "rough": "粗加工", "qt": "调质", "finish": "精加工", "keyway": "键槽", "gear": "齿形",
            "harden": "淬火", "grind": "磨削", "inspect": "检验", "assembly": "装配", "other": "其他"}
KIND_ZH = {"raw": "原材料", "buy": "外购件", "make": "自制件", "sub": "部件", "fg": "成品"}
WEEK = ("周一", "周二", "周三", "周四", "周五", "周六", "周日")
YES, NO = "是", "否"

# 工厂表：(清单路径, 项目, 说明)
FACTORY_ROWS = [
    ("id", "工厂编号", "小写字母、数字和 -，如 acme-sewing"),
    ("company.name", "公司名称", "ERPNext 里建的公司名"),
    ("company.abbr", "公司缩写", "字母数字、不超过 6 位；仓库名后缀用它"),
    ("company.country", "国家", "ERPNext 的英文国名，如 China、United States"),
    ("company.currency", "币种", "如 CNY、USD"),
    ("company.timezone", "时区", "如 Asia/Shanghai、America/New_York"),
    ("demo_data", "装演示数据", "是：装演示工厂的物料、客户、订单（只用于试用）；否：只建工厂本身"),
    ("stores.raw", "原材料库（设备编号）", "设备表里的仓库编号：物流小车从这里取毛坯"),
    ("stores.finished", "成品库（设备编号）", "设备表里的仓库编号：完工的件送到这里"),
    ("floor.length_m", "车间长（米）", "x 方向"),
    ("floor.width_m", "车间宽（米）", "y 方向"),
    ("floor.aisles_y", "横向主通道 y（米）", "物流小车走的通道；多条用逗号分开，如 9, 18.5"),
    ("floor.cross_x", "纵向通道 x（米）", "连接各条横向通道"),
    ("shift.periods", "班次时段", "如 08:00-12:00, 13:00-17:00"),
    ("shift.workdays", "工作日", "如 周一, 周二, 周三, 周四, 周五"),
    ("shift.max_overtime_min", "每天加班上限（分钟）", "0–480"),
    ("kpi.on_time", "准时交付目标", "0–1 的小数，如 0.95"),
    ("kpi.oee", "设备综合效率目标", "0–1 的小数"),
    ("kpi.fpy", "一次合格率目标", "0–1 的小数"),
]

# 列表表：(列名, 清单键, 类型, 说明)；类型：str num int yesno area stage ids
WS_COLS = [
    ("名称", "name", "str", "ERPNext 工作站名，工序表按它引用"),
    ("台数", "count", "int", ""),
    ("折旧费率（每小时）", "cost.0", "num", "币种同公司"),
    ("人工费率（每小时）", "cost.1", "num", ""),
    ("能耗费率（每小时）", "cost.2", "num", ""),
    ("主电机功率（kW）", "power.0", "num", "可空；工艺能耗估算用"),
    ("传动效率", "power.1", "num", "0–1，可空"),
]
UNIT_COLS = [
    ("编号", "id", "str", "小写字母、数字和 -，如 cnc-01"),
    ("区域", "area", "area", "/".join(AREA_ZH.values())),
    ("中文名", "name", "str", ""),
    ("英文名", "name_en", "str", "可空"),
    ("工作站", "workstation", "str", "所属工作站（工作站表的名称）；不在车间里的可空"),
    ("x（米）", "layout.0", "num", "设备中心"),
    ("y（米）", "layout.1", "num", ""),
    ("宽（米）", "layout.2", "num", "x 方向"),
    ("深（米）", "layout.3", "num", "y 方向"),
    ("高（米）", "height_m", "num", "3D 车间用；可空"),
    ("上下料点 x", "dock.0", "num", "可空：取设备靠最近主通道一侧"),
    ("上下料点 y", "dock.1", "num", ""),
    ("停靠点 x", "home.0", "num", "物流小车必填"),
    ("停靠点 y", "home.1", "num", ""),
    ("看板顺序", "_board", "int", "在看板“车间实况”里显示的顺序；空 = 不显示"),
    ("同时加工件数", "capacity", "int", "可空 = 1"),
    ("换型（分钟）", "setup_min", "num", "可空"),
    ("每件刀具磨损", "wear_per_part", "num", "0–1，如 0.01 = 每件磨 1%；可空"),
    ("换刀（分钟）", "tool_change_min", "num", "可空"),
    ("主轴转速（rpm）", "spindle.0", "num", "可空"),
    ("进给（mm/min）", "spindle.1", "num", "可空"),
    ("虚拟 PLC 程序", "plc", "str", "可空"),
]
OP_COLS = [
    ("工序名", "name", "str", "工艺路线里用的名称"),
    ("工作站", "workstation", "str", ""),
    ("可用设备", "units", "ids", "设备编号，多台用逗号分开"),
    ("阶段", "stage", "stage", "/".join(STAGE_ZH.values()) + "；可空"),
    ("检验", "inspection", "yesno", "是 / 否"),
]
SUP_COLS = [
    ("代号", "_code", "str", "物料类别代号，如 castings、steel"),
    ("名称", "name", "str", ""),
    ("交货期（天）", "lead", "int", ""),
]
CUS_COLS = [("名称", "_name", "str", "")]

# 第 2 版：产品几节（一行一条；工艺路线、物料清单按名称、父件连续几行成一组）
ITEM_COLS = [
    ("编号", "code", "str", "物料编号，物料清单、库存、测量特性按它引用"),
    ("名称", "name", "str", "中英文名"),
    ("类别", "kind", "kind", "/".join(KIND_ZH.values())),
    ("单位", "uom", "str", "Nos、Kg、Litre……"),
    ("单价", "price", "num", "原材料、外购件；不知道就空着（不估）"),
    ("供应商", "supplier", "str", "供应商表的代号；原材料、外购件必填"),
    ("检验模板", "inspection", "str", "可空"),
    ("零件库编号", "library", "str", "可空"),
    ("备注", "note", "str", "可空"),
]
ROUTE_COLS = [("工艺路线", "name", "str", "同一条路线的几道工序连续写，名称相同"),
              ("工序", "op", "str", "工序表里的工序名"), ("每件分钟", "min", "num", "")]
BOM_COLS = [("父件", "item", "str", "同一父件的几个子件连续写"), ("工艺路线", "routing", "str", "父件的工艺路线（同组写一次即可）"),
            ("子件", "child", "str", "物料编号"), ("用量", "qty", "num", "每个父件用多少")]
PROD_COLS = [("成品", "item", "str", "可接单的成品（物料编号）"), ("售价", "price", "num", "可空")]
STOCK_COLS = [("物料", "item", "str", ""), ("期初库存", "opening", "num", "可空"), ("安全库存", "safety", "num", "可空"),
              ("关键物料顺序", "key", "int", "看板“关键物料”显示的物料填 1、2、3…（按数字排序），其余空着")]
MEAS_COLS = [("物料", "item", "str", "测哪个物料"), ("工序", "operation", "str", "在哪道工序测"),
             ("代号", "code", "str", "特性代号，如 backlash"), ("名称", "name", "str", ""), ("单位", "unit", "str", "如 mm、arcmin、N·m"),
             ("名义", "nominal", "num", ""), ("下限", "lo", "num", ""), ("上限", "hi", "num", ""),
             ("散布σ", "sigma", "num", "仿真车间合成测量值的散布（示例值）")]
PRODUCT_SHEETS = (("物料", ITEM_COLS), ("工艺路线", ROUTE_COLS), ("物料清单", BOM_COLS), ("成品", PROD_COLS),
                  ("库存", STOCK_COLS), ("测量特性", MEAS_COLS))

GROUP = {"cost": "费率", "power": "主电机功率 / 传动效率", "layout": "x / y / 宽 / 深", "dock": "上下料点 x / y",
         "home": "停靠点 x / y", "spindle": "主轴转速 / 进给"}
LIST_SHEETS = (("工作站", "workstations", WS_COLS), ("设备", "units", UNIT_COLS), ("工序", "operations", OP_COLS),
               ("供应商", "suppliers", SUP_COLS), ("客户", "customers", CUS_COLS))
FIELD_SHEET = {"workstations": "工作站", "units": "设备", "operations": "工序"}


# ---------------------------------------------------------------- 清单 → 表格行
def _get(d, path):
    for k in path.split("."):
        if d is None:
            return None
        if isinstance(d, list):
            d = d[int(k)] if int(k) < len(d) else None
        else:
            d = d.get(k)
    return d


def _fmt_time(t):
    return str(t)[:5]


def _cell_out(kind, v):
    if v is None:
        return None
    if kind == "area":
        return AREA_ZH.get(v, v)
    if kind == "stage":
        return STAGE_ZH.get(v, v)
    if kind == "yesno":
        return YES if v else NO
    if kind == "ids":
        return ", ".join(v)
    if kind == "kind":
        return KIND_ZH.get(v, v)
    return v


def _factory_value(p, path):
    v = _get(p, path)
    if path == "demo_data":
        return YES if v else NO
    if v is None:
        return None
    if path == "floor.aisles_y":
        return ", ".join(_num_text(x) for x in v)
    if path == "shift.periods":
        return ", ".join("{}-{}".format(_fmt_time(a), _fmt_time(b)) for a, b in v)
    if path == "shift.workdays":
        return ", ".join(WEEK[d] for d in v)
    return v


def _num_text(x):
    return str(int(x)) if float(x).is_integer() else str(x)


def rows_of(p):
    """清单 → {表名: [行（列表）]}（不含表头）。"""
    out = {"工厂": [[label, _factory_value(p, path), note] for path, label, note in FACTORY_ROWS]}
    board = {u: i + 1 for i, u in enumerate(p.get("dashboard") or [])}
    for sheet, key, cols in LIST_SHEETS:
        items = p.get(key) or ([] if key != "suppliers" else {})
        if key == "suppliers":
            items = [{"_code": c, **v} for c, v in items.items()]
        if key == "customers":
            items = [{"_name": c} for c in items]
        rows = []
        for it in items:
            it = dict(it, _board=board.get(it.get("id"))) if key == "units" else it
            rows.append([_cell_out(kind, _get(it, k)) for _, k, kind, _ in cols])
        out[sheet] = rows
    if p.get("items"):
        out.update(_product_rows(p))
    return out


def _product_rows(p):
    """第 2 版：产品几节 → 表格行"""
    out = {"物料": [[_cell_out(kind, it.get(k)) for _, k, kind, _ in ITEM_COLS] for it in p.get("items") or []]}
    out["工艺路线"] = [[r["name"], op, m] for r in p.get("routings") or [] for op, m in r["steps"]]
    out["物料清单"] = [[b["item"], b["routing"] if i == 0 else None, c, q] for b in p.get("boms") or []
                    for i, (c, q) in enumerate(b["components"])]
    out["成品"] = [[x["item"], x.get("price")] for x in p.get("products") or []]
    st = p.get("stock") or {}
    keys = list(p.get("key_materials") or [])
    codes = list(dict.fromkeys(list((st.get("opening") or {}).keys()) + list((st.get("safety") or {}).keys()) + keys))
    out["库存"] = [[c, (st.get("opening") or {}).get(c), (st.get("safety") or {}).get(c), keys.index(c) + 1 if c in keys else None]
                 for c in codes]
    out["测量特性"] = [[m["item"], m["operation"], c["code"], c["name"], c.get("unit"), c["nominal"], c["lo"], c["hi"], c["sigma"]]
                    for m in p.get("measurements") or [] for c in m["chars"]]
    return out


# ---------------------------------------------------------------- 模板
README = [
    "工厂清单 Excel 模板（ADR-0064）",
    "",
    "这份模板已经按演示工厂（问渠减速器厂）填好，作为样例。请按自己的车间改：",
    "1. 工厂：公司、车间尺寸、通道、班次、指标目标。",
    "2. 工作站：ERPNext 里的工作站（同类设备一组），台数与每小时费率。",
    "3. 设备：每台设备一行。编号唯一；平面位置是设备中心的 x、y 和占地宽、深（米），原点在车间左下角。",
    "   物流小车（区域填“物流”）要填停靠点。看板顺序决定“车间实况”里显示哪些设备。",
    "4. 工序：工序名与工艺路线一致；“可用设备”填设备编号，多台用逗号分开。",
    "5. 供应商、客户：可空（只在装演示数据时用）。",
    "6. 产品（第 2 版，可整组不填）：物料、工艺路线、物料清单、成品、库存、测量特性六张表。",
    "   工艺路线、物料清单按名称、父件连续几行成一组；外购件的单价、交期不知道就空着，系统标“待确认”，不替你估。",
    "",
    "填好后转成清单并校验：",
    "    python -m factory.plant import 填好的.xlsx plant.yaml",
    "有错会列出“哪张表第几行哪一列”；改好再导入。导入成功后用 install.sh --plant plant.yaml 安装。",
    "",
    "灰色斜体的第 2 行是填写说明，导入时跳过。空格子表示“不填”，不要写 0 代替。",
]


def write_template(path, p=None):
    """把清单（缺省演示工厂）写成 Excel 模板。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    p = p or P.export_demo()
    rows = rows_of(p)
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    for line in README:
        ws.append([line])
    ws.column_dimensions["A"].width = 110
    ws["A1"].font = Font(bold=True, size=14)

    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="3B5B7A")
    note = Font(italic=True, color="808080")

    ws = wb.create_sheet("工厂")
    ws.append(["项目", "值", "说明"])
    for r in rows["工厂"]:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill = head, fill
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 3).font = note
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 22, 46, 60
    yn = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(yn)
    yn.add("B{}".format(1 + [x[0] for x in FACTORY_ROWS].index("demo_data") + 1))

    for sheet, _key, cols in LIST_SHEETS:
        ws = wb.create_sheet(sheet)
        ws.append([c[0] for c in cols])
        ws.append([c[3] for c in cols])
        for c in ws[1]:
            c.font, c.fill = head, fill
            c.alignment = Alignment(wrap_text=True)
        for c in ws[2]:
            c.font = note
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows[sheet]:
            ws.append(r)
        for i, c in enumerate(cols):
            ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = max(10, min(36, len(c[0]) * 2 + 4))
            choice = {"area": AREA_ZH.values(), "stage": STAGE_ZH.values(), "yesno": (YES, NO)}.get(c[2])
            if choice:
                dv = DataValidation(type="list", formula1='"{}"'.format(",".join(choice)), allow_blank=True)
                ws.add_data_validation(dv)
                col = ws.cell(1, i + 1).column_letter
                dv.add("{0}3:{0}500".format(col))
        ws.freeze_panes = "B3"
    for sheet, cols in PRODUCT_SHEETS:
        if sheet not in rows:
            continue
        ws = wb.create_sheet(sheet)
        ws.append([c[0] for c in cols])
        ws.append([c[3] for c in cols])
        for c in ws[1]:
            c.font, c.fill = head, fill
        for c in ws[2]:
            c.font = note
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows[sheet]:
            ws.append(r)
        for i, c in enumerate(cols):
            ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = max(10, min(40, len(c[0]) * 2 + 6))
            if c[2] == "kind":
                dv = DataValidation(type="list", formula1='"{}"'.format(",".join(KIND_ZH.values())), allow_blank=True)
                ws.add_data_validation(dv)
                dv.add("{0}3:{0}2000".format(ws.cell(1, i + 1).column_letter))
        ws.freeze_panes = "B3"
    wb.save(path)


# ---------------------------------------------------------------- 导入
class _Cell:
    def __init__(self, errs, sheet, row, col):
        self.errs, self.where = errs, "{}表第 {} 行“{}”".format(sheet, row, col)

    def bad(self, msg):
        self.errs.append("{}：{}".format(self.where, msg))


def _blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _parse(kind, v, cell):
    """一格 → 清单里的值；空格子 → None。"""
    if _blank(v):
        return None
    if kind == "str":
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).strip()
    if kind in ("num", "int"):
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            try:
                v = float(str(v).strip().replace("，", "."))
            except ValueError:
                return cell.bad("要填数字（现在是“{}”）".format(v))
        if kind == "int":
            if not float(v).is_integer():
                return cell.bad("要填整数（现在是 {}）".format(v))
            return int(v)
        return int(v) if isinstance(v, int) else float(v)
    if kind == "yesno":
        t = str(v).strip().lower()
        if t in ("是", "y", "yes", "true", "1", "√"):
            return True
        if t in ("否", "n", "no", "false", "0"):
            return False
        return cell.bad("填“是”或“否”")
    if kind in ("area", "stage", "kind"):
        table = {"area": AREA_ZH, "stage": STAGE_ZH, "kind": KIND_ZH}[kind]
        t = str(v).strip()
        for code, zh in table.items():
            if t in (code, zh) or t.lower() == code:
                return code
        return cell.bad("可选：{}".format("、".join(table.values())))
    if kind == "ids":
        return [x for x in re.split(r"[,，、;；\s]+", str(v).strip()) if x]
    raise ValueError(kind)


def _put(d, key, val):
    """key 是 'a' 或 'a.0'；把值放进 d。"""
    if "." not in key:
        d[key] = val
        return
    k, i = key.split(".")
    d.setdefault(k, {})[int(i)] = val


def _pack(d, sizes, cell_of):
    """把 {k: {0: v, 1: v}} 变成列表；整组空则去掉，缺一项报错。"""
    for k, n in sizes.items():
        part = d.pop(k, None)
        if not part:
            continue
        vals = [part.get(i) for i in range(n)]
        if any(v is None for v in vals):
            cell_of(k).bad("这一组要填全（{} 项）".format(n))
            continue
        d[k] = vals


def _time(v, cell):
    if isinstance(v, _dt.time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, _dt.datetime):
        return v.strftime("%H:%M:%S")
    t = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})[:：](\d{2})(?:[:：](\d{2}))?", t)
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        cell.bad("时间写成 08:00 这样（现在是“{}”）".format(t))
        return None
    return "{:02d}:{}:{}".format(int(m.group(1)), m.group(2), m.group(3) or "00")


def _factory(ws, errs):
    labels = {label: (path, label) for path, label, _ in FACTORY_ROWS}
    got, rows = {}, {}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or _blank(row[0]):
            continue
        name = str(row[0]).strip()
        if name not in labels:
            errs.append("工厂表第 {} 行：不认识的项目“{}”".format(r, name))
            continue
        got[labels[name][0]] = row[1] if len(row) > 1 else None
        rows[labels[name][0]] = r
    p = {"version": 1, "company": {}, "stores": {}, "floor": {}, "shift": {}, "kpi": {}}
    for path, label, _ in FACTORY_ROWS:
        cell = _Cell(errs, "工厂", rows.get(path, "?"), label)
        v = got.get(path)
        if path not in rows:
            errs.append("工厂表：缺少项目“{}”".format(label))
            continue
        if _blank(v):
            if path != "demo_data":
                cell.bad("必填")
            continue
        if path == "demo_data":
            val = _parse("yesno", v, cell)
        elif path in ("id",) or path.startswith(("company.", "stores.")):
            val = _parse("str", v, cell)
        elif path == "floor.aisles_y":
            parts = [x for x in re.split(r"[,，、;；\s]+", str(v).strip()) if x]
            val = [_parse("num", x, cell) for x in parts]
            val = None if any(x is None for x in val) else val
        elif path == "shift.periods":
            val = []
            for seg in re.split(r"[,，、;；]+", str(v).strip()):
                ab = [x for x in re.split(r"\s*[-–—~～至]\s*", seg.strip()) if x]
                if len(ab) != 2:
                    cell.bad("时段写成 08:00-12:00，多段用逗号分开")
                    val = None
                    break
                val.append([_time(ab[0], cell), _time(ab[1], cell)])
        elif path == "shift.workdays":
            val = []
            for x in re.split(r"[,，、;；\s]+", str(v).strip()):
                if not x:
                    continue
                x = x.replace("星期", "周")
                if x in WEEK:
                    val.append(WEEK.index(x))
                elif x.isdigit() and 1 <= int(x) <= 7:
                    val.append(int(x) - 1)
                else:
                    cell.bad("写周一到周日，逗号分开（现在有“{}”）".format(x))
        elif path == "shift.max_overtime_min":
            val = _parse("int", v, cell)
        else:
            val = _parse("num", v, cell)
        if val is None:
            continue
        if "." in path:
            a, b = path.split(".")
            p[a][b] = val
        else:
            p[path] = val
    p["demo_data"] = bool(p.get("demo_data", False))
    return p, rows


def _list_sheet(ws, sheet, key, cols, errs):
    hdr = [(_blank(c) and "" or str(c).strip()) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    idx = {}
    for name, k, kind, _ in cols:
        if name in hdr:
            idx[k] = (hdr.index(name), kind, name)
        elif not k.startswith(("dock", "home", "power", "spindle")) and k not in ("height_m", "plc", "_board"):
            errs.append("{}表：缺少列“{}”".format(sheet, name))
    items, rownum = [], []
    for r, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if not row or all(_blank(v) for v in row):
            continue
        d = {}
        for k, (ci, kind, name) in idx.items():
            cell = _Cell(errs, sheet, r, name)
            val = _parse(kind, row[ci] if ci < len(row) else None, cell)
            if val is not None:
                _put(d, k, val)
        _pack(d, {"cost": 3, "power": 2, "layout": 4, "dock": 2, "home": 2, "spindle": 2},
              lambda k, r=r: _Cell(errs, sheet, r, GROUP.get(k, k)))
        items.append(d)
        rownum.append(r)
    return items, rownum


def read_workbook(path):
    """Excel → (清单 dict, 错误列表, 行号表)。"""
    from openpyxl import load_workbook

    errs = []
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        return None, ["打不开这个 Excel：{}".format(e)], {}
    for need in ("工厂", "工作站", "设备", "工序"):
        if need not in wb.sheetnames:
            errs.append("缺少表“{}”".format(need))
    if errs:
        return None, errs, {}
    p, frow = _factory(wb["工厂"], errs)
    rows = {"工厂": frow}
    for sheet, key, cols in LIST_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        items, rn = _list_sheet(wb[sheet], sheet, key, cols, errs)
        rows[key] = rn
        if key == "units":
            board = sorted(((u.pop("_board"), i) for i, u in enumerate(items) if u.get("_board") is not None))
            for u in items:
                u.pop("_board", None)
            if board:
                p["dashboard"] = [items[i]["id"] for _, i in board if items[i].get("id")]
            p["units"] = items
        elif key == "suppliers":
            if items:
                p["suppliers"] = {it.pop("_code", ""): it for it in items}
        elif key == "customers":
            if items:
                p["customers"] = [it["_name"] for it in items if it.get("_name")]
        else:
            if key == "operations":
                for o in items:
                    o["inspection"] = bool(o.get("inspection"))
                    if not o["inspection"]:
                        o.pop("inspection")
            p[key] = items
    if "物料" in wb.sheetnames:
        _read_products(wb, p, errs)
    if p.get("items") or any(p.get(k) for k in ("routings", "boms", "products")):
        p["version"] = 2
    else:
        p["version"] = 1 if "物料" not in wb.sheetnames else 2
    return p, errs, rows


def _sheet_rows(wb, sheet, cols, errs):
    """一张产品表 → [(行号, {键: 值})]（跳过表头、说明行和空行）"""
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    hdr = [(_blank(c) and "" or str(c).strip()) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    idx = {}
    for name, k, kind, _ in cols:
        if name in hdr:
            idx[k] = (hdr.index(name), kind, name)
        else:
            errs.append("{}表：缺少列“{}”".format(sheet, name))
    out = []
    for r, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if not row or all(_blank(v) for v in row):
            continue
        d = {}
        for k, (ci, kind, name) in idx.items():
            v = _parse(kind, row[ci] if ci < len(row) else None, _Cell(errs, sheet, r, name))
            if v is not None:
                d[k] = v
        out.append((r, d))
    return out


def _read_products(wb, p, errs):
    p["items"] = [d for _, d in _sheet_rows(wb, "物料", ITEM_COLS, errs)]
    routes = {}
    for r, d in _sheet_rows(wb, "工艺路线", ROUTE_COLS, errs):
        if not d.get("name") or not d.get("op") or d.get("min") is None:
            errs.append("工艺路线表第 {} 行：工艺路线、工序、每件分钟都要填".format(r))
            continue
        routes.setdefault(d["name"], []).append([d["op"], d["min"]])
    p["routings"] = [{"name": k, "steps": v} for k, v in routes.items()]
    boms = {}
    for r, d in _sheet_rows(wb, "物料清单", BOM_COLS, errs):
        if not d.get("item") or not d.get("child") or d.get("qty") is None:
            errs.append("物料清单表第 {} 行：父件、子件、用量都要填".format(r))
            continue
        b = boms.setdefault(d["item"], {"item": d["item"], "routing": None, "components": []})
        if d.get("routing"):
            b["routing"] = d["routing"]
        b["components"].append([d["child"], d["qty"]])
    p["boms"] = list(boms.values())
    p["products"] = [dict({"item": d["item"]}, **({"price": d["price"]} if d.get("price") is not None else {}))
                     for _, d in _sheet_rows(wb, "成品", PROD_COLS, errs) if d.get("item")]
    st, keys = {"opening": {}, "safety": {}}, []
    for _, d in _sheet_rows(wb, "库存", STOCK_COLS, errs):
        for k in ("opening", "safety"):
            if d.get("item") and d.get(k) is not None:
                st[k][d["item"]] = d[k]
        if d.get("item") and d.get("key") is not None:
            keys.append((d["key"], d["item"]))
    p["stock"] = st
    if keys:
        p["key_materials"] = [i for _, i in sorted(keys)]
    meas = {}
    for _, d in _sheet_rows(wb, "测量特性", MEAS_COLS, errs):
        m = meas.setdefault((d.get("item"), d.get("operation")), {"item": d.get("item"), "operation": d.get("operation"), "chars": []})
        m["chars"].append({k: d[k] for k in ("code", "name", "unit", "nominal", "lo", "hi", "sigma") if k in d})
    if meas:
        p["measurements"] = list(meas.values())


def _where(err, rows):
    """validate() 的 “units[3].layout：…” → “设备表第 6 行（layout）：…”。"""
    m = re.match(r"(workstations|units|operations)\[(\d+)\](?:\.(\w+))?：(.*)", err)
    if m and m.group(1) in rows and int(m.group(2)) < len(rows[m.group(1)]):
        key, i, field, msg = m.group(1), int(m.group(2)), m.group(3), m.group(4)
        cols = {k: n for n, k, _, _ in dict(LIST_SHEETS_BY_KEY)[key]}
        name = GROUP.get(field) or cols.get(field)
        col = "“{}”".format(name) if name else ""
        return "{}表第 {} 行{}：{}".format(FIELD_SHEET[key], rows[key][i], col, msg)
    m = re.match(r"([a-z_]+\.[a-z_]+|id|version)：(.*)", err)
    if m:
        label = {path: label for path, label, _ in FACTORY_ROWS}.get(m.group(1))
        if label:
            r = rows.get("工厂", {}).get(m.group(1))
            return "工厂表{}“{}”：{}".format("第 {} 行".format(r) if r else "", label, m.group(2))
    return err


LIST_SHEETS_BY_KEY = [(key, cols) for _, key, cols in LIST_SHEETS]


def import_workbook(path):
    """Excel → 校验过的清单 dict；有错抛 PlantError（错误按表、行、列说明）。"""
    p, errs, rows = read_workbook(path)
    if errs:
        raise P.PlantError(errs)
    errs = [_where(e, rows) for e in P.validate(p)]
    if errs:
        raise P.PlantError(errs)
    return p
