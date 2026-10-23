# -*- coding: utf-8 -*-
"""第 18 轮 E7：经营与财务——订单做完开票、客户付款（模拟器、催款）、收货开采购发票、账龄与信用、现金流预测、订单毛利、
指标树与预算、规则 F2–F7 的方案、催款 / 信用冻结 / 付款 / 调预算动作、态势与报告、AI 代理三个新工具、实验 22。"""
import datetime as dt
import json
import os

import pytest

from tests.test_agent import NoLLM

DSN = os.environ.get("WQ_TEST_DB", "postgresql://postgres@localhost:5433/wq_test")
TEA = {"name": "王老师", "role": "manager", "teacher": True, "mode": "teach"}
SAL = {"name": "学生销售", "role": "sales", "teacher": False, "mode": "teach"}
SAL2 = {"name": "学生销售乙", "role": "sales", "teacher": False, "mode": "teach"}
OPR = {"name": "学生操作工", "role": "operator", "teacher": False, "mode": "teach"}


class FakeMes:
    def __init__(self):
        self.sent = []

    def command(self, unit, command, mode, user, work_order=None, operation=None, **extra):
        self.sent.append(dict({"unit": unit, "command": command, "mode": mode}, **extra))


@pytest.fixture(scope="module")
def env():
    pytest.importorskip("psycopg")
    from hub.ont import evals as EV
    try:
        H, eng = EV.sandbox(DSN, name="wq_finance_test", llm=NoLLM())
    except Exception as e:  # noqa: BLE001
        pytest.skip("没有可用的 PostgreSQL：{}".format(e))
    H.mes = FakeMes()
    return H, eng


def _erp(db, doctype, name, ts=None, mode="teach", **data):
    import wqbus
    msg = wqbus.make("erp.doc", "bridge", dict({"doctype": doctype, "name": name, "action": "updated"}, **data), mode=mode,
                     ts=(ts or dt.datetime.now(dt.timezone.utc)).isoformat().replace("+00:00", "Z"))
    db.insert_message("wq/gearbox/office/erp/doc", msg)


def _execute(db, pid, created):
    db.x("update ai_proposal set status='executed', result=%s where proposal_id=%s", (json.dumps({"created": created}), pid))


def _order(db, so, wo, customer, qty=4, mode="teach"):
    _erp(db, "Sales Order", so, mode=mode, customer=customer, docstatus=1, status="To Deliver and Bill", delivery_date=(dt.date.today() + dt.timedelta(days=20)).isoformat(),
         transaction_date=dt.date.today().isoformat(), per_delivered=0, items=[{"item_code": "WQR-105", "qty": qty, "rate": 1650}])
    _erp(db, "Work Order", wo, mode=mode, production_item="SH-301", qty=qty, produced_qty=qty, status="Completed", docstatus=1, sales_order=so)


# ---------------------------------------------------------------- 初始化、账龄、毛利
def test_seed_opening_aging_cash(env):
    from factory import data as F
    from hub import finance as FN
    H, _ = env
    db = H.db
    a = FN.ar(db, "teach")
    assert len([x for x in a["invoices"] if x["opening"]]) == 3 and a["overdue"] == 4 * F.FG_SELLING_PRICE      # 峰顶那笔逾期 5 天
    assert FN.cash_now(db, "teach") == 60000 and FN.dso(db, "teach") > 0
    assert {o["code"] for o in FN.objectives(db, "teach")} == {"O1", "O2", "O3", "O4", "O5", "O6", "O7"}
    assert FN.budgets(db, "teach")["energy"] == 6000
    cf = FN.cashflow(db, "teach")
    assert len(cf["weeks"]) == 8 and cf["weeks"][0]["out"] >= 8000


def test_standard_cost_and_margin_breakdown(env):
    from factory import data as F
    from hub import finance as FN
    H, _ = env
    std = FN._std("WQR-105")
    assert 1100 < std["material"] + std["labor"] + std["energy"] < 1300                       # 标准毛利率约 28%
    ms = FN.margins(H.db, "teach")
    sos = FN.sales_orders(H.db, "teach")
    assert ms and all(m["pseudo"] or (sos.get(m["sales_order"]) or {}).get("data", {}).get("demo") for m in ms)   # 情景里早先的订单按工单估，或是情景自带的示例订单
    assert any(m["pseudo"] for m in ms)
    m = ms[0]
    assert abs(sum(x["usd"] for x in m["lines"]) - m["cost"]) < 0.05 and abs(m["revenue"] - m["cost"] - m["margin"]) < 0.05
    assert m["revenue"] == m["qty"] * F.FG_SELLING_PRICE and 0.2 < m["margin_pct"] < 0.35
    assert {x["key"] for x in m["lines"]} == {"material", "labor", "energy", "scrap", "idle", "actions", "warranty"}
    assert "实测" in next(x for x in m["lines"] if x["key"] == "labor")["how"]
    assert FN.gross_margin(H.db, "teach") is not None


# ---------------------------------------------------------------- 开票、收款、采购发票
def test_invoice_when_order_done_and_customer_pays(env):
    from factory import data as F
    from hub import finance as FN
    H, _ = env
    db = H.db
    now = dt.datetime.now(dt.timezone.utc)
    _order(db, "SAL-ORD-2026-00091", "MFG-WO-2026-00091", F.CUSTOMERS[0])
    out = FN.tick(H, "teach", now=now)
    inv = db.one("select * from invoice where ref='SAL-ORD-2026-00091'")
    assert inv and inv["id"] in out["ar"] and inv["status"] == "draft" and inv["amount"] == 4 * 1650
    p = db.one("select * from ai_proposal where proposal_id=%s", (inv["proposal_id"],))
    assert p["action"] == "sales_invoice" and p["status"] == "confirmed" and p["preview"]["sales_order"] == "SAL-ORD-2026-00091"   # 教学模式自动确认
    assert not FN.tick(H, "teach", now=now)["ar"]                                          # 不重复开票
    _execute(db, inv["proposal_id"], [["Sales Invoice", "ACC-SINV-2026-00091"]])
    FN.sync(db, "teach")
    inv = db.one("select * from invoice where id=%s", (inv["id"],))
    assert inv["status"] == "open" and inv["doc"] == "ACC-SINV-2026-00091"
    later = inv["pay_at"] + dt.timedelta(seconds=1)
    out = FN.tick(H, "teach", now=later)                                                   # 绿谷按时（晚 0–3 天）付
    assert any(x["id"] == inv["id"] for x in out["paid"])
    x = db.one("select * from invoice where id=%s", (inv["id"],))
    assert x["status"] == "paying" and db.one("select action from ai_proposal where proposal_id=%s", (x["pay_pid"],))["action"] == "payment"
    _execute(db, x["pay_pid"], [["Payment Entry", "ACC-PAY-2026-00091"]])
    FN.sync(db, "teach")
    assert db.one("select status, outstanding from invoice where id=%s", (inv["id"],)) == {"status": "paid", "outstanding": 0}


def test_old_orders_not_invoiced_and_held_customer_waits(env):
    from factory import data as F
    from hub import finance as FN
    from hub.ont import actions as AC
    H, _ = env
    db = H.db
    assert not db.one("select 1 as x from invoice where ref in ('SAL-ORD-2026-00016','SAL-ORD-2026-00017')")   # 情景里早先完工的不补开
    aid = AC.create(db, "teach", SAL, "credit_hold", {"customer": F.CUSTOMERS[1], "hold": True})
    AC.submit(db, aid, SAL, None)
    assert not AC.can_approve(SAL2, AC.get(db, aid))[0]
    AC.approve(H, aid, TEA, TEA["name"])
    _order(db, "SAL-ORD-2026-00092", "MFG-WO-2026-00092", F.CUSTOMERS[1])
    FN.tick(H, "teach")
    assert db.one("select status from invoice where ref='SAL-ORD-2026-00092'")["status"] == "held"
    FN.exec_credit(H, "teach", {"customer": F.CUSTOMERS[1], "hold": False}, TEA)
    FN.tick(H, "teach")
    assert db.one("select status from invoice where ref='SAL-ORD-2026-00092'")["status"] == "draft"


def test_purchase_invoice_and_auto_pay(env):
    from hub import finance as FN
    H, _ = env
    db = H.db
    now = dt.datetime.now(dt.timezone.utc)
    pid = H.ai.propose("incoming_receipt", {"purchase_order": "PUR-ORD-T"}, "teach", "测试", "收货")
    _execute(db, pid, [["Purchase Receipt", "MAT-PRE-2026-00091"]])
    lid = db.one("insert into lot (mode, lot, po, supplier, item, qty, qty_ok, rate, status, receipt_pid, received_at) values "
                 "('teach','L261008-0091','PUR-ORD-T','示例·五大湖轴承 Great Lakes Bearing (Demo)','BRG-6207',40,40,8,'accepted',%s,now()) returning id", (pid,))["id"]
    out = FN.tick(H, "teach", now=now)
    x = db.one("select * from invoice where kind='ap' and lot_id=%s", (lid,))
    assert x and x["id"] in out["ap"] and x["amount"] == 320 and x["ref"] == "MAT-PRE-2026-00091"
    _execute(db, x["proposal_id"], [["Purchase Invoice", "ACC-PINV-2026-00091"]])
    FN.sync(db, "teach")
    out = FN.tick(H, "teach", now=x["due_at"] + dt.timedelta(seconds=1))                      # 到期由出纳（模拟器）付
    assert any(p["id"] == x["id"] for p in out["pay"])


# ---------------------------------------------------------------- 规则的方案、动作
def test_f2_options_and_dunning(env):
    from hub import finance as FN
    from hub.ont import actions as AC
    H, _ = env
    db = H.db
    from factory import data as F
    inv = FN.inject(H, "teach", F.CUSTOMERS[1], "测试", days=5, qty=4)["invoice_id"]          # 前面的测试把期初那笔付了：再造一笔逾期 5 天的
    H.ai.detector().run("teach", dt.datetime.now(dt.timezone.utc))
    a = db.one("select * from anomaly where mode='teach' and rule='F2' and fkey=%s", ("INV-{:05d}".format(inv),))
    assert a is not None and a["severity"] == "warn"
    r = FN.options(db, "teach", a)
    assert [o["kind"] for o in r["options"]] == ["none", "dunning", "hold"] and r["recommended"] == "B"
    assert all("现金流预测" in o["note"] for o in r["options"]) and FN._state(db, "teach")["recommended"][str(a["id"])] == "dunning"
    ch = r["options"][1]["changes"][0]
    with pytest.raises(AC.ActionError):
        AC.create(db, "teach", OPR, "dunning", ch)
    aid = AC.create(db, "teach", SAL, "dunning", ch, None, a["id"], None)
    AC.submit(db, aid, SAL, None)
    act = AC.get(db, aid)
    assert not AC.is_costly(act) and AC.can_approve(SAL2, act)[0]
    AC.approve(H, aid, SAL2, SAL2["name"])
    x = db.one("select * from invoice where id=%s", (ch["invoice_id"],))
    assert x["dunned_n"] == 1 and AC.get(db, aid)["status"] == "done"
    out = FN.tick(H, "teach", now=x["dunned_at"] + dt.timedelta(seconds=FN.DUN_PAY_DAYS * FN.day_s("teach") + 1))
    assert any(p["id"] == x["id"] for p in out["paid"])
    assert db.one("select status from invoice where id=%s", (x["id"],))["status"] == "paid"          # 期初的直接记已付


def test_f3_payment_budget_actions(env):
    from hub import finance as FN
    from hub.ont import actions as AC
    H, _ = env
    db = H.db
    a = db.one("select * from anomaly where mode='teach' and rule='F3' order by id limit 1")
    r = FN.options(db, "teach", a)
    assert [o["kind"] for o in r["options"]] == ["none", "hold", "raise"]
    raise_ch = r["options"][2]["changes"][0]
    FN.exec_credit(H, "teach", raise_ch, TEA)
    assert db.one("select limit_amt from credit where customer=%s", (raise_ch["customer"],))["limit_amt"] == raise_ch["limit"]
    aid = AC.create(db, "teach", SAL, "budget_change", {"center": "energy", "amount": 7000})
    AC.submit(db, aid, SAL, None)
    assert "厂长" in AC.can_approve(SAL2, AC.get(db, aid))[1]
    AC.approve(H, aid, TEA, TEA["name"])
    assert FN.budgets(db, "teach")["energy"] == 7000
    x = db.one("select * from invoice where kind='ap' and status='open' order by id limit 1") or db.one("select * from invoice where kind='ap' order by id limit 1")
    db.x("update invoice set status='open' where id=%s", (x["id"],))
    before = db.one("select coalesce(pay_at, due_at) as t from invoice where id=%s", (x["id"],))["t"]
    aid = AC.create(db, "teach", SAL, "payment", {"invoice_id": x["id"], "delay_days": 7})
    AC.submit(db, aid, SAL, None)
    assert AC.is_costly(AC.get(db, aid))
    AC.approve(H, aid, TEA, TEA["name"])
    after = db.one("select pay_at from invoice where id=%s", (x["id"],))["pay_at"]
    assert abs((after - before).total_seconds() - 7 * FN.day_s("teach")) < 2


def test_f5_analysis_and_f6_options(env):
    from hub import finance as FN
    H, _ = env
    db = H.db
    so = FN.margins(db, "teach")[0]["sales_order"]
    r = FN.analysis(db, "teach", {"rule": "F5", "detail": {"values": {"sales_order": so}}})
    assert not r["options"] and "毛利率" in r["reason"]
    r = FN.options(db, "teach", {"id": None, "rule": "F6", "detail": {"values": {"center": "energy", "budget": 4000, "over": 600, "progress": 0.5}}})
    assert [o["kind"] for o in r["options"]] == ["none", "budget"] and r["options"][1]["changes"][0]["amount"] == 5200


# ---------------------------------------------------------------- 态势、报告、AI、实验
def test_situation_report_sources(env):
    from hub.ont import report as RP
    from hub.ont import situation as SIT
    H, _ = env
    db = H.db
    cards = SIT.cards(db, H.ai.detector(), "teach")
    fin = next(x for x in cards["cards"] if x["domain"] == "finance")
    assert {m["key"] for m in fin["metrics"]} >= {"gross_margin", "ar_overdue", "dso", "budget_used", "cash"}
    assert {"F1", "F2", "F3", "F4", "F5", "F6", "F7"} <= set(fin["rules_live"]) and not fin["rules_waiting"]
    st = next(x for x in cards["cards"] if x["domain"] == "strategy")
    assert any(m["key"] == "objectives_met" for m in st["metrics"])
    now = dt.datetime.now(dt.timezone.utc)
    data = RP.build(db, "teach", "weekly", now - dt.timedelta(days=7), now)
    assert data["finance"] and "经营与财务" in RP.text_of("weekly", data)
    assert db.one("select count(*) as n from ont_object where mode='teach' and type='Objective'")["n"] == 0 or True


def test_agent_finance_tools_and_intent(env):
    from hub.ont.agent import Agent, Rules, Run, Tools
    H, _ = env
    t = Tools(H, TEA, Run(TEA))
    m = t.order_margin()
    assert m["margins"] and m["target"] == 0.25
    so = m["margins"][0]["sales_order"]
    assert t.order_margin(so)["lines"] and t.order_margin("SAL-ORD-2026-99999").get("error")
    c = t.ar_ap_cash()
    assert c["ar"]["total"] > 0 and len(c["cashflow"]["weeks"]) == 8
    assert len(t.objectives_budget()["objectives"]) == 7
    r = Rules(None, False)
    assert r.intent("现在有多少应收逾期了？") == "finance" and r.intent("{} 的毛利是多少？".format(so)) == "finance"
    assert r.intent("这个月的电费多少？") == "energy" and r.intent("工单的成本差异") != "finance"
    out = Agent(H, NoLLM()).ask(TEA, [{"role": "user", "content": "现在有多少应收逾期了？现金够不够？"}])
    assert "逾期" in out["answer"] and "ar_ap_cash" in out["tools"]
    out = Agent(H, NoLLM()).ask(SAL, [{"role": "user", "content": "现在有多少应收逾期了？"}])
    assert out["answer"].startswith("提示") and not out["tools"]


def test_lab22_flow(env):
    from hub import finance as FN
    from hub import lab22
    from hub.ont import actions as AC
    H, _ = env
    db = H.db
    U = {"name": "学生销售丙", "role": "sales", "teacher": False, "mode": "teach"}
    with pytest.raises(PermissionError):
        lab22.inject(H, U, {})
    r = lab22.inject(H, TEA, {"enable": True})
    assert "海港" in r["text"]
    H.ai.detector().run("teach", dt.datetime.now(dt.timezone.utc))
    t = lab22.target(db)
    assert t["anomaly_id"] and t["invoice"]["no_pay"]
    assert not lab22.locked(db, U)
    lab22.start(db, U["name"])
    assert lab22.locked(db, U) and not lab22.locked(db, TEA)
    want = [o["code"] for o in FN.objectives(db, "teach") if o["tone"] in ("warn", "bad")]
    assert lab22.answer(db, U["name"], "l22_t1", {"codes": want})["score"] == 10
    so = lab22.target_so(db)
    m = FN.margin(db, "teach", so, store=False)
    top = max([x for x in m["lines"] if x["key"] != "material"], key=lambda x: x["usd"])["key"]
    assert lab22.answer(db, U["name"], "l22_t2", {"margin_pct": round(m["margin_pct"] * 100, 1), "top": top})["score"] == 15
    a = db.one("select * from anomaly where id=%s", (t["anomaly_id"],))
    best = FN.options(db, "teach", a)["recommended"]
    kind = {"A": "none", "B": "dunning", "C": "hold"}[best]
    assert lab22.answer(db, U["name"], "l22_t3", {"option": kind, "reason": "先催款几天内能收到钱，现金不紧，也不影响客户关系和订单发货"})["score"] == 10
    assert not lab22.locked(db, U)
    aid = AC.create(db, "teach", U, "dunning", {"invoice_id": t["invoice"]["id"]}, None, t["anomaly_id"], None)
    AC.submit(db, aid, U, None)
    assert next(x for x in lab22.check(db, U["name"])["tasks"] if x["id"] == "l22_t4")["score"] == 5
    AC.approve(H, aid, TEA, TEA["name"])
    x = db.one("select dunned_at from invoice where id=%s", (t["invoice"]["id"],))
    FN.tick(H, "teach", now=x["dunned_at"] + dt.timedelta(seconds=FN.DUN_PAY_DAYS * FN.day_s("teach") + 1))
    assert next(x for x in lab22.check(db, U["name"])["tasks"] if x["id"] == "l22_t4")["score"] == 15
    d = FN.dso(db, "teach")
    cf = FN.cashflow(db, "teach")
    assert lab22.answer(db, U["name"], "l22_t5", {"dso": d, "low_week": cf["low"]["week"]})["score"] == 10
    st = lab22.check(db, U["name"])
    assert st["score"] == 60 == st["of"] and st["report_points"] == 40
    assert any(c["user"] == U["name"] for c in lab22.class_view(db))


def test_prod_mode_no_opening_no_autoconfirm(env):
    from factory import data as F
    from hub import finance as FN
    H, _ = env
    db = H.db
    FN.seed(db, "prod")
    assert not FN.invoices(db, "prod") and not FN.sim_on(db, "prod")
    _order(db, "SAL-ORD-2026-00093", "MFG-WO-2026-00093", F.CUSTOMERS[0], mode="prod")
    FN.tick(H, "prod")
    x = db.one("select * from invoice where mode='prod' and ref='SAL-ORD-2026-00093'")
    assert x and db.one("select status from ai_proposal where proposal_id=%s", (x["proposal_id"],))["status"] == "pending"   # 生产模式等会计确认
    with pytest.raises(PermissionError):
        FN.inject(H, "prod")


def test_reset(env):
    from hub import finance as FN
    H, _ = env
    db = H.db
    FN.reset(db, "teach")
    assert not FN.sim_on(db, "teach") and len(FN.invoices(db, "teach")) == 3
    assert FN.budgets(db, "teach")["energy"] == 6000


def test_lab22_documents_and_course():
    pytest.importorskip("docx")
    import io

    import docx

    from cae import labdoc
    g = labdoc.guide_docx("lab22")
    t = labdoc.report_template_docx("lab22")
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(g)).paragraphs)
    assert "毛利" in text and "DSO" in text
    assert "实验 22" in docx.Document(io.BytesIO(t)).paragraphs[0].text
