# -*- coding: utf-8 -*-
"""第 15 轮：数字孪生——物理模型与手算、现场模拟、建档与接收（标准题）。"""
import datetime as dt
import math
import os

import pytest

import wqbus
from twin import engine as E
from twin import field as F
from twin import models as M

DSN = os.environ.get("WQ_TEST_DB", "postgresql://postgres@localhost:5433/wq_test").rsplit("/", 1)[0] + "/wq_twin_test"


# ---------------------------------------------------------------- 模型：手算核对
def test_bearing_l10_constant_load_matches_hand_calc():
    T = 350.0
    # 手算：中间轴转矩、第二级小齿轮圆周力
    t_mid = T / 3.5 / (0.97 * 0.99)
    ft34 = 2 * t_mid * 1000 / 60
    assert abs(M.torques(T)["mid"] - t_mid) < 1e-9
    P = M.bearing_loads(T)
    n_mid = 1450 / 3
    for bid in ("mid-B", "in-A", "out-A"):
        C = M.BRG[bid][3]
        n = M.speeds()[M.BRG[bid][4]]
        L10 = (C / P[bid]) ** 3 * 1e6 / (60 * n)
        D = M.bearing_damage([[T, 1000.0]])[bid]
        assert abs(D - 1000 / L10) / (1000 / L10) < 0.01
    # 输出轴：齿轮合力 = Ft/cos20°，按设计台尺寸分到两个轴承（46、82、124.5 mm）
    F34 = ft34 / math.cos(math.radians(20))
    assert abs(P["out-A"] - F34 * (124.5 - 82) / (124.5 - 46)) < 1e-6
    assert abs(n_mid - M.speeds()["mid"]) < 1e-9


def test_bearing_variable_load_equals_equivalent_load():
    hist = [[250.0, 600.0], [350.0, 300.0], [450.0, 100.0]]
    D = M.bearing_damage(hist)["mid-B"]
    Pm = M.equivalent_load(hist, "mid-B")
    C, n = M.BRG["mid-B"][3], M.speeds()["mid"]
    assert abs(D - 1000 / M.l10_hours(C, Pm, n)) / D < 1e-9
    # 降载 20%：L10 变成 1/0.8³ ≈ 1.95 倍
    D2 = M.bearing_damage([[t * 0.8, h] for t, h in hist])["mid-B"]
    assert abs(D / D2 - 1 / 0.8 ** 3) < 1e-6


def test_shaft_fatigue_matches_round11_fatigue_page():
    """同一条跑合转矩记录：孪生的 Miner 损伤与“仿真与分析”页的疲劳计算一致（< 1%）"""
    from cae import fatigue as FT
    from cae import materials as MAT
    from sim.engine import torque_record
    rec = torque_record("WQR105-T-01", "WO-T", "WQR-105")
    series = rec["samples_nm"]
    mat = MAT.get("45-QT")
    _, s = FT.compute([M.STRESS_REF_MPA], M.STRESS_REF_NM, series, rec["duration_s"], mat, surface="ground", size_factor=0.85, kf=1.0)
    cyc = [[a, m, 1] for a, m in FT.rainflow(series)]
    D = M.shaft_damage(cyc)
    assert s["damage_per_block"] > 0
    assert abs(D - s["damage_per_block"]) / s["damage_per_block"] < 0.01
    # 齿轮位改 Ø44：扭转应力 ∝ 1/d³
    assert abs(M.stress_per_nm({"segments": [[30, 40], [35, 12], [44, 60], [35, 25], [30, 30]]}) / M.stress_per_nm() - (40 / 44) ** 3) < 1e-12


def test_oil_temperature_same_as_round14_heat_balance():
    assert abs(M.oil_temp_model(350, 20) - 89.6) < 0.05           # 第 14 轮：无散热筋箱体，公式 89.6 ℃
    assert abs(M.ETA - 0.97 ** 2 * 0.99 ** 3) < 1e-12
    assert M.oil_equiv_hours(10, 60) == 10 and M.oil_equiv_hours(10, 90) == 40
    assert [M.vib_zone(v) for v in (0.5, 1.0, 2.0, 5.0)] == ["A", "B", "C", "D"]


# ---------------------------------------------------------------- 现场模拟
def _run(site, days, faults=None, serial="WQR105-T-02", d0=dt.date(2026, 3, 1)):
    s, rows = E.new_state(), []
    for i in range(days):
        s, row = E.update(s, F.day_summary(serial, site, d0 + dt.timedelta(days=i), faults))
        rows.append(row)
    return s, rows


def test_field_normal_operation_plausible():
    for site in F.SITES:
        s, rows = _run(site, 60)
        run = [r for r in rows if r["run_h"] > 0]
        assert all(abs(r["resid_c"]) < 4 for r in run)              # 正常时实测油温和模型差几度以内
        assert abs(s["resid"]) < 3
        assert all(30 < r["oil_t_c"] < 95 for r in run)
        assert all(r["vib_mm_s"] < 1.8 for r in run)                 # 正常时在 A、B 区
        assert all(r["torque_max_nm"] <= 2.3 * 350 + 1 for r in run)
    # 同一台、同一天每次一样
    assert F.day_summary("X", "mine", dt.date(2026, 5, 1)) == F.day_summary("X", "mine", dt.date(2026, 5, 1))
    # 一周摘要 ≈ 7 天之和
    wk = F.day_summary("X", "port", dt.date(2026, 5, 7), None, 7)
    dy = [F.day_summary("X", "port", dt.date(2026, 5, 1) + dt.timedelta(days=i)) for i in range(7)]
    assert abs(wk["run_h"] - sum(d["run_h"] for d in dy)) < 0.1 and wk["starts"] == sum(d["starts"] for d in dy)


def test_field_faults_show_in_data():
    d0 = dt.date(2026, 3, 1)
    _, rows = _run("port", 10, [{"kind": "cooling", "start": "2026-03-05"}])
    assert rows[3]["resid_c"] < 4 and rows[7]["resid_c"] > 8      # 散热变差：两天内油温高出模型很多
    _, rows = _run("mine", 40, [{"kind": "bearing", "start": "2026-03-05"}], d0=d0)
    v = {r["day"]: r["vib_mm_s"] for r in rows if r.get("vib_mm_s")}
    assert v["2026-03-04"] < 1.8 and max(v.values()) > 4.5          # 轴承：一个月内从 B 区到 D 区
    s0, _ = _run("mine", 30)
    s1, _ = _run("mine", 30, [{"kind": "overload", "start": "2026-03-01"}])
    assert s1["brg_D"]["mid-B"] / s0["brg_D"]["mid-B"] > 1.9       # 过载 30%：轴承消耗约 1.3³ ≈ 2.2 倍
    assert s1["shaft_D"] > s0["shaft_D"]


def test_summary_remaining_life():
    s, _ = _run("mine", 120)
    sm = E.summary(s)
    w = sm["worst_bearing"]
    assert w["id"] == "mid-B" and 0 < w["used"] < 0.2
    # 剩余天数 ≈ 剩余 / 最近 30 天的消耗速度
    rate = sum(x[5]["mid-B"] for x in s["window"]) / sum(x[1] for x in s["window"])
    assert abs(w["remaining_days"] - (1 - w["used"]) / rate) < 1e-6
    assert sm["oil"]["due_h"] == 500 and sm["oil"]["left_h"] < 0     # 磨合后首次换油已到期
    s2 = E.maintain(s, "oil", "2026-07-01")
    assert E.summary(s2)["oil"]["due_h"] == 5000 and s2["oil_eq_h"] == 0
    s3 = E.maintain(s, "bearing:mid-B", "2026-07-01")
    assert s3["brg_D"]["mid-B"] == 0 and s3["brg_D"]["in-A"] == s["brg_D"]["in-A"]


# ---------------------------------------------------------------- 服务：建档、接收、模拟器
@pytest.fixture
def db():
    psycopg = pytest.importorskip("psycopg")
    try:
        with psycopg.connect(DSN.rsplit("/", 1)[0] + "/postgres", autocommit=True) as c:
            c.execute("drop database if exists wq_twin_test with (force)")
            c.execute("create database wq_twin_test")
    except Exception:  # noqa: BLE001
        pytest.skip("没有可用的 PostgreSQL")
    from hub.db import DB
    d = DB(DSN)
    d.init()
    return d


def _service(db):
    from hub.historian import Historian
    from hub.twin import TwinService
    sent = []
    holder = {}

    def publish(tp, type_, source, data, corr=None, mode=None):
        msg = wqbus.make(type_, source, data, mode=mode or "teach", corr=corr)
        sent.append((tp, msg))
        holder["hist"].handle(tp, msg)
        return msg
    tw = TwinService(db, publish)
    tw.init()
    holder["hist"] = Historian(db, tw.on_message)
    return tw, sent, holder["hist"]


def test_twin_from_runin_and_field(db, monkeypatch):
    from hub import twin as TW
    from sim.engine import torque_record
    monkeypatch.setattr(TW, "FIELD_MODES", ["teach"])
    tw, sent, hist = _service(db)
    # 一根合格的 SH-301（检验站）、一台 WQR-105 跑合
    for code, name, nom, lo, hi in __import__("sim.errors", fromlist=["CHARS"]).CHARS:
        hist.handle("wq/gearbox/quality/qc-01/measurement", wqbus.make("quality.measurement", "sim/qc-01", {
            "part_serial": "SH301-A-01", "item": "SH-301", "work_order": "WO-S", "characteristic": code, "name": name,
            "nominal_mm": nom, "lower_tol_mm": lo, "upper_tol_mm": hi, "value_mm": nom, "result": "pass"}, mode="teach"))
    rec = torque_record("WQR105-A-01", "WO-A", "WQR-105")
    hist.handle("wq/gearbox/machining/test-01/torque", wqbus.make("test.torque", "sim/test-01", rec, mode="teach"))
    r = db.one("select * from twin where serial='WQR105-A-01'")
    assert r["as_built"]["shaft"]["serial"] == "SH301-A-01" and r["as_built"]["runin"]["peak_nm"] > 500
    assert r["as_built"]["design"]["segments"][2][0] == 40
    fu = db.one("select * from field_unit where serial='WQR105-A-01'")
    assert fu["start_day"] == tw.field_today("teach") + dt.timedelta(days=3)
    # 现场时钟往前拨 10 天：模拟器补发 7 条（3 天运输安装）
    db.x("update field_clock set anchor_day = anchor_day + 10 where mode='teach'")
    assert tw.tick("teach") == 7
    s = db.one("select state from twin where serial='WQR105-A-01'")["state"]
    assert s["days"] == 7 and s["run_h"] > 0 and s["brg_D"]["mid-B"] > 0
    assert len(db.q("select * from twin_log where serial='WQR105-A-01'")) == 7
    assert tw.tick("teach") == 0                                  # 不重复补
    # 旧日子的数据（乱序、重复）不再累加
    old = F.day_summary("WQR105-A-01", fu["site"], fu["start_day"])
    hist.handle("wq/gearbox/field/WQR105-A-01/telemetry", wqbus.make("twin.telemetry", "x", old, mode="teach"))
    assert db.one("select state from twin where serial='WQR105-A-01'")["state"]["days"] == 7


def test_real_device_without_record_and_seed(db, monkeypatch):
    from hub import twin as TW
    monkeypatch.setattr(TW, "FIELD_MODES", ["teach"])
    tw, sent, hist = _service(db)
    # 真设备：按格式发一条（没建过档）
    data = {"serial": "REAL-0001", "day": "2026-09-30", "period_h": 24, "run_h": 10, "starts": 3, "n_in_rpm": 1450,
            "torque_mean_nm": 250, "torque_max_nm": 500, "t_amb_c": 25, "oil_t_c": 62, "vib_mm_s": 0.8,
            "rainflow": [[250, 250, 3]], "load_hist": [[250, 10]]}
    hist.handle("wq/gearbox/field/REAL-0001/telemetry", wqbus.make("twin.telemetry", "device/REAL-0001", data, mode="prod"))
    r = db.one("select * from twin where serial='REAL-0001' and mode='prod'")
    assert r["source"] == "field" and r["state"]["days"] == 1
    assert abs(r["state"]["brg_D"]["mid-B"] - M.bearing_damage([[250, 10]])["mid-B"]) < 1e-12
    # 情景设备：12 台，历史补齐
    n = tw.seed("teach")
    assert n > 300 and tw.seed("teach") == 0
    fl = TW.fleet(db, "teach")
    assert len(fl) == 12
    by = {u["serial"]: u for u in fl}
    old = max(fl, key=lambda u: u["run_h"])
    assert old["summary"]["worst_bearing"]["used"] > 0.3            # 用了两年多的矿山设备
    brg = [u for u in fl if u["summary"]["vib"] and u["summary"]["vib"] > 1.8]
    hot = [u for u in fl if (u["summary"]["resid_c"] or 0) > 6]
    assert len(brg) == 1 and brg[0]["site"] == "mine" and len(hot) == 1 and hot[0]["site"] == "port"
    assert sum(1 for u in fl if u["summary"]["oil"]["left_h"] < 0) == 1          # 只有一台换油过期
    d = TW.detail(db, "teach", old["serial"])
    assert any(e["kind"] == "maint" for e in d["events"])
    assert d["log"] and d["as_built"]["shaft"]["meas"] and "runin_samples" not in d["as_built"]
    assert by
    tw.reset("teach", reseed=False)
    assert not db.q("select * from twin where mode='teach'")


# ---------------------------------------------------------------- 第 2 步：异常判断
def _diag_run(site, days, faults=None, serial="WQR105-T-09", d0=dt.date(2026, 3, 1)):
    from twin import diagnose as DG
    s, logs, hist = E.new_state(), [], []
    for i in range(days):
        s, row = E.update(s, F.day_summary(serial, site, d0 + dt.timedelta(days=i), faults))
        logs.append(row)
        hist.append(DG.diagnose(s, logs[-30:]))
    return s, hist


def test_no_false_alarm_in_normal_month():
    for site in F.SITES:
        for serial in ("WQR105-T-11", "WQR105-T-12", "WQR105-T-13"):
            _, hist = _diag_run(site, 31, serial=serial)
            # 保养到期（磨合后首次换油）是该报的，不算误报
            bad = [(i, f["kind"], f["title"]) for i, fs in enumerate(hist) for f in fs if f["level"] != "info" and f["kind"] != "oil"]
            assert not bad, (site, serial, bad[:3])


def test_cooling_fault_detected_within_a_day():
    _, hist = _diag_run("port", 12, [{"kind": "cooling", "start": "2026-03-06"}])
    assert not any(f["kind"] == "cooling" for fs in hist[:5] for f in fs)
    hit = next(i for i, fs in enumerate(hist) if any(f["kind"] == "cooling" for f in fs))
    assert hit - 5 <= 1                                              # 故障开始后 1 个现场日内
    f = next(f for f in hist[hit] if f["kind"] == "cooling")
    assert "散热变差" in f["title"] and f["action"]["kind"] == "cooling"
    assert not any(f["kind"] == "bearing" for fs in hist for f in fs)


def test_bearing_fault_detected_and_located():
    _, hist = _diag_run("mine", 40, [{"kind": "bearing", "start": "2026-03-06"}])
    first = next(i for i, fs in enumerate(hist) if any(f["kind"] == "bearing" for f in fs))
    assert 5 < first < 25
    f = next(f for f in hist[first] if f["kind"] == "bearing")
    assert f["level"] == "warn" and "中间轴轴承 B" in f["cause"] and f["action"]["kind"] == "bearing:mid-B"
    assert any(f["kind"] == "bearing" and f["level"] == "critical" for f in hist[-1])
    assert not any(f["kind"] == "cooling" and "散热变差" in f["title"] for fs in hist[:first] for f in fs)


def test_overload_detected():
    _, hist = _diag_run("mine", 10, [{"kind": "overload", "start": "2026-03-01"}])
    assert any(f["kind"] == "overload" and f["level"] == "warn" for f in hist[-1])
    _, hist = _diag_run("mine", 10)
    assert not any(f["kind"] == "overload" for fs in hist for f in fs)


def test_fleet_health_and_repair(db, monkeypatch):
    from hub import twin as TW
    monkeypatch.setattr(TW, "FIELD_MODES", ["teach"])
    tw, sent, hist = _service(db)
    tw.seed("teach")
    fl = {u["serial"]: u for u in TW.fleet(db, "teach")}
    bad = {s: [a["kind"] for a in u["alerts"]] for s, u in fl.items() if u["health"] != "green"}
    kinds = sorted(k for v in bad.values() for k in v)
    assert "bearing" in kinds and "cooling" in kinds and "oil" in kinds, bad
    assert not [m for _, m in sent if m["type"] == "ai.alert"]       # 情景历史不发提醒
    brg = next(s for s, v in bad.items() if "bearing" in v)
    # 修好：换中间轴轴承 B → 寿命复位、故障从现场消失；之后的数据不再报警
    tw.repair("teach", brg, "bearing:mid-B")
    st = db.one("select state from twin where serial=%s", (brg,))["state"]
    assert st["brg_D"]["mid-B"] == 0 and not db.one("select faults from field_unit where serial=%s", (brg,))["faults"]
    db.x("update field_clock set anchor_day = anchor_day + 3 where mode='teach'")
    tw.tick("teach")
    assert "bearing" not in [a["kind"] for a in next(u for u in TW.fleet(db, "teach") if u["serial"] == brg)["alerts"]]
    # 老师制造故障 → 报警并发 AI 提醒
    ok = next(s for s, u in fl.items() if u["health"] == "green" and u["site"] == "port")
    tw.set_fault("teach", ok, "cooling")
    db.x("update field_clock set anchor_day = anchor_day + 3 where mode='teach'")
    tw.tick("teach")
    u = next(u for u in TW.fleet(db, "teach") if u["serial"] == ok)
    assert u["health"] != "green" and any(a["kind"] == "cooling" for a in u["alerts"])
    assert any(m["type"] == "ai.alert" and ok in m["data"]["title"] for _, m in sent)
    with pytest.raises(ValueError):
        tw.set_fault("teach", ok, "nope")


def test_whatif_load_fan_material():
    from twin import whatif as W
    s, logs = E.new_state(), []
    for i in range(40):
        s, row = E.update(s, F.day_summary("WQR105-T-21", "mine", dt.date(2026, 3, 1) + dt.timedelta(days=i)))
        logs.append(row)
    base = W.compare(s, logs)
    assert base["base"]["bearing_days"] == base["alt"]["bearing_days"]
    r = W.compare(s, logs, load=0.8)
    # 剩余寿命 = (1 − 已用) / 速度：速度变成 0.8³ 倍
    assert abs(r["alt"]["bearing_days"] / r["base"]["bearing_days"] - 1 / 0.8 ** 3) < 1e-6
    f = W.compare(s, logs, fan=True)
    assert f["alt"]["oil_t_c"] < f["base"]["oil_t_c"] - 10 and f["alt"]["oil_interval_days"] >= f["base"]["oil_interval_days"]
    assert f["alt"]["bearing_days"] == f["base"]["bearing_days"]
    m = W.compare(s, logs, mat="40Cr-QT")
    assert m["alt"]["shaft_days"] > m["base"]["shaft_days"]


def test_api_order_done_explain_feedback(db, monkeypatch):
    """网页接口：机队、详情、假如、AI 说明（规则）、开维修单（提议 → 确认）、维修完成（寿命复位、ERPNext 工单改为已解决）、反馈设计"""
    import importlib
    from fastapi.testclient import TestClient
    monkeypatch.setenv("WQ_HUB_NO_START", "1")
    monkeypatch.setenv("WQ_SECRET", "t")
    import hub.app as app_mod
    app_mod = importlib.reload(app_mod)
    from hub import twin as TW
    from hub.ai import Assistant
    from hub.historian import Historian
    monkeypatch.setattr(TW, "FIELD_MODES", ["teach"])
    sent = []

    def publish(tp, type_, source, data, corr=None, mode=None):
        msg = wqbus.make(type_, source, data, mode=mode or "teach", corr=corr)
        sent.append((tp, msg))
        app_mod.H.hist.handle(tp, msg)
        return msg
    app_mod.H.db = db
    app_mod.H.twin = TW.TwinService(db, publish)
    app_mod.H.twin.init()
    app_mod.H.hist = Historian(db, app_mod.H.twin.on_message)
    app_mod.H.ai = Assistant(db, publish)

    class _Teach:
        def start(self, name):
            pass
    app_mod.H.teach = _Teach()
    app_mod.H.twin.seed("teach")
    c = TestClient(app_mod.app)
    h = {"x-wq-token": c.post("/api/login", json={"name": "老师", "role": "manager", "mode": "teach"}).json()["token"]}
    fl = c.get("/api/twin/fleet", headers=h).json()
    assert len(fl["units"]) == 12 and fl["clock"]["simulated"]
    brg = next(u for u in fl["units"] if any(a["kind"] == "bearing" for a in u["alerts"]))
    d = c.get("/api/twin/units/" + brg["serial"], headers=h).json()
    assert d["model"]["l10"]["mid-B"] > 10000
    ev = next(e for e in d["events"] if e["kind"] == "bearing" and e["status"] == "open")
    ex = c.post(f"/api/twin/units/{brg['serial']}/explain", headers=h).json()
    assert ex["engine"] == "rules" and "中间轴轴承 B" in ex["text"]
    wi = c.post(f"/api/twin/units/{brg['serial']}/whatif", json={"load": 0.8}, headers=h).json()
    assert wi["alt"]["bearing_days"] > wi["base"]["bearing_days"]
    assert c.post(f"/api/twin/units/{brg['serial']}/whatif", json={"load": 3}, headers=h).status_code == 400
    pv = c.post(f"/api/twin/units/{brg['serial']}/order", json={"action": ev["detail"]["action"]["kind"], "event_id": ev["id"], "preview_only": True}, headers=h).json()
    assert "中间轴轴承 B" in pv["subject"] and pv["priority"] in ("一般", "紧急")
    r = c.post(f"/api/twin/units/{brg['serial']}/order", json={"action": ev["detail"]["action"]["kind"], "event_id": ev["id"]}, headers=h).json()
    p = db.one("select * from ai_proposal where proposal_id=%s", (r["proposal_id"],))
    assert p["status"] == "confirmed" and p["action"] == "twin_service_order"
    assert any(m["type"] == "ai.proposal" and m["data"]["status"] == "confirmed" for _, m in sent)
    # 桥接写进 ERPNext 后发回 erp.doc（这里模拟）
    app_mod.H.ai.on_erp_result(wqbus.make("erp.doc", "bridge", {"doctype": "Issue", "name": "ISS-2026-00001", "action": "created",
                                                                 "proposal_done": True, "created": [["Issue", "ISS-2026-00001"]]},
                                          mode="teach", corr=r["proposal_id"]))
    d = c.get("/api/twin/units/" + brg["serial"], headers=h).json()
    o = next(e for e in d["events"] if e["kind"] == "order")
    assert o["status"] == "ordered" and next(e for e in d["events"] if e["id"] == ev["id"])["status"] == "ordered"
    assert c.post(f"/api/twin/units/{brg['serial']}/order", json={"action": "nope"}, headers=h).status_code == 400
    done = c.post(f"/api/twin/units/{brg['serial']}/done", json={"order_id": o["id"], "note": "外圈剥落，已换"}, headers=h).json()
    assert done["issue"] == "ISS-2026-00001"
    p2 = db.one("select * from ai_proposal where action='twin_service_done'")
    assert p2["status"] == "confirmed" and p2["preview"]["issue"] == "ISS-2026-00001" if isinstance(p2["preview"], dict) else "ISS-2026-00001" in p2["preview"]
    d = c.get("/api/twin/units/" + brg["serial"], headers=h).json()
    assert d["state"]["brg_D"]["mid-B"] == 0
    assert next(e for e in d["events"] if e["id"] == ev["id"])["status"] == "resolved"
    assert c.post(f"/api/twin/units/{brg['serial']}/done", json={"order_id": o["id"]}, headers=h).status_code == 409
    fb = c.get("/api/twin/feedback", headers=h).json()
    by = {s["site"]: s for s in fb["sites"]}
    assert by["mine"]["eq_vs_rated"] > by["port"]["eq_vs_rated"] > by["food"]["eq_vs_rated"]
    assert by["mine"]["worst_bearing"] == "mid-B" and any("6306" in n for n in fb["notes"])


# ---------------------------------------------------------------- 第 5 步：机械臂关节孪生
def test_harmonic_reducer_life_formula():
    import numpy as np
    from twin import robot as R
    t = np.linspace(0, 10, 2001)
    w = np.full_like(t, 0.5)                                     # 关节 0.5 rad/s
    tau = np.full_like(t, 60.0)
    r = R.joint_life(t, tau, w, 137.0, 100)
    n_in = 0.5 * 60 / (2 * math.pi) * 100
    assert abs(r["n_av_in"] - n_in) / n_in < 1e-6 and abs(r["T_av"] - 60) < 1e-6
    assert abs(r["L10_h"] - 7000 * (137 / 60) ** 3 * (2000 / n_in)) / r["L10_h"] < 1e-6
    # 两段：力矩 40 N·m 转 1 rad/s 一半时间，80 N·m 转 0.5 rad/s 另一半 → 按转速加权的立方平均
    w2 = np.where(t < 5, 1.0, 0.5)
    tau2 = np.where(t < 5, 40.0, 80.0)
    r2 = R.joint_life(t, tau2, w2, 137.0, 100)
    expect = ((1.0 * 40 ** 3 + 0.5 * 80 ** 3) / 1.5) ** (1 / 3)
    assert abs(r2["T_av"] - expect) / expect < 0.01
    ser = {"t": t, "drive.j1": tau2, "qd.j1": w2, "qdd.j1": np.gradient(w2, t)}
    out = R.reducer_life(ser, ["j1"])
    j = out["joints"][0]
    assert j["gear"].startswith("谐波") and j["L10_h"] > 0 and out["cycle_s"] == 10


def test_robot_twin_accumulates(db, monkeypatch):
    from hub import twin as TW
    monkeypatch.setattr(TW, "FIELD_MODES", ["teach"])
    tw, sent, hist = _service(db)
    life = {"job": "J1", "cycle_s": 12, "joints": [{"joint": "shoulder", "L10_h": 20000, "gear": "谐波 32 号"},
                                                   {"joint": "wrist", "L10_h": 80000, "gear": "谐波 17 号"}]}
    rid = tw.add_robot("teach", "UR5e 搬运", life, 16)
    db.x("update field_clock set anchor_day = anchor_day + 30 where mode='teach'")
    tw.tick_robots("teach")
    r = TW.robots(db, "teach")[0]
    assert r["id"] == rid and r["first"] == "shoulder"
    sh = next(j for j in r["joints"] if j["joint"] == "shoulder")
    assert 30 * 16 * 0.85 / 20000 * 0.7 < sh["used"] < 30 * 16 / 20000 * 1.35
    assert abs(sh["remaining_days"] - (1 - sh["used"]) / (sh["used"] / 30)) < 1e-6


def test_real_device_http_entry(db, monkeypatch):
    """真设备经 HTTPS 发运行摘要（钥匙校验、格式校验），和模拟器走同一条路进孪生"""
    import importlib
    from fastapi.testclient import TestClient
    monkeypatch.setenv("WQ_HUB_NO_START", "1")
    monkeypatch.setenv("WQ_TWIN_DEVICE_KEY", "dev-key")
    import hub.app as app_mod
    app_mod = importlib.reload(app_mod)
    from hub import twin as TW
    from hub.historian import Historian
    app_mod.H.db = db
    app_mod.H.twin = TW.TwinService(db, lambda *a, **k: None)
    app_mod.H.twin.init()
    app_mod.H.hist = Historian(db, app_mod.H.twin.on_message)
    c = TestClient(app_mod.app)
    d = F.day_summary("REAL-0002", "port", dt.date(2026, 9, 1))
    assert c.post("/api/twin/telemetry", json=d).status_code == 401
    assert c.post("/api/twin/telemetry", json=d, headers={"Authorization": "Bearer nope"}).status_code == 401
    bad = dict(d, day="2026-9-1")
    assert c.post("/api/twin/telemetry", json=bad, headers={"Authorization": "Bearer dev-key"}).status_code == 422
    r = c.post("/api/twin/telemetry", json=dict(d, mode="teach"), headers={"Authorization": "Bearer dev-key"}).json()
    assert r["topic"] == "wq/gearbox/field/REAL-0002/telemetry"
    t = db.one("select * from twin where serial='REAL-0002' and mode='teach'")
    assert t["source"] == "field" and t["state"]["days"] == 1
    assert db.one("select count(*) as n from bus_message where type='twin.telemetry'")["n"] == 1


def test_lab13_documents():
    import io
    import docx
    from cae import labdoc
    g = docx.Document(io.BytesIO(labdoc.guide_docx("lab13")))
    text = "\n".join(p.text for p in g.paragraphs)
    assert "实验 13" in text and "数字孪生" in text and "ISO 281" in text
    t = docx.Document(io.BytesIO(labdoc.report_template_docx("lab13")))
    cells = " ".join(c.text for tb in t.tables for r in tb.rows for c in r.cells)
    assert "2990" in cells and "ERPNext 服务工单号" in cells
