# -*- coding: utf-8 -*-
"""第 18 轮 E6：能源与环境——仿真车间的分档功率与 15 分钟读数、自动待机与预热、错峰、教学问题（忘关、漏气、劣化、峰时开炉）、
读数入库、峰平谷电费、需量、碳核算、车间环境、单耗与空转、规则 N1–N6、方案推演、节能设定 / 错峰动作、态势与日报、AI 代理三个新工具、实验 21。"""
import datetime as dt
import json
import os

import pytest

from tests.test_agent import NoLLM

DSN = os.environ.get("WQ_TEST_DB", "postgresql://postgres@localhost:5433/wq_test")
TEA = {"name": "王老师", "role": "manager", "teacher": True, "mode": "teach"}
PLN = {"name": "学生计划", "role": "planner", "teacher": False, "mode": "teach"}
PLN2 = {"name": "学生计划乙", "role": "planner", "teacher": False, "mode": "teach"}
OPR = {"name": "学生操作工", "role": "operator", "teacher": False, "mode": "teach"}


class FakeMes:
    def __init__(self):
        self.sent = []

    def command(self, unit, command, mode, user, work_order=None, operation=None, **extra):
        self.sent.append(dict({"unit": unit, "command": command, "mode": mode}, **extra))


def _run(eng, clock, hours, step=60):
    for _ in range(int(hours * 3600 / step)):
        clock.t += step
        eng.run_due()


def _readings(out):
    return [m["data"] for _, m in out if m["type"] == "energy.reading"]


# ---------------------------------------------------------------- 仿真车间
def test_energy_off_changes_nothing_and_on_keeps_production():
    from tests_helpers import make_engine

    def cycles(en):
        eng, clock, out = make_engine(auto=True)
        if en:
            eng.set_energy({"enabled": True})
        eng.at(0, "release", work_order="WO-T", item="SH-301", qty=8)
        _run(eng, clock, 12)
        return [(m["data"]["event"], m["data"].get("part_serial"), m["data"].get("cycle_time_s")) for _, m in out if m["type"] == "machine.event"], out
    a, out_a = cycles(False)
    b, out_b = cycles(True)
    assert a == b                                                         # 计量只记数，不改车间节拍
    assert not _readings(out_a) and len(_readings(out_b)) >= 47           # 关闭时没有读数；打开时每 15 分钟一个


def test_reading_energy_equals_state_time_power():
    from sim import energy as SE
    from tests_helpers import make_engine
    eng, clock, out = make_engine(auto=True)
    eng.set_energy({"enabled": True})
    eng.at(0, "release", work_order="WO-T", item="SH-301", qty=6)
    _run(eng, clock, 24)
    rd = _readings(out)
    assert rd and all(r["seconds"] <= 900 + 1e-6 for r in rd)
    for u, (run, idle, stby) in SE.POWER.items():
        if u == "ht-01":
            continue
        m = [r["meters"][SE.meter_of(u)] for r in rd]
        s = {}
        for x in m:
            for k, v in x["s"].items():
                s[k] = s.get(k, 0) + v
        want = (s.get("run", 0) * run + (s.get("idle", 0) + s.get("setup", 0)) * idle + (s.get("standby", 0) + s.get("down", 0)) * stby) / 3600
        assert abs(sum(x["value"] for x in m) - want) < 0.05 * max(1, want), u
    tot = sum(r["meters"]["E-TOTAL"]["value"] for r in rd)
    parts = sum(r["meters"][k]["value"] for r in rd for k in r["meters"] if k.startswith("E-") and k != "E-TOTAL")
    assert abs(tot - parts - sum(1 for _ in rd) * 0) < tot * 0.15                      # 总表 = 各设备 + 空压机 + 照明
    assert any(r["meters"]["E-GRD-01"]["parts"] for r in rd)                            # 每件的运行电量按物料记下
    assert {r["bucket"] for r in rd} <= {"peak", "flat", "valley"}


def test_standby_saves_energy_with_warmup():
    from sim import energy as SE
    from tests_helpers import make_engine
    res = {}
    for sb in (None, 10):
        eng, clock, out = make_engine(auto=True)
        eng.set_energy({"enabled": True, "standby": {u: sb for u in SE.POWER} if sb else {}})
        eng.at(0, "release", work_order="WO-T", item="SH-301", qty=6)
        _run(eng, clock, 24)
        res[sb] = SE.summarize({}, eng.energy.totals())
    assert res[10]["kwh"] < res[None]["kwh"] * 0.8 and res[10]["idle"] < res[None]["idle"] * 0.1 and res[10]["standby"] > 0
    eng, clock, out = make_engine(auto=True)
    eng.set_energy({"enabled": True, "standby": {"saw-01": 10}})
    _run(eng, clock, 1)
    assert eng._epower(eng.machines["saw-01"], clock.t)[0] == "standby"
    eng.at(clock.t, "release", work_order="WO-W", item="SH-301", qty=1)
    clock.t += 1
    eng.run_due()
    assert eng.machines["saw-01"].state == "setup" and eng.machines["saw-01"].setup_until - clock.t == pytest.approx(SE.WARM_MIN * 60, abs=2)   # 先预热


def test_off_when_nobody_on_shift_and_forgot():
    from sim import staffing as STF
    from tests_helpers import make_engine
    eng, clock, out = make_engine(auto=True)
    local = eng.factory_dt(0).astimezone(eng.calendar.tzi)
    d, sh = STF.shift_of(local)
    lv = {"{}|{}|hmc-01".format((d + dt.timedelta(days=i)).isoformat(), s): 0 for i in (-1, 0, 1, 2) for s in ("day", "night")}
    eng.set_staffing({"enabled": True, "levels": lv})
    eng.set_energy({"enabled": True})
    _run(eng, clock, 1)
    assert eng.energy.cur["hmc-01"][1] == "off"                           # 这一班没人：下班的人关了
    eng.set_energy({"inject": [{"kind": "forgot", "unit": "hmc-01", "hours": 2}]})
    assert eng.energy.cur["hmc-01"][1] == "idle"                          # 忘关：照样空转
    _run(eng, clock, 2.5)
    assert eng.energy.cur["hmc-01"][1] == "off"                           # 到期后关了


def test_leak_degrade_and_ht_peak_with_tou():
    from sim import energy as SE
    from tests_helpers import make_engine
    res = {}
    for name, setting, inj in (("base", {}, []), ("leak", {}, [{"kind": "leak"}]), ("degrade", {}, [{"kind": "degrade", "unit": "grd-01", "pct": 0.4}]),
                               ("peak", {}, [{"kind": "ht_peak"}]), ("peak_tou", {"tou": ["ht-01"]}, [{"kind": "ht_peak"}])):
        eng, clock, out = make_engine(auto=True)
        # 固定从周一 10:00 开始（工厂时间原来取真实时间：在不同时刻跑，错峰前后的电量对比会落在不同的电价时段，测试时过时不过）
        eng.epoch = dt.datetime(2026, 10, 5, 10, 0, tzinfo=eng.calendar.tzi).astimezone(dt.timezone.utc)
        eng.set_energy(dict({"enabled": True}, **setting))
        eng.set_energy({"inject": inj})
        eng.at(0, "release", work_order="WO-T", item="SH-301", qty=8)
        _run(eng, clock, 30)
        rd = _readings(out)
        grd = [p for r in rd for p in r["meters"]["E-GRD-01"]["parts"].values()]
        res[name] = dict(SE.summarize({}, eng.energy.totals()), air=sum(r["meters"]["E-AIR"]["value"] for r in rd),
                         grd_pp=sum(k for n, k in grd) / max(1, sum(n for n, k in grd)))
    assert res["leak"]["air"] > res["base"]["air"] + 100                   # 漏气：空压机一天多 6 kW × 24 h 左右
    assert res["degrade"]["grd_pp"] > res["base"]["grd_pp"] * 1.2          # 劣化：磨床每件多用电
    assert res["peak"]["peak"] > res["base"]["peak"] + 200 and res["peak"]["cost"] > res["base"]["cost"]
    assert res["peak_tou"]["peak"] < res["peak"]["peak"] - 200 and abs(res["peak_tou"]["kwh"] - res["peak"]["kwh"]) < 50   # 错峰：电量一样，挪出峰时
    assert res["peak_tou"]["cost"] < res["peak"]["cost"]


def test_tou_blocks_new_batches_in_peak():
    from tests_helpers import make_engine
    eng, clock, out = make_engine(auto=True)
    eng.set_energy({"enabled": True, "tou": ["ht-01"]})
    m = eng.machines["ht-01"]
    local = eng.factory_dt(0).astimezone(eng.calendar.tzi)
    peak = local.replace(hour=9, minute=0, second=0, microsecond=0)
    if peak <= local:
        peak += dt.timedelta(days=1)
    clock.t = eng.sim_t(peak.astimezone(dt.timezone.utc))
    assert eng._tou_closed(m, clock.t) and eng._tou_wait["ht-01"] > clock.t     # 峰时不开新炉，排好峰后再开
    clock.t += 3 * 3600
    assert not eng._tou_closed(m, clock.t)


def test_snapshot_roundtrip_keeps_energy():
    from sim import energy as SE
    from sim.engine import Engine
    from tests_helpers import FakeClock, make_engine
    eng, clock, out = make_engine(auto=True)
    eng.set_energy({"enabled": True, "standby": {"cnc-l01-b": 10}, "inject": [{"kind": "degrade", "unit": "grd-01"}]})
    eng.at(0, "release", work_order="WO-T", item="SH-301", qty=8)
    _run(eng, clock, 5)
    snap = json.loads(json.dumps(eng.snapshot()))
    assert snap["energy"]["enabled"] and snap["energy"]["standby"] == {"cnc-l01-b": 10.0}
    c2 = FakeClock()
    c2.t = clock.t
    e2 = Engine.from_snapshot(snap, lambda *a: None, c2)
    for _ in range(600):
        clock.t += 60
        eng.run_due()
        c2.t += 60
        e2.run_due()
    assert SE.summarize({}, eng.energy.totals())["kwh"] == pytest.approx(SE.summarize({}, e2.energy.totals())["kwh"], rel=1e-6)


def test_plan_energy_changes_and_capture():
    from hub.ont import plan as P
    from tests_helpers import make_engine
    assert "energy_setting" in P.CHANGE_TYPES and "tou_rule" in P.CHANGE_TYPES
    eng, clock, out = make_engine(auto=True)
    eng.at(0, "release", work_order="WO-T", item="SH-301", qty=6)
    _run(eng, clock, 1)
    snap = json.loads(json.dumps(eng.snapshot()))
    ctx = {"energy": {"price": {"peak": 0.15, "flat": 0.1, "valley": 0.05}, "factor": 0.5306}}
    a = P.simulate(snap, [], ctx)
    b = P.simulate(snap, [{"type": "energy_setting", "units": ["vmc-01", "hmc-01", "cnc-l01-b"], "standby_min": 10}], ctx)
    assert a["energy"]["hours"] == 24 and b["energy"]["kwh"] < a["energy"]["kwh"] - 50
    assert "energy" not in P.simulate(snap, [], {})                       # 不问电量的推演不算


# ---------------------------------------------------------------- 平台（沙盘：过去 24 小时的读数）
@pytest.fixture(scope="module")
def env():
    pytest.importorskip("psycopg")
    from hub.ont import evals as EV
    try:
        H, eng = EV.sandbox(DSN, name="wq_energy_test", llm=NoLLM())
    except Exception as e:  # noqa: BLE001
        pytest.skip("没有可用的 PostgreSQL：{}".format(e))
    H.mes = FakeMes()
    return H, eng


def test_meters_readings_overview_cost(env):
    from hub import energy as EN
    H, _ = env
    db = H.db
    assert len(EN.meters(db, "teach")) == 18 and db.one("select count(distinct meter) as n from energy_read where mode='teach'")["n"] == 16
    o = EN.overview(db, "teach")
    assert o["hours"] >= 6 and o["kwh"] > 0
    assert abs(sum(o["by_bucket"].values()) - o["kwh"]) < 0.5
    want = sum(o["by_bucket"][b] * o["price"][b] for b in ("peak", "flat", "valley"))
    assert o["cost"] == pytest.approx(want, abs=0.5) and o["co2_kg"] == pytest.approx(o["kwh"] * 0.5306, abs=0.5)
    hm = next(u for u in o["units"] if u["unit"] == "hmc-01")
    assert hm["by_state"].get("idle", 0) > 0                              # 卧加下班忘关
    assert EN.demand(db, "teach")["kw"] > 0 and EN.kpis(db, "teach")["kwh"] == o["kwh"]
    c = EN.carbon(db, "teach")
    assert c["t_co2e"] > 0 and c["factor"] == 0.5306 and any(p["item"] == "SH-301" for p in c["products"])
    e = EN.env_now(db, "teach")
    assert set(e) == {"mist", "noise"} and e["noise"]["limit"] == 85.0


def test_params_set_and_validate(env):
    from hub import energy as EN
    H, _ = env
    p = EN.set_params(H.db, "teach", {"contract_kw": 250, "price": {"peak": 0.2}})
    assert p["contract_kw"] == 250 and p["price"] == {"peak": 0.2, "flat": 0.1, "valley": 0.05}
    with pytest.raises(ValueError):
        EN.set_params(H.db, "teach", {"co2_factor": -1})
    EN.set_params(H.db, "teach", {"contract_kw": 220, "price": {"peak": 0.15}})


def test_anomalies_from_scene(env):
    H, _ = env
    rules = {r["rule"]: r for r in H.db.q("select rule, subject_key, title from anomaly where mode='teach' and rule like 'N%%' and status in ('open','ack','doing')")}
    assert {"N1", "N2", "N3", "N4"} <= set(rules)                         # 磨床劣化、空转、峰时开炉、漏气
    assert rules["N1"]["subject_key"] == "grd-01"


def test_idle_options_and_energy_setting_action(env):
    from hub import energy as EN
    from hub.ont import actions as AC
    from hub.ont import plan as P
    H, _ = env
    db = H.db
    a = db.one("select * from anomaly where mode='teach' and rule='N2' and subject_key='hmc-01' order by id limit 1") or \
        db.one("select * from anomaly where mode='teach' and rule='N2' order by id limit 1")
    _, st, _ = H.plan_snaps.take("teach")
    r = EN.options(db, "teach", a, st, P.context(db, "teach", dt.datetime.fromisoformat(st["factory_now"])))
    kinds = [o["kind"] for o in r["options"]]
    assert kinds == ["none", "standby10", "standby30", "all10"]
    for o in r["options"][1:]:
        assert o["saving"]["kwh"] > 0 and o["result"]["cost_total"] < 0 and "每天省" in o["note"]
    assert r["recommended"] != "A" and EN._state(db, "teach")["recommended"][str(a["id"])] == next(o["kind"] for o in r["options"] if o["recommended"])
    ch = r["options"][1]["changes"][0]
    with pytest.raises(AC.ActionError):
        AC.create(db, "teach", OPR, "energy_setting", ch)                 # 操作工不能发起
    tr = AC.trial(db, H.plan_snaps, "teach", [ch], st)
    assert tr["energy"]["kwh"] > 0
    aid = AC.create(db, "teach", PLN, "energy_setting", ch, None, a["id"], tr)
    AC.submit(db, aid, PLN, tr)
    act = AC.get(db, aid)
    assert not AC.can_approve(PLN2, act)[0] and "厂长" in AC.can_approve(PLN2, act)[1]
    n0 = len(H.mes.sent)
    AC.approve(H, aid, TEA, TEA["name"])
    assert AC.get(db, aid)["status"] == "done"
    assert EN.settings(db, "teach")["standby"][ch["units"][0]] == 10
    c = H.mes.sent[n0]
    assert c["command"] == "set_energy" and c["standby"][ch["units"][0]] == 10


def test_tou_rule_action_and_n3_options(env):
    from hub import energy as EN
    from hub.ont import actions as AC
    from hub.ont import plan as P
    H, _ = env
    db = H.db
    a = db.one("select * from anomaly where mode='teach' and rule='N3' order by id limit 1")
    _, st, _ = H.plan_snaps.take("teach")
    r = EN.options(db, "teach", a, st, P.context(db, "teach", dt.datetime.fromisoformat(st["factory_now"])))
    assert [o["kind"] for o in r["options"]] == ["none", "tou"] and r["options"][1]["changes"][0]["type"] == "tou_rule"
    aid = AC.create(db, "teach", PLN, "tou_rule", r["options"][1]["changes"][0], None, a["id"], None)
    AC.submit(db, aid, PLN, None)
    AC.approve(H, aid, TEA, TEA["name"])
    assert EN.settings(db, "teach")["tou"] == ["ht-01"]


def test_fix_options_and_clearing_faults(env):
    from hub import energy as EN
    H, _ = env
    db = H.db
    for rule, typ in (("N1", "repair_order"), ("N4", "rectify_order")):
        a = db.one("select * from anomaly where mode='teach' and rule=%s order by id limit 1", (rule,))
        r = EN.fix_options(db, "teach", a)
        assert r["options"][0]["changes"][0]["type"] == typ and r["options"][-1]["kind"] == "none"
    r = EN.inject(H, "teach", "leak", "老师")
    assert H.mes.sent[-1]["inject"] == [{"kind": "leak"}] and "压缩空气泄漏" in r["text"]
    with pytest.raises(ValueError):
        EN.inject(H, "teach", "forgot", "老师", unit="ht-01")
    with pytest.raises(PermissionError):
        EN.inject(H, "prod", "leak", "老师")
    db.x("insert into action (mode, type, title, params, status, executed_at) values ('teach','rectify_order','查漏', %s, 'done', now() + interval '1 second')",
         (json.dumps({"energy_fix": "leak", "unit": "air"}),))
    out = EN.tick(H, "teach", now=dt.datetime.now(dt.timezone.utc) + dt.timedelta(seconds=5))
    assert "leak" in out["cleared"] and H.mes.sent[-1].get("clear") == [{"kind": "leak", "unit": None}]


def test_simulator_injects_and_reset(env):
    import random
    from hub import energy as EN
    H, _ = env
    db = H.db
    EN.set_sim(db, "teach", True)
    st = EN._state(db, "teach")
    st["auto_at"] = 1.0
    EN._save_state(db, "teach", st)
    out = EN.tick(H, "teach", rng=random.Random(3))
    assert out["injected"] or EN.live_faults(db, "teach")
    st = EN._state(db, "teach")
    st["auto_at"] = dt.datetime.now(dt.timezone.utc).timestamp()
    EN._save_state(db, "teach", st)
    assert not EN.tick(H, "teach")["injected"]                            # 30 分钟内不再出
    assert EN.measuring(db, "prod") is False and EN.sim_on(db, "prod") is False


def test_situation_report_and_sources(env):
    from hub.ont import report as RP
    from hub.ont import situation as SIT
    H, _ = env
    db = H.db
    d = H.ai.detector()
    cards = SIT.cards(db, d, "teach")
    c = next(x for x in cards["cards"] if x["domain"] == "energy")
    assert {m["key"] for m in c["metrics"]} >= {"kwh", "cost", "peak_share", "idle_share", "co2_kg"} and not c["waiting"]
    assert {"N1", "N2", "N3", "N4", "N5", "N6"} <= set(c["rules_live"])
    now = dt.datetime.now(dt.timezone.utc)
    data = RP.build(db, "teach", "daily", now - dt.timedelta(days=1), now)
    assert data["energy"]["kwh"] > 0 and "能源与环境" in RP.text_of("daily", data)
    assert db.one("select count(*) as n from ont_object where mode='teach' and type='Meter'")["n"] == 18


def test_agent_energy_tools_and_intent(env):
    from hub.ont.agent import Agent, Rules, Run, Tools
    H, _ = env
    t = Tools(H, TEA, Run(TEA))
    o = t.energy_overview()
    assert o["kwh"] > 0 and "series" not in o and o["units"][0]["ref"]
    assert t.energy_overview("2026-13-01").get("error")
    u = t.unit_energy("grd-01")
    assert u["unit_energy"] and all(x["unit"] == "grd-01" for x in u["unit_energy"]) and u["air"]["excess_kw"] > 3
    c = t.carbon_account()
    assert c["t_co2e"] > 0 and t.carbon_account("2026/10").get("error")
    r = Rules(None, False)
    assert r.intent("空压机是不是漏气了？") == "energy" and r.intent("这个月的碳排放是多少？") == "energy"
    assert r.intent("危废暂存超期了吗？") == "safety" and r.intent("昨天电费多少？") == "energy"
    out = Agent(H, NoLLM()).ask(TEA, [{"role": "user", "content": "昨天用了多少电、电费多少？"}])
    assert "电费" in out["answer"] and "energy_overview" in out["tools"]
    out = Agent(H, NoLLM()).ask(PLN, [{"role": "user", "content": "昨天用了多少电、电费多少？"}])
    assert out["answer"].startswith("提示") and not out["tools"]


def test_lab21_flow(env):
    from hub import energy as EN
    from hub import lab21
    from hub.ont import actions as AC
    H, _ = env
    db = H.db
    U = {"name": "学生计划丙", "role": "planner", "teacher": False, "mode": "teach"}
    with pytest.raises(PermissionError):
        lab21.inject(H, U, {})
    r = lab21.inject(H, TEA, {"enable": True})
    assert "HMC-01" in r["text"] and EN._state(db, "teach")["faults"]["forgot:hmc-01"]["lab"]
    assert not lab21.locked(db, U)
    lab21.start(db, U["name"])
    assert lab21.locked(db, U) and not lab21.locked(db, TEA)
    top = sorted(EN.idle_by_unit(db, "teach"), key=lambda x: -x["idle_kwh"])[0]["unit"]
    assert lab21.answer(db, U["name"], "l21_t1", {"unit": top})["score"] == 10
    kwh, co2 = lab21.expected_t2(db)
    assert kwh == pytest.approx(10.892, abs=0.01) and co2 == pytest.approx(5.779, abs=0.01)
    assert lab21.answer(db, U["name"], "l21_t2", {"kwh": 10.9, "co2_kg": 5.78})["score"] == 15
    a = lab21.target(db)
    assert a is not None
    best = (EN._state(db, "teach").get("recommended") or {}).get(str(a["id"])) or "standby10"
    assert lab21.answer(db, U["name"], "l21_t3", {"option": best, "reason": "每天省一百多度电和十几美元，预热三分钟不影响订单交期"})["score"] == 10
    assert not lab21.locked(db, U)
    ch = {"type": "energy_setting", "units": ["hmc-01"], "standby_min": 10}
    aid = AC.create(db, "teach", U, "energy_setting", ch, None, a["id"], None)
    AC.submit(db, aid, U, None)
    assert next(t for t in lab21.check(db, U["name"])["tasks"] if t["id"] == "l21_t4")["score"] == 5
    AC.approve(H, aid, TEA, TEA["name"])
    assert next(t for t in lab21.check(db, U["name"])["tasks"] if t["id"] == "l21_t4")["score"] == 15
    e = lab21.expected_t5(db)
    assert e["tou_usd"] == pytest.approx(363.0) and e["standby_co2"] == pytest.approx(1012.4, abs=0.1)
    assert lab21.answer(db, U["name"], "l21_t5", {"tou_usd": 363, "tou_co2": 0, "standby_co2": 1012})["score"] == 10
    st = lab21.check(db, U["name"])
    assert st["score"] == 60 == st["of"] and st["report_points"] == 40
    assert any(c["user"] == U["name"] for c in lab21.class_view(db))


def test_reset_keeps_measuring_turns_sim_off(env):
    from hub import energy as EN
    H, _ = env
    db = H.db
    EN.reset(db, "teach")
    assert EN.measuring(db, "teach") and not EN.sim_on(db, "teach")
    assert EN.settings(db, "teach") == {"standby": {}, "tou": [], "quiet": []}
    assert db.one("select count(*) as n from energy_read where mode='teach'")["n"] == 0


def test_lab21_documents_and_course():
    pytest.importorskip("docx")
    import io

    import docx

    from cae import labdoc
    g = labdoc.guide_docx("lab21")
    t = labdoc.report_template_docx("lab21")
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(g)).paragraphs)
    assert "空转" in text and "0.5306" in text
    tt = docx.Document(io.BytesIO(t))
    assert "实验 21" in tt.paragraphs[0].text


def test_idle_rule_always_reports_the_forgotten_machine(monkeypatch):
    """N2 只报空转最多的 3 台；老师 / 情景注入“下班忘关”的那台够条件就一定报（带锯床功率小，会被数控车床挤掉）"""
    from types import SimpleNamespace
    from hub import energy as EN
    from hub.ont.rules import energy as R
    rows = [dict(unit=u, name=u, kwh=k / 0.9, idle_kwh=k, idle_h=12.0, idle_share=0.9) for u, k in
            (("hmc-01", 72), ("cnc-l01-a", 66), ("vmc-01", 60), ("cnc-l01-b", 54), ("saw-01", 24))]
    monkeypatch.setattr(R, "_end", lambda c: dt.datetime.now(dt.timezone.utc))
    monkeypatch.setattr(EN, "idle_by_unit", lambda db, mode, hours, end: [dict(x) for x in rows])
    monkeypatch.setattr(EN, "params", lambda db, mode: {"price": {"flat": 0.1}})
    c = SimpleNamespace(db=None, mode="teach")
    p = {"hours": 12, "share": 0.3, "min_h": 1, "min_kwh": 10, "top": 3}
    monkeypatch.setattr(EN, "_state", lambda db, mode: {"faults": {}})
    assert [f.fkey for f in R.idle_energy(c, p, {})] == ["hmc-01", "cnc-l01-a", "vmc-01"]
    monkeypatch.setattr(EN, "_state", lambda db, mode: {"faults": {"forgot:saw-01": {"lab": True}}})
    assert [f.fkey for f in R.idle_energy(c, p, {})] == ["hmc-01", "cnc-l01-a", "vmc-01", "saw-01"]
