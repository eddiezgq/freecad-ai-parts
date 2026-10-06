"""第 11 轮：教材案例件 RJ-201、LS-101 的第 1 版模型——可以直接在“仿真与分析”里读入；尺寸与教材程序一致"""
import os
import sys

import pytest

from cae import geometry as G
from hub import case_parts as C

TEXTBOOK = os.path.join(os.path.dirname(__file__), "..", "..", "..", "textbook", "mechdesign", "ch33", "code")


@pytest.mark.parametrize("item,n_cyl", [("RJ-201", 5), ("LS-101", 5)])
def test_case_part_steps(item, n_cyl):
    step = C.step_bytes(item)
    faces, _, _ = G.faces(step)
    cyl = [f for f in faces if f["kind"] == "cylinder"]
    assert len(cyl) >= n_cyl                      # 各段外圆（RJ-201 还有内孔）
    assert any(f["kind"] == "torus" for f in faces) or len(faces) > n_cyl + 2   # 台阶圆角
    assert C.step_bytes("SH-999") is None


@pytest.mark.skipif(not os.path.isdir(TEXTBOOK), reason="没有教材目录")
def test_same_dimensions_as_textbook():
    sys.path.insert(0, os.path.abspath(TEXTBOOK))
    import _ls
    import _rj
    assert [(d, L) for d, L, _ in C.RJ_SEGS] == [(d, L) for d, L, _ in _rj.shaft_segments(_rj.SEAT)] and C.RJ_BORE == _rj.BORE
    assert [(d, L) for d, L, _ in C.LS_SEGS] == [(d, L) for d, L, _ in _ls.SEGS] and C.LS_FILLET == _ls.FILLET


def test_lab12_documents():
    import io
    import docx
    from cae import labdoc
    g = docx.Document(io.BytesIO(labdoc.guide_docx("lab12")))
    assert any("只限中间一圈" in p.text for p in g.paragraphs)
    t = docx.Document(io.BytesIO(labdoc.report_template_docx("lab12")))
    cells = {c.text for tb in t.tables for r in tb.rows for c in r.cells}
    assert {"6010", "只限中间一圈（铰支）", "RJ-201"} <= cells
