# -*- coding: utf-8 -*-
"""第 16 轮：流体分析 → 温度场的对流行；AI 一句话填表和结果解释（规则部分，不需要 OpenFOAM）"""
from hub import cfd_ai
from hub.cfd_api import thermal_rows


FACES = [{"id": 1, "area_mm2": 45006, "p_pa": 12.3, "h_w_m2k": 16.9, "heat_w": 30.0},
         {"id": 2, "area_mm2": 123017, "p_pa": -8.1, "h_w_m2k": 20.1, "heat_w": 111.0},
         {"id": 4, "area_mm2": 123017, "p_pa": -9.1, "h_w_m2k": 20.2, "heat_w": 112.0},
         {"id": 5, "area_mm2": 61500, "p_pa": 0.1, "h_w_m2k": 3.0, "heat_w": 5.0},
         {"id": 7, "area_mm2": 900, "p_pa": 0.0, "h_w_m2k": None, "heat_w": None}]


def test_thermal_rows_merge_and_skip():
    ex = {"thermal": [{"type": "heat_flux", "faces": [9], "q_w_m2": 100.0}, {"type": "convection", "faces": [1], "h_w_m2k": 17.45}],
          "groups": {"bottom": [5], "inner": [9]}}
    rows = thermal_rows(FACES, 25.0, ex)
    assert rows[0]["type"] == "heat_flux"                       # 示例的发热保留，查表的对流行不保留
    conv = [r for r in rows if r["type"] == "convection"]
    assert all(r["source"] == "cfd" and r["t_inf_c"] == 25.0 for r in conv)
    assert [r["faces"] for r in conv] == [[2, 4], [1]]          # 20.1 和 20.2 都取 20.0，合成一行；底面 5、没有 h 的 7 跳过
    assert conv[0]["h_w_m2k"] == 20.0 and conv[1]["h_w_m2k"] == 17.0
    assert [r["faces"] for r in thermal_rows(FACES, 20.0)] == [[2, 4], [1], [5]]   # 一般零件：每个有 h 的面都给对流


def test_ai_setup_rules():
    r = cfd_ai.rules_setup("轴端风扇 5 m/s 沿长边吹，空气 25 ℃，箱体表面 70 ℃，放在机座上", bbox=[0, 0, 0, 150, 410, 300])
    f = r["form"]
    assert f == {"velocity_ms": 5.0, "air_t_c": 25.0, "wall_t_c": 70.0, "heat": True, "ground": True, "direction": "+y"}
    assert not r["unmatched"]
    r = cfd_ai.rules_setup("36 km/h 从下往上，只算阻力")
    assert r["form"]["velocity_ms"] == 10.0 and r["form"]["direction"] == "+z" and r["form"]["heat"] is False
    assert cfd_ai.rules_setup("风量 30 CFM")["unmatched"]           # 看不懂的带数字的话交给模型（没有模型就提示学生）


def test_ai_explain_rules():
    st = {"velocity_ms": 5, "direction": "+x", "L_char_mm": 410, "Re": 131271, "turbulence": "kOmegaSST", "drag_n": 0.683,
          "drag_pressure_n": 0.689, "drag_friction_n": -0.006, "Cd": 1.03, "frontal_m2": 0.045, "q_dyn_pa": 14.8,
          "heat_w": 357.6, "h_mean": 20.0, "wall_t_c": 70, "air_t_c": 25, "converged": False}
    t = cfd_ai.explain(None, st, FACES)
    assert t["engine"] == "rules"
    assert "湍流" in t["text"] and "压差阻力" in t["text"] and "面 4" in t["text"] and "面 5" in t["text"] and "收敛" in t["text"]


def test_h_rad():
    from hub.cfd_api import h_rad
    assert abs(h_rad(70, 25) - 6.76) < 0.05                    # 涂漆铸铁 70 ℃ 对 25 ℃ 空气
    assert h_rad(70, 25, eps=0.1) < 1


def test_agv_box_example():
    """机器人例子：AGV 电控箱（前后风口三种布置、箱内驱动器和工控机），面分组和到温度场的行"""
    import pytest
    pytest.importorskip("build123d")
    from cae import geometry as G
    from cae import thermal_parts as TP
    p = TP.params_of("agv_box", {"layout": 2})
    faces = G.faces(TP.build("agv_box", p))[0]
    ex = TP.example("agv_box", p, faces)
    g = ex["groups"]
    assert len(g["drive"]) == 5 and len(g["pc"]) == 5 and len(g["bottom"]) == 1 and g["inner"] and g["outer"]
    assert ex["thermal"][0]["power_w"] == 60 and ex["formula"]["dt_air_c"] > 0
    fs = [{"id": f["id"], "area_mm2": 1000.0, "h_w_m2k": 10.0} for f in faces]
    rows = thermal_rows(fs, 25.0, dict(ex, key="agv_box"))
    conv = {i for r in rows if r["type"] == "convection" for i in r["faces"]}
    assert set(g["inner"]) <= conv and set(g["drive"]) <= conv and not set(g["bottom"]) & conv
    with pytest.raises(ValueError):
        TP.params_of("agv_box", {"vent_h": 200})


def test_lab14_documents():
    import io
    import docx
    from cae import labdoc
    g = docx.Document(io.BytesIO(labdoc.guide_docx("lab14")))
    text = "\n".join(p.text for p in g.paragraphs)
    assert "实验 14" in text and "散热系数" in text and "CFD" in text
    t = docx.Document(io.BytesIO(labdoc.report_template_docx("lab14")))
    cells = " ".join(c.text for tb in t.tables for r in tb.rows for c in r.cells)
    assert "17.45" in cells and "AGV 风口布置" in cells


def test_cfd_report():
    import io
    import docx
    from cae import cfd_report
    job = {"id": "x1", "title": "箱体 5 m/s", "owner_name": "甲", "created": 0,
           "setup": {"flow": {"velocity_ms": 5, "direction": "+x", "air_t_c": 25}, "heat": {"wall_t_c": 70}, "domain": {"ground": True}},
           "stats": {"cells": 100952, "cell_mm": 37.6, "level": 2, "turbulence": "kOmegaSST", "iterations": 800, "converged": False, "L_char_mm": 410,
                     "Re": 131271, "drag_n": 0.683, "drag_pressure_n": 0.689, "drag_friction_n": -0.006, "Cd": 1.03, "frontal_m2": 0.045,
                     "q_dyn_pa": 14.8, "h_mean": 20.0, "heat_w": 357.6, "area_m2": 0.3975, "wall_t_c": 70, "air_t_c": 25, "seconds": 222}}
    aux = {"faces": [dict(f, force_n=[0.1, 0, 0], heat_w=f["heat_w"] or 0) for f in FACES]}
    d = docx.Document(io.BytesIO(cfd_report.build(job, aux, [], "第一句。\n第二句。")))
    text = "\n".join(p.text for p in d.paragraphs)
    cells = " ".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "流体分析报告" in text and "第二句" in text and "k-ω SST" in cells and "20.00" in cells and "面 4" in cells
