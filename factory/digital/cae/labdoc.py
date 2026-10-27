# -*- coding: utf-8 -*-
"""实验 8 的 Word 文件（第 11 轮 F8）：指导书（由仓库里的 Markdown 转成 Word）和实验报告模板。"""
import io
import os
import re

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GUIDE_MD = os.path.join(HERE, "实验8_输出轴强度与疲劳校核.md")
GUIDES = {"lab8": GUIDE_MD, "lab9": os.path.join(HERE, "实验9_机械臂关节力矩与电机选型.md"),
          "lab10": os.path.join(HERE, "实验10_输出轴数控车削与键槽铣削编程.md"),
          "lab11": os.path.join(HERE, "实验11_减速器热平衡与输出轴轻量化.md"),
          "lab12": os.path.join(HERE, "实验12_一根轴的数字化设计与两个典型案例.md"),
          "lab13": os.path.join(HERE, "实验13_减速器数字孪生与预测性维护.md"),
          "lab14": os.path.join(HERE, "实验14_减速器风扇冷却的流体与传热分析.md"),
          "lab15": os.path.join(HERE, "实验15_车间异常的追因推演与决策.md"),
          "lab16": os.path.join(HERE, "实验16_客户投诉的追溯召回与服务决策.md"),
          "lab17": os.path.join(HERE, "实验17_安全隐患的识别评估与停机决策.md"),
          "lab18": os.path.join(HERE, "实验18_设备维护策略保养时机的决策.md"),
          "lab19": os.path.join(HERE, "实验19_排班与技能缺人时的决策.md"),
          "lab20": os.path.join(HERE, "实验20_供应链来料不合格与供应商选择.md"),
          "lab21": os.path.join(HERE, "实验21_能源空转错峰与单件碳足迹.md"),
          "lab22": os.path.join(HERE, "实验22_经营订单毛利应收与现金.md"),
          "lab23": os.path.join(HERE, "实验23_两周情景一张急单的全厂决策.md"),
          "lab24": os.path.join(HERE, "实验24_一台协作机器人从订单到交付.md"),
          "lab25": os.path.join(HERE, "实验25_关节背隙异常的追因与8D.md")}


def _doc():
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        d.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
    return d


def _runs(p, text):
    """**粗体** 和 `代码` 两种行内格式"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            p.add_run(part[1:-1]).font.name = "Consolas"
        else:
            p.add_run(part)


def md_to_docx(md):
    """够用的 Markdown → Word：标题、段落、列表、表格、粗体"""
    d = _doc()
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln:
            i += 1
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            t = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    cell = t.cell(r, c)
                    cell.text = ""
                    _runs(cell.paragraphs[0], v)
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            d.add_paragraph()
            continue
        m = re.match(r"(#{1,4})\s+(.*)", ln)
        if m:
            level = len(m.group(1))
            d.add_heading(m.group(2), 0 if level == 1 else level - 1)
        elif re.match(r"\s*\d+\.\s", ln):
            _runs(d.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s", "", ln))
        elif re.match(r"\s*[-*]\s", ln):
            _runs(d.add_paragraph(style="List Bullet 2" if ln.startswith("   ") else "List Bullet"), re.sub(r"^\s*[-*]\s", "", ln))
        else:
            _runs(d.add_paragraph(), ln.strip())
        i += 1
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def guide_docx(lab="lab8"):
    return md_to_docx(open(GUIDES[lab], encoding="utf-8").read())


def report_template_docx(lab="lab8"):
    if lab == "lab9":
        return lab9_template()
    if lab == "lab10":
        return lab10_template()
    if lab == "lab11":
        return lab11_template()
    if lab == "lab12":
        return lab12_template()
    if lab == "lab13":
        return lab13_template()
    if lab == "lab14":
        return lab14_template()
    if lab == "lab15":
        return lab15_template()
    if lab == "lab16":
        return lab16_template()
    if lab == "lab17":
        return lab17_template()
    if lab == "lab18":
        return lab18_template()
    if lab == "lab19":
        return lab19_template()
    if lab == "lab20":
        return lab20_template()
    if lab == "lab21":
        return lab21_template()
    if lab == "lab22":
        return lab22_template()
    if lab == "lab23":
        return lab23_template()
    if lab in ("lab24", "lab25"):
        return robot_template(lab)
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 8　输出轴强度与疲劳校核　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、实验目的", 1)
    hint("（用自己的话写 3–5 条）")
    for _ in range(3):
        d.add_paragraph("", style="List Number")

    d.add_heading("二、计算模型", 1)
    table([["项目", "设置", "理由"],
           ["零件与版本", "SH-301 rev __", ""],
           ["材料", "", ""],
           ["约束 1", "面 __：", ""],
           ["约束 2", "面 __：", ""],
           ["约束 3", "面 __：", ""],
           ["载荷", "面 __：", ""],
           ["网格", "单元尺寸 __ mm，__ 个单元", ""]])
    hint("贴设置截图（模型上的彩色面）。说明：为什么轴承只限制径向、联轴器只限制转动？试“左端面固定”时结果有什么不同？")

    d.add_heading("三、结果与手算核对", 1)
    table([["项目", "网格“中”", "网格“细”", "变化 %"],
           ["最大 Von Mises 应力 / MPa", "", "", ""],
           ["最大应力位置（面）", "", "", ""],
           ["安全系数", "", "", ""],
           ["最大位移 / mm", "", "", ""],
           ["Ø35 光滑段应力（平均）/ MPa", "", "", ""]])
    hint("手算：τ = 16T/(πd³) = ________ MPa；σ = √3·τ = ________ MPa；与有限元误差 ______ %。")
    hint("讨论：网格加密后光滑段与键槽根部的应力为什么变化不同？（应力集中、尖角奇异）")

    d.add_heading("四、疲劳寿命", 1)
    table([["载荷谱", "Miner 规则", "S_D / MPa", "每块循环数", "寿命 / h", "寿命 / 次"],
           ["恒幅 0～350 N·m", "Haibach", "", "", "", ""],
           ["恒幅 0～350 N·m", "原始（低于 S_D 不损伤）", "", "", "", ""],
           ["跑合试验台记录 ______", "Haibach", "", "", "", ""],
           ["跑合试验台记录 ______", "原始", "", "", "", ""]])
    hint("手算核对一个循环：σa = ____，σm = ____，σa,eq = σa/(1 − σm/σb) = ____ MPa；N = N_D·(S_D/σa,eq)^k = ____ 次；平台给出 ____ 次。")
    hint("讨论：损伤主要来自曲线的哪一段？两种 Miner 规则哪个偏安全，为什么？")

    d.add_heading("五、改进方案对比", 1)
    table([["方案", "改动", "安全系数", "疲劳寿命 / h", "下料重量 / kg", "评价（成本、加工）"],
           ["原设计", "45 钢，Ø40", "", "", "", ""],
           ["方案 A", "换 40Cr", "", "", "", ""],
           ["方案 B", "键槽轴段 Ø44", "", "", "", ""],
           ["方案 C", "台阶圆角 1.0 mm（B 版做法）", "", "", "", ""]])
    table([["尖角与圆角", "台阶根部最大应力 MPa", "最大应力位置", "中网格 → 细网格变化"], ["A 版（尖角）", "", "", ""], ["加 1.0 mm 圆角", "", "", ""]])
    hint("AI 解释哪里说得对、哪里需要补充？你推荐哪个方案，理由是什么？尖角处的应力为什么随网格一直变大？")

    d.add_heading("六、结论", 1)
    d.add_paragraph("")
    d.add_heading("七、附件", 1)
    d.add_paragraph("平台生成的计算报告（Word）。", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab9_template():
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 9　机械臂关节力矩与电机选型　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、静态重力矩（任务 1）", 1)
    table([["关节", "平台 3 kg / N·m", "平台 5 kg / N·m", "手算 3 kg / N·m", "误差 %"],
           ["J2 大臂", "", "", "", ""], ["J3 小臂", "", "", "", ""], ["J4 腕 1", "", "", "", ""]])
    hint("手算过程：J2 之后各构件的质量、质心到 J2 轴线的水平距离，τ = Σ m·g·r = …")
    d.add_heading("二、搬运结果（任务 2、3）", 1)
    table([["方案", "J1 峰值 N·m", "J2 峰值 N·m", "J1 峰值功率 W", "J2 平均功率 W", "峰值时刻 s"],
           ["1.2 秒，加速段 25%", "", "", "", "", ""],
           ["1.8 秒（放长 1.5 倍）", "", "", "", "", ""],
           ["1.2 秒，加速段 40%", "", "", "", "", ""]])
    hint("贴驱动力矩、速度曲线。说明峰值出现在梯形速度的哪一段；J2 平均功率为负的原因；重力部分与惯性部分各占多少。")
    d.add_heading("三、电机选型（任务 4）", 1)
    table([["关节", "峰值 N·m", "均方根 N·m", "最高 r/min", "推荐（安全系数 1.2）", "最紧的一项", "安全系数 1.5 时"]] +
          [[j, "", "", "", "", "", ""] for j in ("J1", "J2", "J3", "J4", "J5", "J6")])
    hint("手算核对 J2：电机侧均方根力矩 τ_rms/(i·η) = ____ N·m（额定 ____）；电机转速 ω_max·i = ____ r/min（上限 ____）。")
    d.add_heading("四、曲柄滑块（任务 5）", 1)
    table([["项目", "600 r/min", "60 r/min"], ["驱动力矩峰值 N·m", "", ""], ["θ = 90° 时平台值 N·m", "", ""],
           ["θ = 90° 时虚功原理手算 N·m", "", ""], ["连杆最大受力 N", "", ""], ["连杆安全系数", "", ""], ["连杆疲劳寿命", "", ""]])
    d.add_heading("五、结论与讨论", 1)
    d.add_paragraph("")
    d.add_heading("六、附件", 1)
    d.add_paragraph("平台生成的动力学报告、有限元报告（Word）。", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab10_template():
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 10　输出轴数控车削与键槽铣削编程　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、粗车编程单（任务 1）", 1)
    table([["部位", "图纸尺寸", "粗车工序尺寸", "编程直径（平台）", "编程直径（手算）", "说明"],
           ["轴承位", "Ø35 k6", "Ø36.5 0/−0.25", "", "", ""], ["齿轮位", "Ø40 k6", "Ø41.5 0/−0.25", "", "", ""],
           ["轴伸", "Ø30", "（未列）", "", "", "其余直径加的余量 ="]])
    table([["项目", "平台", "手算"], ["端面余量（每端）mm", "", ""], ["车端面刀数", "", ""], ["Ø50 处转速 r/min", "", ""],
           ["轴端编程直径处转速 r/min", "", ""], ["G50 限速什么时候起作用", "", ""]])
    hint("说明为什么编程直径取公差带中间；为什么一道粗车工序要两次装夹、两个程序。")
    d.add_heading("二、程序读懂（任务 1、3、4）", 1)
    table([["程序段（行号）", "在干什么", "关键指令和参数"], ["", "", ""], ["", "", ""], ["", "", ""], ["", "", ""], ["", "", ""]])
    hint("至少说明：安全行、换刀与主轴（G50、G96）、分层粗车、沿轮廓去台阶、精车轮廓、键槽斜线下刀、程序结尾。")
    d.add_heading("三、试切检查（任务 2）", 1)
    table([["项目", "O1201 右端", "O1202 左端", "合计"], ["仿真与本工序尺寸的差 mm", "", "", ""], ["切削时间", "", "", ""],
           ["快移时间", "", "", ""], ["换刀时间", "", "", ""], ["程序合计", "", "", ""]])
    table([["对比", "分钟"], ["程序合计", ""], ["基本时间公式 t_b = L·i/(n·f)", ""], ["工艺规程工时", "18"], ["差值主要来自", ""]])
    table([["功率", "ap 2.5、f 0.3", "ap 6、f 0.5"], ["F_c = k_c1.1·a_p·f^(1−m_c)  N", "", ""], ["P_c = F_c·v_c/60000  kW", "", ""],
           ["P_c / 0.8  kW（机床 11 kW）", "", ""], ["平台检查结果", "", ""]])
    d.add_heading("四、精车与一句话编程（任务 3）", 1)
    table([["项目", "结果"], ["精车毛坯从哪来", ""], ["精车编程直径", ""], ["仿真差 mm", ""], ["一句话：AI 改了哪几项", ""],
           ["多留 0.3 后粗车目标直径", ""], ["提交审批时的提示", ""]])
    d.add_heading("五、铣键槽（任务 4）", 1)
    table([["项目", "平台", "手算"], ["槽深 mm", "", ""], ["主轴转速 S r/min", "", ""], ["进给 F mm/min", "", ""],
           ["加工时间（默认参数）", "", "—"], ["加工时间（分 4 层、fz 0.04）", "", "—"], ["工艺规程工时", "10 分", "—"]])
    hint("哪组参数更合理？为什么键槽在磨削前铣、却按磨削后的尺寸 d − t₁ 控制？斜线下刀的好处？")
    d.add_heading("六、下发（任务 5）", 1)
    table([["项目", "记录"], ["提交号 / 工艺规程版本", ""], ["AI 评审意见与处理", ""], ["批准人", ""], ["下发的程序号", ""],
           ["ERPNext 附件 / 3D 回放 / 派工指令（贴图）", ""]])
    d.add_heading("七、结论与讨论", 1)
    d.add_paragraph("")
    d.add_heading("八、附件", 1)
    d.add_paragraph("平台生成的程序（.nc）、截图。", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab11_template():
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 11　减速器热平衡与输出轴轻量化　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、热平衡手算与有限元（任务 1）", 1)
    table([["项目", "手算", "平台"], ["总效率 η = 0.97² × 0.99³", "", "0.913"], ["输出功率 P_out  W", "", ""], ["输入功率 P_in  W", "", ""],
           ["发热 Φ = P_in(1 − η)  W", "", "482.6"], ["散热面积 A  m²（不计底面）", "", ""], ["油温 t = t₀ + Φ/(K_s·A)  ℃", "", ""]])
    table([["有限元结果", "数值", "位置 / 说明"], ["最高温度 ℃", "", ""], ["外表面平均温度 ℃", "", ""], ["热平衡：进来 / 散走 W", "", ""],
           ["与公式的差 ℃", "", ""]])
    hint("最高温度为什么比平均高？有限元比公式多告诉了你什么？")
    d.add_heading("二、降温方案比较（任务 2）", 1)
    table([["方案", "散热系数 / 面积", "公式油温 ℃", "外表面平均 ℃", "最高 ℃", "代价"], ["无散热筋，通风良好", "", "", "", "", ""],
           ["每侧 8 条散热筋", "", "", "", "", ""], ["轴端装风扇", "", "", "", "", ""]])
    table([["一句话设置", "记录"], ["写的那句话", ""], ["AI 选的面、散热系数", ""], ["需要手动改的地方", ""]])
    d.add_heading("三、热应力（任务 3）", 1)
    table([["项目", "手算 / 结果"], ["线膨胀系数 α（10⁻⁶/K）", ""], ["自由热伸长 ΔL = α·L·Δt  mm", ""], ["有限元最大热变形 mm", ""],
           ["最大热应力 MPa / 位置", ""], ["安全系数", ""]])
    hint("热应力从哪来？实际减速器怎样减小热应力？")
    d.add_heading("四、输出轴轻量化（任务 4）", 1)
    table([["项目", "现行设计", "最优方案"], ["齿轮位直径 mm", "40", ""], ["轴伸直径 mm", "30", ""], ["质量 kg（平台）", "", ""],
           ["质量 kg（手算，不扣键槽、倒角）", "", "—"], ["安全系数", "", ""], ["疲劳寿命", "", ""]])
    table([["问题", "回答"], ["一句话：AI 填了哪些", ""], ["第几次以后很难再变好", ""], ["不满足要求的几次被哪条要求挡住", ""],
           ["卡住最优解的要求", ""], ["安全系数要求提到 1.5 会怎样", ""], ["送设计台后校核结果、键要不要换", ""]])
    d.add_heading("五、选做：CPU 散热片 / 拓扑优化（任务 5、6）", 1)
    table([["项目", "结果"], ["散热片 h = 50 最高温度 ℃", ""], ["散热片 h = 8 最高温度 ℃", ""], ["散热片优化：最轻方案", ""],
           ["拓扑优化 50%：柔度、形状", ""], ["拓扑优化 30%：柔度、形状", ""], ["拉伸成板件后安全系数", ""]])
    d.add_heading("六、结论与讨论", 1)
    d.add_paragraph("")
    d.add_heading("七、附件", 1)
    d.add_paragraph("平台生成的温度场报告（Word）、优化结果截图。", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab12_template():
    """实验 12（第 11 轮《机械设计》第 33 章）：零件库 → 建模 → 强度、刚度、固有频率、疲劳 → 发布；RJ-201、LS-101 两个案例"""
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 12　一根轴的数字化设计与两个典型案例　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、零件库选件（任务 1）", 1)
    table([["轴承", "d mm", "D mm", "B mm", "r_s mm", "用在"], ["6207", "", "", "", "", "SH-301"], ["6010", "", "", "", "", "RJ-201"],
           ["6201", "", "", "", "", "LS-101"]])
    table([["CSF-25-100", "额定 N·m", "启停峰值 N·m", "平均负载 N·m", "瞬时最大 N·m", "输入最高 r/min"], ["零件库", "", "", "", "", ""]])
    hint("台阶圆角是否小于轴承内圈倒角？不小于会怎样？")
    d.add_heading("二、SH-301 B 版建模与发布（任务 2）", 1)
    table([["项目", "记录"], ["各段直径 × 长度", ""], ["台阶圆角", ""], ["键槽", ""], ["下料重量（设计台）kg", ""],
           ["教材 33.9.4 节的质量 kg", ""], ["发布的版本号", ""]])
    d.add_heading("三、RJ-201：扭转刚度与关节力矩载荷谱（任务 3）", 1)
    table([["项目", "手算", "有限元"], ["J2 最大力矩 N·m（运动与动力分析）", "", "—"], ["扭转角 θ μrad", "", ""], ["差别 % 及原因", "", ""],
           ["末端偏移 δ = θ × 850 mm", "", ""], ["是否 ≤ 0.20 mm", "", ""]])
    table([["疲劳", "记录"], ["载荷谱（计算编号、关节）", ""], ["最大应力幅 MPa", ""], ["修正后疲劳极限 MPa", ""], ["寿命 h", ""], ["为什么这么长", ""]])
    table([["材料", "扭转角 μrad", "末端偏移 mm", "质量 kg"], ["40Cr（调质）", "", "", ""], ["7075-T6", "", "", ""]])
    hint("铝轴要达到同样刚度需要多大直径？减重主要来自材料还是来自直径？")
    d.add_heading("四、LS-101：轴承支承与固有频率（任务 4）", 1)
    table([["约束方式", "一阶 Hz", "二阶 Hz", "三阶 Hz"], ["整个轴承位限径向", "", "", ""], ["只限中间一圈（铰支）", "", "", ""],
           ["梁理论（简支光轴）", "", "—", "—"]])
    table([["项目", "记录"], ["5 000 r/min 对应 Hz", ""], ["一阶固有频率 / 转频", ""], ["哪种约束与梁理论一致、为什么", ""],
           ["强度计算用哪种约束", ""]])
    d.add_heading("五、LS-101：CSV 载荷谱（任务 5）", 1)
    table([["项目", "记录"], ["CSV 点数、最大 |T| N·m", ""], ["一转的时间 s", ""], ["扭转疲劳寿命", ""], ["弯曲：σa、σm MPa（手算 / 程序 33.11.1）", ""],
           ["无限寿命安全系数 S", ""]])
    hint("平台的载荷谱为什么表达不了“轴在转、力也在转”？")
    d.add_heading("六、小结：三根轴的尺寸由什么决定", 1)
    table([["轴", "决定尺寸的因素", "依据（本实验的数据）"], ["SH-301", "", ""], ["RJ-201", "", ""], ["LS-101", "", ""]])
    d.add_heading("七、附件", 1)
    d.add_paragraph("平台生成的计算报告（Word）：RJ-201 扭转与疲劳、LS-101 固有频率（两种约束）、LS-101 扭转疲劳。", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab13_template():
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 13　减速器数字孪生与预测性维护　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、读懂一台的分身（任务 1）", 1)
    table([["项目", "记录"], ["序列号 / 客户 / 起运行日期", ""], ["工单、输出轴序列号、设计版本", ""], ["跑合峰值转矩 N·m", ""],
           ["运行小时 / 起动次数", ""], ["最先到期的轴承、已用 L10", ""], ["输出轴疲劳损伤", ""], ["换油剩余 / 油温残差 / 振动", ""]])
    hint("为什么总是中间轴轴承 B 最先到期？")
    d.add_heading("二、手算核对（任务 2）", 1)
    table([["项目", "手算", "平台"], ["中间轴转矩 T₂ N·m", "", ""], ["第二级 Ft / Fr / F N", "", ""], ["输入轴转矩 T₁ N·m", "", ""],
           ["第一级 Ft / Fr N", "", ""], ["中间轴轴承 B 径向力 P N", "", "2990"], ["中间轴转速 r/min", "", ""],
           ["L10 h（额定）", "", "约 10 800"], ["85% 负载时寿命倍数 1/0.85³", "", ""], ["某天热平衡油温 ℃", "", ""]])
    hint("为什么不能直接用平均转矩算轴承寿命？")
    d.add_heading("三、诊断（任务 3）", 1)
    table([["设备", "异常的数据", "判断的原因", "依据", "AI 的说法与你的异同"], ["港口（红）", "", "", "", ""], ["矿山（振动）", "", "", "", ""],
           ["港口（换油过期）", "", "", "", ""]])
    d.add_heading("四、预测性维护（任务 4）", 1)
    table([["项目", "记录"], ["维修单标题 / 优先级", ""], ["ERPNext 服务工单号", ""], ["维修记录", ""],
           ["完成前后：轴承已用、振动、报警状态", ""], ["ERPNext 工单状态", ""], ["老师制造的故障：哪一台、什么问题、怎么发现的", ""]])
    d.add_heading("五、假如与反馈设计（任务 5）", 1)
    table([["假如", "最先到期的轴承约剩", "输出轴约剩", "换油间隔", "平均油温"], ["现在", "", "", "", ""], ["负载 80%", "", "", "", ""],
           ["加风扇", "", "", "", ""], ["输出轴 40Cr", "", "", "", ""]])
    table([["机队反馈", "记录"], ["轴承当量转矩最高的客户 / 占额定", ""], ["中间轴轴承每年用掉 L10", ""], ["平台的设计建议", ""], ["你的意见", ""]])
    d.add_heading("六、选做（任务 6、7）", 1)
    d.add_paragraph("")
    d.add_heading("七、结论与讨论", 1)
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab14_template():
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 14　减速器风扇冷却的流体与传热分析　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、标准题核对（任务 1）", 1)
    table([["题", "教材公式 / 文献", "平台", "差"], ["平行平板层流压降", "", "", ""], ["平板层流平均 Nu", "", "", ""],
           ["湍流通道 C_f", "", "", ""], ["湍流通道 Nu", "", "", ""], ["绕圆柱 C_d（Re 40）", "", "", ""]])
    d.add_paragraph("第 2 题手算过程：")
    hint("为什么平板换热的算例要在平板前面留一段来流区？")
    d.add_heading("二、箱体风扇冷却（任务 2）", 1)
    table([["项目", "手算 / 记录"], ["Re = UL/ν", ""], ["层流边界层 h", ""], ["全湍流 h", ""], ["平台平均 h", ""],
           ["阻力 / 压差 / 摩擦 N", ""], ["阻力系数 C_d", ""]])
    table([["面", "h W/(m²·K)", "平均压力 Pa", "为什么大 / 小"], ["迎风端面", "", "", ""], ["侧面（前）", "", "", ""], ["侧面（后）", "", "", ""],
           ["顶面", "", "", ""], ["背风端面", "", "", ""]])
    d.add_paragraph("截图：表面压力、表面 h、流线（各一张）")
    d.add_paragraph("AI 解释里对和不对的地方：")
    d.add_heading("三、送到温度场（任务 3）", 1)
    table([["散热系数来源", "K_s W/(m²·℃)", "公式油温 ℃", "有限元外表面平均 ℃", "有限元最高 ℃"], ["查表“通风良好”", "17.45", "", "", ""],
           ["查表“风扇冷却”", "31", "", "", ""], ["CFD + 辐射", "", "", "", ""]])
    hint("CFD 的 h 为什么比“风扇冷却 31–38”小？去掉辐射会差多少？")
    d.add_heading("四、改一个条件（任务 4）", 1)
    table([["改了什么", "平均 h", "总散热 W", "阻力 N", "解释"], ["基准（任务 2）", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""]])
    d.add_heading("五、选做（任务 5、6）", 1)
    table([["AGV 风口布置", "驱动器平均 h", "工控机平均 h", "元件最高温度 ℃"], ["1 前下进、后上出", "", "", ""],
           ["2 前上进、后上出", "", "", ""], ["3 前下进、后下出", "", "", ""]])
    d.add_heading("六、结论与讨论", 1)
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab15_template():
    """实验 15 报告模板（第 17 轮 D 期 D9）"""
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 15　车间异常的追因、推演与决策　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、出了什么事（任务 1）", 1)
    table([["项目", "记录"], ["异常编号与标题", ""], ["发现时间（工厂时间）", ""], ["你从哪几页、哪些数据发现的", ""]])
    d.add_heading("二、我的因果链与中枢的因果链", 1)
    table([["环节", "我手工追的", "中枢的因果链", "一致吗？不一致的原因"], ["起因（设备 / 物料 / 质量）", "", "", ""], ["受影响的工序 / 工单", "", "", ""],
           ["受影响的订单、客户", "", "", ""], ["不处理时晚几天、多花多少", "", "", ""]])
    hint("每一环写依据：哪一页、哪张单、哪条总线消息（#编号）。")
    d.add_heading("三、推演的方案（任务 2）", 1)
    table([["场景", "改动", "订单预计交付 / 晚几天", "多花的钱", "加班", "校验"], ["SC-", "不处理", "", "", "", ""], ["SC-", "", "", "", "", ""],
           ["SC-", "", "", "", "", ""]])
    d.add_heading("四、我的选择（任务 3）", 1)
    table([["项目", "记录"], ["选的方案", ""], ["理由（交期、成本、风险）", ""], ["代价", ""]])
    d.add_heading("五、和中枢的建议比较", 1)
    table([["项目", "记录"], ["中枢建议的方案与理由", ""], ["和我的选择一样吗", ""], ["不一样时，谁更合理？为什么", ""],
           ["AI 分析（“让 AI 分析”）说了什么、哪句话有依据、哪句没有", ""]])
    d.add_heading("六、执行与复盘（任务 4、5）", 1)
    table([["项目", "记录"], ["动作编号、批准人、执行结果", ""], ["推演的完工时刻 / 实际完工时刻", ""], ["偏差（小时、%）", ""],
           ["复盘找到的原因", ""], ["“如果当时选别的”：哪个更好", ""]])
    d.add_heading("七、结论与讨论", 1)
    hint("推演靠得住吗？什么情况下推演会偏？人和 AI 在决策里各该做什么？")
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab16_template():
    """实验 16 报告模板（第 18 轮 E1-12）"""
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 16　客户投诉的追溯、召回与服务决策　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、投诉（任务 1）", 1)
    table([["项目", "记录"], ["投诉编号、客户、现象", ""], ["设备序列号、投运日期、是否在保修期内", ""], ["合同响应时限 / 实际首次响应时间", ""]])
    d.add_heading("二、追溯链：我追的与系统的", 1)
    table([["环节", "我手工追的", "系统的因果链", "一致吗？不一致的原因"], ["设备 → 出厂总成 → 输出轴", "", "", ""], ["输出轴 → 工单、磨削砂轮", "", "", ""],
           ["同一片砂轮磨的其他轴", "", "", ""], ["装在哪些设备、哪些客户（排查范围）", "", "", ""]])
    hint("每一环写依据：哪一页、哪个对象卡、哪条链接。")
    table([["排查范围里的设备", "客户", "健康度", "最先到期的轴承已用", "保修内？", "高风险？（任务 2）"], ["", "", "", "", "", ""], ["", "", "", "", "", ""],
           ["", "", "", "", "", ""]])
    d.add_heading("三、三个服务方案（任务 3）", 1)
    table([["方案", "设备 / 客户数", "工程师", "几天做完", "备件", "费用（保修内 / 向客户收）", "漏掉的高风险设备"], ["A 只修出问题的", "", "", "", "", "", ""],
           ["B 修 + 高风险点检", "", "", "", "", "", ""], ["C 整批排查", "", "", "", "", "", ""]])
    d.add_heading("四、我的选择与中枢的建议", 1)
    table([["项目", "记录"], ["选的方案", ""], ["理由（费用、风险、客户）", ""], ["代价", ""], ["中枢建议的方案与理由", ""],
           ["一样吗？不一样时谁更合理", ""], ["AI 助手怎么回答的，哪句有依据、哪句没有", ""]])
    d.add_heading("五、执行与回访（任务 4、5）", 1)
    table([["项目", "记录"], ["动作编号、批准人、执行结果", ""], ["ERPNext 服务单（Issue）、备件出库单编号", ""], ["客户通知（发给谁、内容要点）", ""],
           ["计划上门日 / 实际完成日", ""], ["回访分数与扣分原因", ""], ["保修内工厂实际承担的费用", ""]])
    d.add_heading("六、结论与讨论", 1)
    hint("召回范围怎么定？什么时候值得整批排查？保修费用算在哪里、怎样从源头减少？")
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab17_template():
    """实验 17 报告模板（第 18 轮 E2-9）"""
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 17　安全隐患的识别、评估与停机决策　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、风险评估（任务 1）", 1)
    table([["项目", "记录"], ["隐患编号、设备、描述", ""], ["对应的危险源（编号、名称）", ""], ["原来的可能性 × 严重度", ""],
           ["防护缺失后的可能性 × 严重度、风险值、等级", ""], ["为什么改的是可能性而不是严重度", ""]])
    d.add_heading("二、影响链：我追的与系统的（任务 2）", 1)
    table([["环节", "我手工追的", "系统的因果链", "一致吗？不一致的原因"], ["设备 → 派工队列里的工序", "", "", ""], ["工序 → 工单", "", "", ""],
           ["工单 → 订单、客户、交期", "", "", ""], ["整改需要多久（分钟）", "", "", ""]])
    d.add_heading("三、处理方案（任务 3）", 1)
    table([["方案", "什么时候停、停多久", "订单预计交付 / 晚几天", "多花的钱（整改、临时措施、加班）", "安全上的问题"], ["A 不处理（对照）", "", "", "", ""],
           ["B 立即停机挂牌", "", "", "", ""], ["C 限时整改", "", "", "", ""], ["D 停机 + 改派 / 加班（如有）", "", "", "", ""]])
    d.add_heading("四、我的选择与中枢的建议", 1)
    table([["项目", "记录"], ["选的方案", ""], ["理由（安全、交期、成本）", ""], ["代价", ""], ["中枢建议的方案与理由", ""],
           ["一样吗？不一样时谁更合理", ""], ["AI 助手怎么回答的，哪句有依据、哪句没有", ""]])
    d.add_heading("五、挂牌上锁全过程（任务 4、5）", 1)
    table([["步骤", "时间（工厂时间）", "谁", "记录"], ["起草、提交停机挂牌动作", "", "", "动作编号："], ["批准", "", "", ""],
           ["车间停机（许可证号）", "", "", "PTW-"], ["整改完成", "", "", ""], ["验收关闭、解锁", "", "", ""], ["订单的实际影响", "", "", ""]])
    d.add_heading("六、结论与讨论", 1)
    hint("限时整改什么时候可以接受？安全为什么不和交期、成本放在同一个天平上？挂牌上锁为什么要“两个人”？")
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab18_template():
    """实验 18 报告模板（第 18 轮 E3-9）"""
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 18　设备维护策略：保养时机的决策　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、手算故障概率（任务 1）", 1)
    table([["项目", "记录"], ["设备、保养周期", ""], ["β、η、可靠性系数 k", ""], ["自上次保养运行 t（小时）", ""],
           ["((t+80)k/η)^β − (tk/η)^β", ""], ["未来 80 运行小时坏掉的概率 P", ""], ["修理费 + 停机损失", ""], ["预期损失 = P × 损失", ""],
           ["一次保养的花费", ""]])
    hint("写出每一步的数，和平台算的比较。")
    d.add_heading("二、三个方案（任务 2）", 1)
    table([["方案", "什么时候停、停多久", "订单预计交付 / 晚几天", "确定的花费", "预期故障损失", "合计"], ["A 不做（对照）", "", "", "", "", ""],
           ["B 现在保养", "", "", "", "", ""], ["C 下班后保养", "", "", "", "", ""]])
    d.add_heading("三、我的选择与中枢的建议", 1)
    table([["项目", "记录"], ["选的方案", ""], ["理由（概率、损失、花费、交期）", ""], ["中枢建议的方案与理由", ""], ["一样吗？不一样时谁更合理", ""],
           ["AI 助手怎么回答的，哪句有依据、哪句没有", ""]])
    d.add_heading("四、保养全过程（任务 3）", 1)
    table([["步骤", "时间（工厂时间）", "谁", "记录"], ["起草、提交保养工单", "", "", "动作编号："], ["批准", "", "", ""], ["挂牌上锁停机", "", "", "PTW-"],
           ["保养完成", "", "", "PM-"], ["验收解锁", "", "", ""], ["ERPNext 保养记录", "", "", ""]])
    d.add_heading("五、对比与平均故障间隔（任务 4、5）", 1)
    table([["项目", "记录"], ["我保养的设备，保养后故障次数", ""], ["逾期最多的另一台，同期故障次数", ""], ["近 30 天全厂故障次数", ""],
           ["MTBF 估计（运行小时）", ""], ["这一次的结果和概率预测一致吗？为什么可能不一致", ""]])
    d.add_heading("六、结论与讨论", 1)
    hint("β = 1 时定期保养还有意义吗？只看预期损失什么时候不够？保养周期太短、太长各有什么代价？")
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab19_template():
    """实验 19 报告模板（第 18 轮 E4-10）"""
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 19　排班与技能：缺人时的决策　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、单点依赖（任务 1）", 1)
    table([["设备", "唯一熟练的人", "他在哪个班", "这台设备现在开几个班"], ["", "", "", ""], ["", "", "", ""], ["", "", "", ""], ["", "", "", ""]])
    d.add_heading("二、缺人一班的影响（任务 2）", 1)
    table([["项目", "记录"], ["哪天、哪个班、哪台设备、谁请假", ""], ["停多少小时", ""], ["每件分钟", ""], ["少做多少件（写算式）", ""],
           ["影响的工单、订单（因果链）", ""]])
    d.add_heading("三、四个方案（任务 3）", 1)
    table([["方案", "谁来、怎么排", "对别的班的影响", "订单预计交付 / 晚几天", "花费"], ["A 不处理（对照）", "", "", "", ""],
           ["B 调人顶岗", "", "", "", ""], ["C 学徒带教", "", "", "", ""], ["D 加班顶班", "", "", "", ""]])
    table([["项目", "记录"], ["选的方案", ""], ["理由", ""], ["中枢建议的方案与理由", ""], ["一样吗？不一样时谁更合理", ""]])
    d.add_heading("四、执行（任务 4）", 1)
    table([["步骤", "时间", "谁", "记录"], ["起草、提交", "", "", "动作编号："], ["批准", "", "", ""], ["排班表里这一格", "", "", ""],
           ["车间实况里这台设备", "", "", ""]])
    d.add_heading("五、技能培训计划（任务 5）", 1)
    table([["设备", "谁来学", "为什么选他", "带教小时", "花费", "多久见效", "培训期间排班"], ["", "", "", "", "", "", ""], ["", "", "", "", "", "", ""],
           ["", "", "", "", "", "", ""], ["", "", "", "", "", "", ""]])
    hint("花费 = 带教小时 × 工时 $30 ×（1 + 师傅分出去的 25%）。")
    d.add_heading("六、结论与讨论", 1)
    hint("调人顶岗为什么不总是最好？加班顶班看不见的代价？健康的技能矩阵每台设备至少几个熟练的人？")
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab20_template():
    """实验 20 报告模板（第 18 轮 E5-8）"""
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 20　供应链：来料不合格与供应商选择　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、单一来源（任务 1）", 1)
    table([["物料类别", "唯一的供应商", "有没有可开发的备选", "风险"], ["", "", "", ""], ["", "", "", ""], ["", "", "", ""], ["", "", "", ""]])
    d.add_heading("二、来料检验（任务 2）", 1)
    table([["项目", "记录"], ["批号、物料、数量、供应商", ""], ["检验模板与上下限", ""], ["抽样件数", ""], ["超差的件和读数", ""], ["判定", ""],
           ["ERPNext 检验单号、收货单号", ""]])
    d.add_heading("三、四个方案（任务 3）", 1)
    table([["方案", "到料日", "订单预计交付 / 晚几天", "花费", "质量风险"], ["A 不处理（对照）", "", "", "", ""], ["B 退货等补发", "", "", "", ""],
           ["C 备选急单", "", "", "", ""], ["D 让步接收全检", "", "", "", ""]])
    table([["项目", "记录"], ["选的方案", ""], ["理由", ""], ["中枢建议的方案与理由", ""], ["一样吗？不一样时谁更合理", ""]])
    d.add_heading("四、执行（任务 4）", 1)
    table([["步骤", "时间", "谁", "记录"], ["起草、提交", "", "", "动作编号："], ["批准", "", "", ""], ["ERPNext 单据", "", "", ""], ["这一批的状态", "", "", ""]])
    d.add_heading("五、总拥有成本（任务 5）", 1)
    table([["供应商", "单价 × 价格指数", "(1 − 合格率) × $4", "(1 − 准时率) × $3", "每件总拥有成本"], ["默认", "", "", "", ""], ["备选", "", "", "", ""]])
    hint("以后默认向哪家下单？为什么还要保留另一家？")
    d.add_heading("六、结论与讨论", 1)
    hint("抽样的风险；让步接收什么时候可以接受；单一来源一定不好吗？")
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab21_template():
    """实验 21 报告模板（第 18 轮 E6-10）"""
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 21　能源：空转、错峰与单件碳足迹　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、空转（任务 1）", 1)
    table([["设备", "近 12 小时电量（度）", "其中空转（度）", "空转占比", "原因"], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""]])
    d.add_heading("二、单件电量与碳足迹（任务 2）", 1)
    table([["工序", "设备", "运行功率 kW", "每件工时（分钟）", "每件电量（度）"], ["下料", "", "", "", ""], ["粗车", "", "", "", ""], ["精车", "", "", "", ""],
           ["铣键槽", "", "", "", ""], ["磨外圆", "", "", "", ""], ["检验", "", "", "", ""], ["合计", "", "", "", ""]])
    table([["项目", "记录"], ["排放因子（kg CO2 / 度）", ""], ["每件碳排放（kg）", ""], ["平台实测的单件电量（对照）", ""], ["差别的原因", ""]])
    d.add_heading("三、四个方案（任务 3）", 1)
    table([["方案", "未来 24 小时电量（度）", "电费", "碳排放", "订单交期"], ["A 不处理（对照）", "", "", "", ""], ["B 10 分钟待机", "", "", "", ""],
           ["C 30 分钟待机", "", "", "", ""], ["D 全部机加设备 10 分钟", "", "", "", ""]])
    table([["项目", "记录"], ["选的方案", ""], ["理由", ""], ["中枢建议的方案与理由", ""], ["一样吗？不一样时谁更合理", ""]])
    d.add_heading("四、执行（任务 4）", 1)
    table([["步骤", "时间", "谁", "记录"], ["起草、提交", "", "", "动作编号："], ["批准", "", "", ""], ["批准后 HMC-01 的电量条", "", "", ""]])
    d.add_heading("五、一个月省多少（任务 5）", 1)
    table([["项目", "算式", "结果"], ["错峰一个月省的电费", "", ""], ["错峰一个月少排的碳", "", ""], ["HMC-01 每晚待机一个月少排的碳", "", ""]])
    hint("错峰和待机，哪个是“省钱”，哪个是“省电”？")
    d.add_heading("六、结论与讨论", 1)
    hint("自动待机设多长合适；错峰对碳排放有没有用；开工率低时每件分摊的电为什么多")
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab22_template():
    """实验 22 报告模板（第 18 轮 E7-10）"""
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 22　经营：订单毛利、应收与现金　实验报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
    d.add_heading("一、没达标的经营指标（任务 1）", 1)
    table([["指标", "现在的值", "目标", "数从哪个领域来", "该去哪里找原因"], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""]])
    d.add_heading("二、订单毛利（任务 2）", 1)
    table([["项目", "金额 $", "来源（实测 / 标准）"], ["收入", "", ""], ["料", "", ""], ["工（折旧 + 人工）", "", ""], ["电", "", ""], ["废品返工", "", ""],
           ["急单、加班、外协、加急", "", ""], ["保修", "", ""], ["毛利 / 毛利率", "", ""], ["按标准的毛利率", "", ""]])
    hint("和按标准比差在哪？报价、工艺、供应商各能改什么？")
    d.add_heading("三、三个方案（任务 3）", 1)
    table([["方案", "预计几天收到钱", "现金最低的周与金额", "对客户关系的影响"], ["A 不处理（对照）", "", "", ""], ["B 发催款函", "", "", ""], ["C 信用冻结", "", "", ""]])
    table([["项目", "记录"], ["选的方案", ""], ["理由", ""], ["中枢建议的方案与理由", ""], ["一样吗？不一样时谁更合理", ""]])
    d.add_heading("四、执行（任务 4）", 1)
    table([["步骤", "时间", "谁", "记录"], ["起草、提交", "", "", "动作编号："], ["批准", "", "", ""], ["客户付款（收款单）", "", "", ""]])
    d.add_heading("五、应收周转与现金（任务 5）", 1)
    table([["项目", "数值"], ["应收余额", ""], ["近 30 天开票额", ""], ["DSO", ""], ["现金最紧的周与金额", ""], ["DSO 降到 30 天能多出的现金", ""]])
    d.add_heading("六、结论与讨论", 1)
    hint("毛利低的原因该找谁改；什么时候不催款、不冻结；为什么赚钱的工厂也会倒闭")
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def lab23_template():
    """实验 23 报告模板（第 19 轮 J-7）：一组一份"""
    from docx.shared import Pt
    d = _doc()
    d.add_heading("实验 23　两周情景：一张急单的全厂决策　课程设计报告", 0)

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    table([["班级", "", "小组", ""], ["日期", "", "订单号", ""]], head=False)
    d.add_heading("一、小组分工（任务 1）", 1)
    table([["姓名 / 学号", "角色", "负责的幕"], ["", "", ""], ["", "", ""], ["", "", ""], ["", "", ""], ["", "", ""], ["", "", ""]])
    d.add_heading("二、八幕的决定（任务 2）", 1)
    rows = [["幕", "发生什么", "厂情中枢报", "各方案多花 / 交期 / 风险", "我们选", "理由", "中枢推荐", "一样吗"]]
    for n, t in ((1, "客户下急单"), (2, "缺料"), (3, "来料不合格"), (4, "缺人"), (5, "安全隐患"), (6, "保养逾期"), (7, "空转与峰时"), (8, "应收逾期")):
        rows.append([str(n), t, "", "", "", "", "", ""])
    table(rows)
    d.add_heading("三、执行与单据（任务 3）", 1)
    table([["幕", "动作编号", "发起人", "批准", "ERPNext 里的单据"], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""]])
    d.add_heading("四、复盘：这张订单赚了多少（任务 4）", 1)
    table([["项目", "金额 $", "来自哪一幕"], ["收入", "", ""], ["料", "", ""], ["工（折旧 + 人工）", "", ""], ["电", "", ""], ["空转电（分摊）", "", ""],
           ["废品返工", "", ""], ["急单、加班、外协、加急", "", ""], ["保修", "", ""], ["毛利 / 毛利率", "", ""], ["按标准的毛利率", "", ""]])
    hint("把因果链截图贴在这里：从“应收逾期”或订单毛利一路追回到哪几个源头异常。")
    d.add_heading("五、如果重来", 1)
    hint("会改哪一幕的决定？估一估毛利、交期会怎么变。")
    d.add_paragraph("")
    d.add_heading("六、讨论", 1)
    hint("哪个决定对毛利影响最大；交期、成本、安全冲突时怎么排先后；真实工厂里哪些钱追不到源头")
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def robot_template(lab):
    """实验 24、25 报告模板（第 47 轮，机器人厂）"""
    from docx.shared import Pt
    d = _doc()

    def table(rows, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
                if head and r == 0:
                    for run in t.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    def hint(text):
        p = d.add_paragraph(text)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    if lab == "lab24":
        d.add_heading("实验 24　一台协作机器人从订单到交付　实验报告", 0)
        table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
        d.add_heading("一、订单与算料（任务 1）", 1)
        table([["项目", "记录"], ["销售订单号、客户、交期", ""], ["工单张数，分别是哪些物料", ""], ["关节模组：大 / 中 / 小各几台", ""],
               ["请购了什么（没有请购的说明为什么库存够）", ""]])
        d.add_heading("二、物料清单与下达顺序（任务 2）", 1)
        hint("画出 WQ-CR5 的多层物料清单（至少到关节壳体、空心轴这一层），在每个自制件、部件旁写工单号和下达时刻。")
        table([["层", "物料", "工单号", "下达时刻", "它用到的下一层完工入库时刻"], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""]])
        d.add_heading("三、关节测试与整机标定（任务 3、4）", 1)
        table([["关节序列号", "背隙（′）", "空载摩擦力矩（N·m）", "零位残差（″）", "判定与理由"], ["", "", "", "", ""], ["", "", "", "", ""]])
        table([["整机", "位置重复精度（mm）", "绝对定位误差（mm）", "入库时间", "销售发票号、金额"], ["WQ-CR5", "", "", "", ""]])
        d.add_heading("四、成本（任务 5）", 1)
        table([["工单（物料）", "标准成本", "实际成本", "差额", "差在哪里（调整、换刀、超时）"], ["", "", "", "", ""], ["", "", "", "", ""],
               ["整台（含材料）", "", "", "", ""]])
        d.add_heading("五、复盘与讨论（任务 6）", 1)
        hint("和实验 7（减速器厂的一根输出轴）比，造一台机器人多了哪些环节、哪些风险？如果整机交期提前一周，你会先动哪一层？")
    else:
        d.add_heading("实验 25　关节背隙异常的追因与 8D　实验报告", 0)
        table([["班级", "", "组号", ""], ["姓名 / 学号", "", "日期", ""]], head=False)
        d.add_heading("一、发现（任务 1）", 1)
        table([["项目", "记录"], ["出问题的特性（物料 · 特性）", ""], ["触发的判异规则、从第几台开始", ""], ["当时的超差台数", ""]])
        hint("贴控制图截图，标出触发判异的那几点。")
        d.add_heading("二、分层（任务 2）", 1)
        table([["分组方式", "组", "台数", "均值", "超差台数"], ["装配工位", "ja-1", "", "", ""], ["", "ja-2", "", "", ""],
               ["谐波减速器批次", "", "", "", ""], ["", "", "", "", ""], ["测试顺序（前一半 / 后一半）", "", "", "", ""]])
        table([["项目", "记录"], ["结论：问题出在哪一组", ""], ["依据（哪两组的差多少，远大于随机散布吗）", ""]])
        d.add_heading("三、8D（任务 3–5）", 1)
        table([["步骤", "内容"], ["D1 小组", ""], ["D2 问题描述（5W2H，带数据）", ""], ["D3 临时措施（已做的关节怎么办）", ""],
               ["D4 根本原因（5 Why）", ""], ["D5 纠正措施（选了哪一项，为什么别的不对症）", ""],
               ["D6 验证：措施前后台数、均值、超差、Cpk", ""], ["D7 预防再发生（改了哪些文件、加了什么防错）", ""], ["D8 总结", ""]])
        d.add_heading("四、讨论", 1)
        hint("如果三种原因同时存在，分层表会是什么样子？出厂检验 100% 复测能代替纠正措施吗？")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
