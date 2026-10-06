# -*- coding: utf-8 -*-
"""流体分析的 Word 报告（第 16 轮 F6）：设置、方法、结果卡片、三维截图（浏览器截）、每个面的 h 表、AI 解释、局限。"""
import base64
import datetime as dt
import io

DIR = {"+x": "+X", "-x": "−X", "+y": "+Y", "-y": "−Y", "+z": "+Z", "-z": "−Z"}


def build(job, aux, images=(), ai_text=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    st, setup = job["stats"], job.get("setup") or {}
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for name in ("Title", "Heading 1"):
        d.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)

    def table(rows):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for k, v in enumerate(r):
                t.cell(i, k).text = str(v)
                if i == 0:
                    for run in t.cell(i, k).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    d.add_heading("流体分析报告：{}".format(job.get("title") or "零件"), 0)
    d.add_paragraph("任务 {}　提交人 {}　{}　问渠数字工厂 · 流体分析（OpenFOAM）".format(
        job["id"], job.get("owner_name") or "", dt.datetime.fromtimestamp(job.get("created") or 0).strftime("%Y-%m-%d %H:%M")))

    d.add_heading("一、设置", 1)
    fl, ht = setup.get("flow") or {}, setup.get("heat")
    ex = setup.get("example")
    table([["项目", "取值"],
           ["零件", (job.get("item") or "") + ("示例 {}".format(ex["key"]) if ex else "") or "上传的 STEP"],
           ["来流", "{} m/s，方向 {}，空气 {} ℃".format(fl.get("velocity_ms"), DIR.get(fl.get("direction"), fl.get("direction")), fl.get("air_t_c"))],
           ["摆放", "放在地面 / 机座上（底面不吹风）" if (setup.get("domain") or {}).get("ground") else "悬空，四面来风"],
           ["换热", "零件表面定温 {} ℃".format(ht.get("wall_t_c")) if ht else "不算"],
           ["网格", "{} 个单元（背景 {} mm，表面加密 {} 级）".format(st.get("cells"), st.get("cell_mm"), st.get("level"))],
           ["流动模型", "层流" if st.get("turbulence") == "laminar" else "k-ω SST 湍流模型 + 壁面函数"],
           ["迭代", "{} 步，{}".format(st.get("iterations"), "收敛" if st.get("converged") else "到步数上限（平均量可用，局部值看趋势）")]])
    d.add_paragraph("方法：STEP → 按 CAD 面分开的表面三角形 → 长方体虚拟风洞（来流前 2 倍、后 5 倍、侧面 1.5 倍特征尺寸）→ "
                    "blockMesh 背景网格 + snappyHexMesh 贴体加密 → simpleFoam 定常不可压流场 → 在流场上自写有限体积解温度"
                    "（迎风格式，等效扩散 ν/Pr + ν_t/Pr_t）→ 按面积分得每个面的压力、摩擦力、散热量；h = Q /（A·ΔT）。")

    d.add_heading("二、结果", 1)
    rows = [["量", "数值"],
            ["雷诺数 Re（特征尺寸 {} mm）".format(st.get("L_char_mm")), "{:.0f}".format(st.get("Re") or 0)],
            ["阻力 N", "{:.4f}（压差 {:.4f} + 摩擦 {:.4f}）".format(st.get("drag_n") or 0, st.get("drag_pressure_n") or 0, st.get("drag_friction_n") or 0)],
            ["阻力系数 C_d", "{:.3f}（迎风面积 {:.1f} cm²、动压 {:.2f} Pa）".format(st.get("Cd") or 0, (st.get("frontal_m2") or 0) * 1e4, st.get("q_dyn_pa") or 0)]]
    if st.get("heat_w") is not None:
        rows += [["平均散热系数 h W/(m²·K)", "{:.2f}".format(st["h_mean"])],
                 ["总散热量 W", "{:.2f}（表面 {:.0f} cm²，{} ℃ 对 {} ℃）".format(st["heat_w"], (st.get("area_m2") or 0) * 1e4, st.get("wall_t_c"), st.get("air_t_c"))]]
    rows.append(["计算时间", "{} 秒".format(st.get("seconds"))])
    table(rows)

    for img in images:
        try:
            raw = base64.b64decode(img["data"].split(",", 1)[-1])
        except Exception:  # noqa: BLE001
            continue
        d.add_picture(io.BytesIO(raw), width=Cm(15.5))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = d.add_paragraph(img.get("caption") or "")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].font.size = Pt(9) if cap.runs else None

    faces = sorted(aux.get("faces") or [], key=lambda f: -f["area_mm2"])[:40]
    if faces:
        d.add_heading("三、每个面", 1)
        head = ["面", "面积 mm²", "平均压力 Pa", "压力合力 N"] + (["h W/(m²·K)", "散热 W"] if st.get("heat_w") is not None else [])
        rows = [head]
        for f in faces:
            r = ["面 {}".format(f["id"]), "{:.0f}".format(f["area_mm2"]), "{:.2f}".format(f["p_pa"]),
                 "{:.4f}".format(sum(x * x for x in (f.get("force_n") or [0])) ** 0.5)]
            if st.get("heat_w") is not None:
                r += ["{:.1f}".format(f.get("h_w_m2k") or 0), "{:.2f}".format(f.get("heat_w") or 0)]
            rows.append(r)
        table(rows)
        d.add_paragraph("压力是相对来流的表压（迎风面为正）。按面积排序，最多列 40 个面。")

    if ai_text:
        d.add_heading("四、结果解释", 1)
        for line in ai_text.split("\n"):
            if line.strip():
                d.add_paragraph(line.strip())

    d.add_heading("五、局限", 1)
    for t in ("定常、不可压、空气常物性；不算浮力（自然对流）和辐射；风扇叶片用均匀来流代替。",
              "网格为教学版（上限 30 万单元），阻力和 h 看前两三位有效数字；风洞侧面 1.5 倍特征尺寸，阻塞较大时阻力系数偏大几个百分点。",
              "方法已用四道标准题核对（平行平板层流压降、平板层流换热、湍流通道摩擦与换热、绕圆柱阻力），误差 3% 以内。"):
        d.add_paragraph(t, style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
